"""Offline build validation: readback, DependencyGraph recompute, every finding.

Engine (B) check. ``validate_offline_build_artifacts`` evaluates every build
finding — compile stage, workbook readback, DependencyGraph recompute,
projected integrity, and, when the publisher supplies one, the external
``NativeRecalculation`` of the rendered workbook — and returns the Store's
``OfflineValidation`` carrying the one ``ValidationReport``
(``schema/build_model_product_contracts.py``).
``workbook_readback_mismatches`` compares the rendered workbook's cells to the
semantic model; ``DependencyGraph.compute_typed`` is the in-process evaluator
of the semantic model, and only a ``NativeRecalculation`` evaluates the
rendered formulas themselves (the report's recalculation-evidence fields are
``None`` without one).
"""

from __future__ import annotations

from datetime import datetime
from dataclasses import dataclass
import io
import math
import numbers
from pathlib import Path
import re
import time
from typing import Any, Callable, Collection, Iterable, Mapping, NamedTuple, Sequence

from .build_findings import (
    CompileStageEvidence,
    apply_declinations,
    bear_case_above_spot_findings,
    bounded_finding_message,
    collapse_findings,
    compile_stage_findings,
    diagnostic_report_payload,
    fold_semantic_value_findings,
    judgment_info_findings,
    make_finding,
    source_diagnostic_findings,
)
from .build_model_product_contracts import (
    BuildDiagnostic,
    ModelSpec,
    OfflineValidation,
    ProvenanceMap,
    SourceBundle,
    ValidationReport,
    _COMPILED_LITERAL_CONTROL_IDS,
    _WORKBOOK_METADATA_ITEM_IDS,
    decode_source_bundle_payload,
)
from .finding_vocabulary import (
    Declination,
    DeclinationAuthor,
    Finding,
    FindingBasis,
    FindingScalar,
    computed_model_findings,
    derive_build_verdict,
)
from .models import FormulaSpec, FormulaType, ItemType, LineItemRef, shift_period
from .presentation_captions import WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
from .projected_financial_integrity import PrintedUnit
from .refs import line_item_ref_from_obj
from .renderer_columns import AbsoluteColumnMapper
from value_semantics_core import CONVENTIONS_IR_V1

_SEMANTIC_TOLERANCE = 1e-9


def provenance_diagnostics_for_map(
    provenance: ProvenanceMap,
    *,
    render_plan: Any | None = None,
    model: Any | None = None,
) -> tuple[BuildDiagnostic, ...]:
    """Check the ``ProvenanceMap`` emitted by ``target_build.project_provenance_map``.

    Every rendered cell with ``lineage_status == 'missing_source'`` emits one
    ``missing_source`` ``BuildDiagnostic``; ``source_absence_evidenced``, cited
    inclusions and ``business_model_node_unresolved`` (the node's own finding is the
    record) are never a missing-source diagnostic.
    """
    period_by_cell: dict[tuple[str, str], int] = {}
    required_role_by_item: dict[str, str | None] = {}
    if render_plan is not None:
        for entry in getattr(render_plan, "semantic_cells", ()):
            if getattr(entry, "period", 0):
                period_by_cell[(str(entry.sheet), str(entry.cell).upper())] = int(entry.period)
    if model is not None:
        for sheet_name, sheet in model.sheets.items():
            layout = sheet.layout
            mapper = AbsoluteColumnMapper(
                model.time_structure,
                layout.first_data_column if layout and layout.first_data_column else "B",
                period_scope=layout.period_scope if layout else "all",
            )
            col_to_period = {
                mapper.col_for_period(period): int(period)
                for period in mapper.all_periods()
            }
            for section in sheet.sections:
                for item in section.line_items:
                    required_role_by_item[item.id] = item.required_role
                    if item.column is not None:
                        periods = (
                            tuple(item.formula_periods or ())
                            or tuple(model.time_structure.projection_periods)
                            or tuple(mapper.all_periods())
                        )[:1]
                        if periods:
                            period_by_cell.setdefault(
                                (sheet_name, f"{item.column}{item.row}".upper()),
                                int(periods[0]),
                            )
                    else:
                        for col_letter, period in col_to_period.items():
                            period_by_cell.setdefault(
                                (sheet_name, f"{col_letter}{item.row}".upper()),
                                period,
                            )
    diagnostics: list[BuildDiagnostic] = []
    for cell in provenance.cells:
        if cell.lineage_status != "missing_source":
            continue
        period = period_by_cell.get((str(cell.sheet), str(cell.cell).upper()))
        formula = cell.formula
        observed: dict[str, FindingScalar] = {
            "formula_type": formula.type.value if formula is not None else None,
            "formula_subtype": formula.subtype if formula is not None else None,
            "formula_note": formula.note if formula is not None else None,
            "compiler_origin": (
                cell.compiler_origin.kind if cell.compiler_origin is not None else None
            ),
            "required_role": required_role_by_item.get(cell.semantic_item_id),
        }
        diagnostics.append(
            BuildDiagnostic(
                code="missing_source",
                message=(
                    "Semantic cell rests only on an accepted historical "
                    "driver seed whose operands are analyst literals with "
                    "no numeric source evidence."
                    if cell.historical_driver_seed_refs
                    else bounded_finding_message(
                        f"Semantic cell {cell.semantic_item_id} lacks resolvable "
                        "observation, assumption, or source-evidence provenance: ",
                        f"formula_type={observed['formula_type']}; "
                        f"formula_subtype={observed['formula_subtype']}; "
                        f"compiler_origin={observed['compiler_origin']}; "
                        f"required_role={observed['required_role']}; "
                        f"formula_note={observed['formula_note']}.",
                    )
                ),
                semantic_item_id=cell.semantic_item_id,
                period_id=f"FY{period}" if period is not None else None,
                observed=observed,
            )
        )
    return tuple(diagnostics)



class _SemanticCell(NamedTuple):
    sheet: str
    cell: str
    item: Any
    period: int
    spec: Any


def _semantic_cells(
    model: Any,
    render_plan: Any | None,
    *,
    owned: Callable[[Any, int], bool] | None = None,
) -> list[_SemanticCell]:
    """Document semantic coordinates, narrowed to actual writes when a plan is supplied.

    ``owned(item, period)`` narrows a document without a plan to the coordinates the document
    owns as cells (``publish.document_owns_cell``), the set a render plan would write."""
    raw_writes = getattr(render_plan, "writes", ())
    writes = (
        {(str(sheet), str(cell).upper()): value for (sheet, cell), value in raw_writes.items()}
        if isinstance(raw_writes, Mapping)
        else {(write.sheet, write.cell.upper()): write.value for write in raw_writes}
    )
    historical = set(model.time_structure.historical_periods)
    projection = set(model.time_structure.projection_periods)
    cells: list[_SemanticCell] = []
    for sheet_name, sheet in model.sheets.items():
        layout = sheet.layout
        mapper = AbsoluteColumnMapper(
            model.time_structure,
            layout.first_data_column if layout and layout.first_data_column else "B",
            period_scope=layout.period_scope if layout else "all",
        )
        for section in sheet.sections:
            for item in section.line_items:
                if (
                    item.item_type == ItemType.header
                    or item.id == "tpl.a.header.year_header"
                ):
                    continue
                periods = tuple(mapper.all_periods())
                if item.column is not None:
                    periods = (
                        tuple(item.formula_periods or ())
                        or tuple(model.time_structure.projection_periods)
                        or periods
                    )[:1]
                for period in periods:
                    cell = (
                        f"{item.column}{item.row}"
                        if item.column is not None
                        else f"{mapper.col_for_period(period)}{item.row}"
                    ).upper()
                    if render_plan is not None and (sheet_name, cell) not in writes:
                        continue
                    if owned is not None and not owned(item, period):
                        continue
                    spec = (item.overrides or {}).get(period)
                    if spec is None and (
                        item.formula_periods is None or period in item.formula_periods
                    ):
                        spec = (
                            item.historical
                            if period in historical
                            else item.projected
                            if period in projection
                            else item.projected or item.historical
                        )
                    if (
                        spec is None
                        and item.id in _COMPILED_LITERAL_CONTROL_IDS
                        and (sheet_name, cell) in writes
                    ):
                        spec = FormulaSpec(
                            type=FormulaType.constant,
                            params={"value": writes[(sheet_name, cell)]},
                        )
                    cells.append(_SemanticCell(sheet_name, cell, item, period, spec))
    return cells


def _cascade_root_cells(
    model: Any,
    required_unavailable: Collection[tuple[str, int]],
) -> frozenset[tuple[str, int]]:
    """The compact root domain shared by check and publication, without rendering.

    An owned blank input is not a published provenance cell. Reuse the
    renderer's input selector; explicit required gaps are its other write arm.
    """

    from .renderer import cell_render_input

    return frozenset(
        (item.id, int(period))
        for _sheet, _cell, item, period, _spec in _semantic_cells(model, None)
        if item.item_type not in {ItemType.header, ItemType.spacer}
        and (
            (item.id, int(period)) in required_unavailable
            or (
                (value := cell_render_input(model, item, int(period))) is not None
                and value != ""
            )
        )
    )


def _formula_resolves_to_text_metadata(
    model: Any,
    item: Any,
    period: int,
    spec: FormulaSpec,
    *,
    visiting: frozenset[tuple[str, int]] = frozenset(),
) -> bool:
    """Return whether a formula is only presentation text, never a number."""
    identity = (str(item.id), int(period))
    if identity in visiting:
        return False
    if spec.type is FormulaType.constant:
        return isinstance((spec.params or {}).get("value"), str)
    if spec.type is not FormulaType.ref:
        return False

    params = spec.params or {}
    if set(params) != {"source"}:
        return False
    source = line_item_ref_from_obj(params.get("source"))
    if source is None:
        return False
    shifted = shift_period(
        int(period),
        int(source.t),
        model.time_structure.period_mode,
    )
    if shifted is None:
        return False
    try:
        source_item = model.get_item(source.id)
    except KeyError:
        return False

    source_spec = (source_item.overrides or {}).get(int(shifted))
    if source_spec is None:
        if (
            source_item.formula_periods is not None
            and int(shifted) not in source_item.formula_periods
        ):
            return False
        historical = set(model.time_structure.historical_periods)
        projection = set(model.time_structure.projection_periods)
        source_spec = (
            source_item.historical
            if int(shifted) in historical
            else source_item.projected
            if int(shifted) in projection
            else source_item.projected or source_item.historical
        )
    if source_spec is None:
        return False
    return _formula_resolves_to_text_metadata(
        model,
        source_item,
        int(shifted),
        source_spec,
        visiting=visiting | {identity},
    )


def owned_cell_gaps(model: Any, source_bundle: SourceBundle | None) -> dict[tuple[str, int], object]:
    """Every coordinate a row owns beyond its ``formula_periods``, the ``gaps`` of ``item_owns_cell``.

    The typed lag-input gaps (``lag_input_gap_readers``), the pinned profile's cited
    ``not_applicable`` cells (``evidenced_absence_cells``) and the history a role with an authored
    forecast still owes (``owed_history_cells``): each leaves a coordinate with no formula
    that the build still renders and the check still enumerates. No bundle declares none.
    """
    if source_bundle is None:
        return {}
    from .build_required_lines import evidenced_absence_cells
    from .required_role_catalog import load_required_role_catalog
    from .role_disposition import items_by_role

    profile = source_bundle.coverage_profile
    cited = evidenced_absence_cells(
        profile.role_dispositions,
        items_by_role(load_required_role_catalog(), ticker=profile.focal_ticker),
    )
    return {
        **lag_input_gap_readers(model, source_bundle),
        **dict.fromkeys(cited),
        **dict.fromkeys(owed_history_cells(model, source_bundle)),
    }


def owed_history_cells(model: Any, source_bundle: SourceBundle) -> frozenset[tuple[str, int]]:
    """The blank historical cells of a role whose forecast the model authors but nothing declares history for.

    ``resolve_build_dispositions`` states that role ``unbound(authored_forecast_without_history)``: the
    forecast participates and its history stays owed and missing. A line of the role that owns no
    formula, override or value for such a year is the one root of every reader's non-known cell there, so
    its coordinate is a cell: required, ``#N/A``, and the subject of the finding the readers fold into.
    A year the pin binds, cites absent or cites as an inclusion has its cell written, and a role no
    authored forecast owns is ``unowed``: neither is owed here.
    """
    from .renderer import _spec_for_period
    from .required_role_catalog import load_required_role_catalog
    from .role_disposition import items_by_role, resolve_build_dispositions

    role_items = items_by_role(load_required_role_catalog(), ticker=str(model.company.ticker))
    table = resolve_build_dispositions(
        source_bundle.coverage_profile.role_dispositions, model, role_items,
        projected=tuple(model.time_structure.projection_periods),
    )
    cells: set[tuple[str, int]] = set()
    for row in table.rows:
        if row.state != "unbound" or row.unbound_reason != "authored_forecast_without_history":
            continue
        period = int(row.period_id.removeprefix("FY"))
        for item_id in role_items[row.required_role].all_items():
            item = model._index.get(item_id)
            if item is None or _spec_for_period(model, item, period) is not None:
                continue
            if item.values is not None and period in item.values.values:
                continue
            cells.add((item_id, period))
    return frozenset(cells)


def lag_input_gap_readers(
    model: Any, source_bundle: SourceBundle | None
) -> dict[tuple[str, int], BuildDiagnostic]:
    """Map each historical cell an off-axis data gap leaves formula-less to that gap.

    A gap is a typed diagnostic of the SourceBundle; no bundle (inputs unresolved) declares none.
    """
    if source_bundle is None:
        return {}
    from .renderer import _spec_for_period
    from .required_role_catalog import _ROLE_RE as _CATALOG_ROLE_RE
    from .required_role_catalog import load_required_role_catalog

    def refs(obj: Any) -> Iterable[Any]:
        ref = line_item_ref_from_obj(obj)
        if ref is not None:
            yield ref
        elif isinstance(obj, dict):
            for value in obj.values():
                yield from refs(value)
        elif isinstance(obj, (list, tuple)):
            for value in obj:
                yield from refs(value)

    historical_periods = {
        int(period) for period in model.time_structure.historical_periods
    }
    rendered_periods = historical_periods | {
        int(period) for period in model.time_structure.projection_periods
    }
    items = tuple(
        item
        for sheet in model.sheets.values()
        for section in sheet.sections
        for item in section.line_items
    )
    role_catalog = load_required_role_catalog()
    readers: dict[tuple[str, int], BuildDiagnostic] = {}
    for diagnostic in source_bundle.diagnostics:
        gap_period = re.fullmatch(r"FY(\d{4})", str(diagnostic.period_id or ""))
        if gap_period is None or int(gap_period.group(1)) in rendered_periods:
            continue
        gap_year = int(gap_period.group(1))
        gap_role = str(diagnostic.semantic_item_id or "")
        # An issuer leaf (``operating.*``) is no catalog role: its workbook item
        # is the one its pinned requirement names.
        gap_items = (
            frozenset(
                role_catalog.population_item_ids(gap_role, ticker=source_bundle.ticker)
            )
            if _CATALOG_ROLE_RE.fullmatch(gap_role)
            else frozenset(
                requirement.model_item
                for requirement in source_bundle.requirements
                if requirement.required_role == gap_role
            )
        )
        for item in items:
            for spec in (item.historical, *(item.overrides or {}).values()):
                if spec is None:
                    continue
                for ref in refs(spec.params):
                    period = gap_year - int(ref.t)
                    if (
                        ref.id in gap_items
                        and period in historical_periods
                        and _spec_for_period(model, item, period) is None
                    ):
                        readers.setdefault((str(item.id), period), diagnostic)
    return readers


def workbook_readback_mismatches(
    workbook_path: str | Path | bytes | bytearray,
    expected_writes: Mapping[tuple[str, str], object],
) -> list[tuple[str, str, object, object]]:
    """Every expected (sheet, CELL) literal/formula the workbook bytes differ on."""

    from openpyxl import load_workbook
    from openpyxl.utils.datetime import to_excel

    workbook_source = (
        io.BytesIO(bytes(workbook_path))
        if isinstance(workbook_path, (bytes, bytearray))
        else Path(workbook_path)
    )
    workbook = load_workbook(workbook_source, data_only=False, read_only=True)
    # Single pass per sheet: random cell access on a read-only workbook
    # re-parses the sheet XML on every lookup, which is quadratic in cells.
    observed_cells: dict[tuple[str, str], object] = {}
    for worksheet in workbook.worksheets:
        for row_cells in worksheet.iter_rows():
            for cell_obj in row_cells:
                if cell_obj.value is not None:
                    observed_cells[(worksheet.title, cell_obj.coordinate)] = (
                        cell_obj.value
                    )
    workbook.close()
    mismatches: list[tuple[str, str, object, object]] = []
    for (sheet, cell), expected in expected_writes.items():
        observed = observed_cells.get((sheet, cell))
        matches = observed == expected
        if (
            not matches
            and isinstance(expected, (int, float))
            and not isinstance(expected, bool)
        ):
            if isinstance(observed, datetime):
                matches = float(expected) == float(to_excel(observed))
            elif isinstance(observed, (int, float)) and not isinstance(observed, bool):
                matches = math.isclose(
                    float(expected),
                    float(observed),
                    rel_tol=1e-12,
                    abs_tol=1e-12,
                )
        if not matches:
            mismatches.append((sheet, cell, expected, observed))
    return mismatches


_FISCAL_YEAR_PERIOD_RE = re.compile(r"FY(\d{4})")


def pin_unbound_accessions(source_bundle: SourceBundle) -> dict[int, str]:
    """Each fiscal year the bundle's pin leaves unbound, with its filing.

    ``SourceBundle.unbound_periods`` is the statement-fact projection's gap
    (``model_build.corpus_facts.project_statement_facts``): a historical
    period whose as-filed filing no approved statement binding window covers —
    the same gap the HistoricalFoundation fact snapshot names
    ``period_unbound_in_pin`` — rendered as a declared gap, never a zero.
    """

    unbound: dict[int, str] = {}
    for row in source_bundle.unbound_periods:
        match = _FISCAL_YEAR_PERIOD_RE.fullmatch(row.period_id)
        if match is not None:
            unbound[int(match.group(1))] = row.accession
    return unbound


def printed_unit_dollars_by_period(
    source_bundle: SourceBundle, *, total_assets_item_id: str,
) -> dict[int, PrintedUnit]:
    """The dollar size of one printed unit of each historical column's total assets.

    The bound total-assets fact of a period is the statement fact its role's pin
    names (``AnnualSourceRequirement.pinned_columns``). The requirement is found
    by the role the required-role catalog declares for the template total-assets
    row (``statement.bs.total_assets``), because a bundle keys its requirements
    by that role: a sealed bundle may carry the role itself as ``model_item``
    where a freshly built one carries the template item id, while
    ``required_role`` names the role in both. Its printed unit is the
    precision the filer declared: XBRL ``decimals`` carried on the fact's provenance
    (``VendorProvenance.decimals``), one unit being ``10**-decimals`` dollars (``-6`` is one
    million, ``0`` one dollar, ``INF`` exact so zero) — ``basis = printed_unit``. A
    sealed bundle whose fact declares no ``decimals`` states the unit the figures
    were displayed in instead, one unit of ``display.scale`` — ``basis =
    display_scale``. The integrity evaluator converts it to the model scale of the
    row. A period whose total assets are not pinned to a filing-native fact is
    absent: there is no filed precision to attribute a non-zero delta to, so the
    evaluator reports an imbalance with ``bound_basis = none``.
    """

    from .required_role_catalog import RequiredRoleUnknown, load_required_role_catalog

    try:
        total_assets_role = load_required_role_catalog().role_for_template_item(
            total_assets_item_id
        ).required_role
    except RequiredRoleUnknown:
        return {}

    dollars = {row.numeric_scale: float(row.multiplier) for row in CONVENTIONS_IR_V1.scale_factors}
    observed = {
        observation.line: observation.fact
        for observation in source_bundle.statement_facts
        if observation.status == "ok" and observation.fact is not None
    }
    by_period: dict[int, PrintedUnit] = {}
    for requirement in source_bundle.requirements:
        if requirement.required_role != total_assets_role:
            continue
        for line in requirement.pinned_columns:
            fact = observed.get(line)
            match = _FISCAL_YEAR_PERIOD_RE.fullmatch(line.period.canonical_label())
            if fact is None or fact.grade != "filing_native" or match is None:
                continue
            decimals = None if fact.provenance is None else fact.provenance.decimals
            by_period[int(match.group(1))] = (
                PrintedUnit(dollars[fact.value.display.scale], "display_scale")
                if decimals is None
                else PrintedUnit(0.0 if decimals == "INF" else 10.0 ** -int(decimals), "printed_unit")
            )
    return by_period


def split_unbound_period_diagnostics(
    diagnostics: Iterable[BuildDiagnostic],
    *,
    unbound_periods: Collection[int],
) -> tuple[tuple[BuildDiagnostic, ...], list[tuple[str, int, str]]]:
    """Cell-lineage diagnostics split into bound ones and unbound-period facts.

    A ``missing_source`` cell in a period the pin binds nothing in becomes an
    ``(item_id, period, "missing_source")`` row for
    :func:`period_unbound_in_pin_findings`; every other diagnostic is kept.
    """

    kept: list[BuildDiagnostic] = []
    unbound_facts: list[tuple[str, int, str]] = []
    for diagnostic in diagnostics:
        period_match = _FISCAL_YEAR_PERIOD_RE.fullmatch(diagnostic.period_id or "")
        if (
            diagnostic.code == "missing_source"
            and diagnostic.semantic_item_id is not None
            and period_match is not None
            and int(period_match.group(1)) in unbound_periods
        ):
            unbound_facts.append(
                (
                    diagnostic.semantic_item_id,
                    int(period_match.group(1)),
                    "missing_source",
                )
            )
        else:
            kept.append(diagnostic)
    return tuple(kept), unbound_facts


def period_unbound_in_pin_findings(
    unbound_facts: Iterable[tuple[str, int, str]],
    *,
    unbound_accessions: Mapping[int, str],
) -> tuple[Finding, ...]:
    """One ``period_unbound_in_pin`` warning per item, naming its window.

    ``unbound_facts`` are ``(item_id, period, fact)`` rows, ``fact`` the code
    the cell would carry on a bound period (``missing_source``,
    ``semantic_value_not_known``); its period is one the pin leaves unbound
    (``unbound_accessions``: year to the filing no window covers), so the
    cause is the unbound window, not a missing disclosure. The finding names
    the window, its filings and the door that binds it: ``corpus_state``
    reads each filing (``research_file_id``, ``accession``), and a successor
    statement binding adding the window lands through
    ``fms_persist_correspondence``.
    """

    by_item: dict[str, tuple[set[int], set[str]]] = {}
    for item_id, period, fact in unbound_facts:
        periods, facts = by_item.setdefault(item_id, (set(), set()))
        periods.add(int(period))
        facts.add(fact)
    findings: list[Finding] = []
    for item_id, (periods, facts) in sorted(by_item.items()):
        window = ", ".join(f"FY{year}" for year in sorted(periods))
        accessions = ", ".join(
            f"FY{year}={unbound_accessions[year]}" for year in sorted(periods)
        )
        findings.append(
            make_finding(
                "period_unbound_in_pin",
                basis="rendered" if "missing_source" in facts else "semantic",
                message=(
                    f"No approved statement binding window covers {window} "
                    f"({accessions}), so the pin binds no value of {item_id} "
                    "there; its history is a declared gap, not a missing "
                    "disclosure. Read each filing through corpus_state "
                    "(research_file_id, accession) and propose a successor "
                    "statement binding that adds its window through "
                    "fms_persist_correspondence."
                ),
                item_id=item_id,
                periods=periods,
                observed={
                    "facts": ", ".join(sorted(facts)),
                    "accessions": accessions,
                    "next_call": "corpus_state",
                    "repair_door": "fms_persist_correspondence",
                },
            )
        )
    return tuple(findings)



def unclaimed_line_findings(
    records: Iterable[Mapping[str, Any]],
) -> tuple[Finding, ...]:
    """One blocking ``filed_line_unclaimed`` finding per build row
    (``schema.filed_line_unclaimed.unclaimed_line_record``): the line's statement, role
    and amounts as scalars, the fiscal years it carries as ``periods``."""

    from .filed_line_unclaimed import unclaimed_line_message

    return tuple(
        make_finding(
            "filed_line_unclaimed",
            message=unclaimed_line_message(record),
            item_id=str(record["line_key"]),
            periods=[int(row["fiscal_year"]) for row in record["periods"]],
            observed={
                "statement": str(record["statement"]),
                "line_qname": str(record["line_qname"]),
                "presentation_role_name": str(record["presentation_role_name"]),
                "dimensions": ",".join(
                    f"{row['axis_qname']}={row['member_qname']}" for row in record["dimensions"]
                ),
                "accessions": ",".join(record["accessions"]),
                "amounts": ", ".join(
                    f"FY{row['fiscal_year']}={row['amount']}" for row in record["periods"]
                ),
                "candidate_roles": ",".join(record["candidate_roles"]),
                "repair_owner": str(record["repair_owner"]),
            },
        )
        for record in records
    )


def _source_absence_evidenced_identities(
    provenance: Any,
    render_plan: Any | None,
    *,
    model: Any | None = None,
) -> set[tuple[str, int]]:
    """(semantic_item_id, period) cells whose lineage is a cited absence.

    Period and item come from the document's semantic coordinates. A cited absence the
    workbook never writes (a coordinate outside the row's ``formula_periods``) is a lineage
    cell with no ``render_plan.semantic_cells`` entry, so the plan only overlays the mapping.
    """

    mapped: dict[tuple[str, str], tuple[str, int]] = {}
    if model is not None:
        mapped.update(
            ((sheet, cell), (item.id, int(period)))
            for sheet, cell, item, period, _spec in _semantic_cells(model, None)
        )
    for entry in getattr(render_plan, "semantic_cells", ()):
        period = getattr(entry, "period", 0)
        if not period:
            continue
        mapped[(str(entry.sheet), str(entry.cell).upper())] = (
            str(entry.semantic_item_id),
            int(period),
        )
    identities: set[tuple[str, int]] = set()
    for cell in getattr(provenance, "cells", ()):
        if getattr(cell, "lineage_status", None) != "source_absence_evidenced":
            continue
        identity = mapped.get((str(cell.sheet), str(cell.cell).upper()))
        if identity is not None:
            identities.add(identity)
    return identities


def _unavailable_only_from_source_absence(
    value: Any,
    evidenced: set[tuple[str, int]],
) -> bool:
    """True only when every cause names a cited source-absence identity.

    An empty cause list, or any cause that is merely uncited missing, stays
    a non-known finding.
    """

    causes = getattr(value, "causes", None)
    if not causes:
        return False
    for cause in causes:
        period = getattr(cause, "period", None)
        fiscal_year = getattr(period, "fiscal_year", None)
        if fiscal_year is None:
            return False
        if (cause.node_id, int(fiscal_year)) not in evidenced:
            return False
    return True


@dataclass(frozen=True, slots=True)
class NativeRecalculation:
    """An external engine's recalculation of the rendered workbook.

    ``values`` is every cell the engine produced, keyed by upper-cased
    ``(sheet, cell)``: an error cell as its text (``#N/A``) and a formula cell
    the engine left uncomputed as ``None``. ``failure`` is set when the engine
    produced no recalculation (``values`` is then empty).
    """

    engine_name: str
    engine_version: str
    values: Mapping[tuple[str, str], Any]
    failure: str | None = None


@dataclass(frozen=True, slots=True)
class _NativeEvidence:
    """What one successful ``NativeRecalculation`` compared."""

    engine_name: str
    engine_version: str
    formula_cell_count: int
    unexpected_error_count: int
    divergence_count: int


_EXCEL_ERRORS = frozenset(
    {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
)


def _recalculated_value_matches(
    expected: object,
    observed: object,
    *,
    tolerance: float,
) -> bool:
    if (
        isinstance(expected, numbers.Real)
        and not isinstance(expected, bool)
        and isinstance(observed, numbers.Real)
        and not isinstance(observed, bool)
    ):
        return math.isclose(
            float(expected),
            float(observed),
            rel_tol=float(tolerance),
            abs_tol=_SEMANTIC_TOLERANCE,
        )
    return expected == observed


def _native_recalculation_evidence(
    native: NativeRecalculation,
    *,
    contracts: Mapping[tuple[str, str], Any],
    formula_contract_count: int,
    structural_item_ids: Collection[str],
    final_writes: Mapping[tuple[str, str], Any],
    all_semantic_contracts: Mapping[tuple[str, str], Any],
    model: Any,
    render_plan: Any,
) -> tuple[list[Finding], _NativeEvidence | None]:
    """The findings of one external recalculation and what it evaluated.

    Every semantic cell of the rendered workbook is compared with the value the
    external engine computed: a formula whose recalculated value differs from
    the render contract's expected value is a ``recalculation_divergence``; an
    error value other than an intentional ``NA()`` is ``excel_error_unexpected``;
    an engine that produced nothing is ``recalculation_engine_failed``. The
    evidence is ``None`` when the engine failed: nothing was evaluated.
    """

    if native.failure is not None:
        return (
            [
                make_finding(
                    "recalculation_engine_failed",
                    message=bounded_finding_message(
                        "External workbook recalculation failed with "
                        f"{native.engine_name} {native.engine_version}: ",
                        native.failure,
                    ),
                )
            ],
            None,
        )
    caption_literals = {
        (cell.semantic_item_id, int(cell.period)): final_writes[identity]
        for identity, cell in all_semantic_contracts.items()
        if cell.semantic_item_id in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
        and isinstance(final_writes.get(identity), str)
        and not final_writes[identity].startswith("=")
    }
    caption_specs = {
        (item.id, int(period)): spec
        for _sheet, _cell, item, period, spec in _semantic_cells(model, render_plan)
        if item.id in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS and spec is not None
    }
    unexpected_errors: list[tuple[str, str, str]] = []
    divergences: list[tuple[str, str, object, object]] = []
    for contract in contracts.values():
        if (
            contract.semantic_item_id in structural_item_ids
            and contract.semantic_item_id
            not in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
        ):
            continue
        identity = (contract.sheet.upper(), contract.cell.upper())
        recalculated = native.values.get(identity)
        if contract.semantic_item_id in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS:
            spec = caption_specs.get((contract.semantic_item_id, int(contract.period)))
            source = (
                spec.params.get("source")
                if spec is not None and spec.type.value == "ref"
                else None
            )
            source_period = (
                shift_period(
                    int(contract.period),
                    int(source.t),
                    model.time_structure.period_mode,
                )
                if isinstance(source, LineItemRef)
                else None
            )
            expected_caption = (
                caption_literals.get((source.id, int(source_period)))
                if isinstance(source, LineItemRef) and source_period is not None
                else None
            )
            if not isinstance(recalculated, str) or recalculated != expected_caption:
                divergences.append(
                    (contract.sheet, contract.cell, expected_caption, recalculated)
                )
            continue
        if isinstance(recalculated, str) and recalculated in _EXCEL_ERRORS:
            rendered_formula = final_writes.get((contract.sheet, contract.cell.upper()))
            # An authored ``NA()`` standing for a non-known value is the
            # intended error; any other error value is not.
            if not (
                recalculated == "#N/A"
                and contract.expected_state != "known"
                and isinstance(rendered_formula, str)
                and "NA()" in rendered_formula.upper()
            ):
                unexpected_errors.append((contract.sheet, contract.cell, recalculated))
            continue
        if contract.cell_kind == "unavailable":
            divergences.append((contract.sheet, contract.cell, "#N/A", recalculated))
            continue
        if contract.cell_kind != "formula" or contract.expected_state != "known":
            continue
        if recalculated is None or not _recalculated_value_matches(
            contract.expected_value,
            recalculated,
            tolerance=contract.tolerance,
        ):
            divergences.append(
                (
                    contract.sheet,
                    contract.cell,
                    contract.expected_value,
                    recalculated,
                )
            )

    findings: list[Finding] = []
    if unexpected_errors:
        sheet, cell, error = unexpected_errors[0]
        findings.append(
            make_finding(
                "excel_error_unexpected",
                message=(
                    f"{len(unexpected_errors)} unexpected Excel error value(s) "
                    f"after external recalculation; first {sheet}!{cell}: {error}."
                ),
                observed={
                    "error_count": len(unexpected_errors),
                    "first_cell": f"{sheet}!{cell}",
                    "first_error": error,
                },
            )
        )
    if divergences:
        sheet, cell, expected, observed = divergences[0]
        findings.append(
            make_finding(
                "recalculation_divergence",
                message=(
                    f"{len(divergences)} externally recalculated "
                    f"value divergence(s); first {sheet}!{cell}: expected "
                    f"{expected!r}, observed {observed!r}."
                ),
                observed={
                    "divergence_count": len(divergences),
                    "first_cell": f"{sheet}!{cell}",
                },
            )
        )
    return findings, _NativeEvidence(
        engine_name=native.engine_name,
        engine_version=native.engine_version,
        formula_cell_count=formula_contract_count,
        unexpected_error_count=len(unexpected_errors),
        divergence_count=len(divergences),
    )



def _computed_working_model(
    model: Any,
    values: Mapping[str, Mapping[int, float]],
    native: _NativeEvidence | None,
    findings: Iterable[Finding],
    *,
    requires_native: bool,
) -> bool:
    """Read numeric statements, DCF and identity evidence from this evaluation.

    WORKING-MODEL-POINTER-AND-PROVENANCE-SEVERITY uses the same semantic
    criteria for document checks and publication. RENDERED-RECHECK-AT-PUBLISH
    additionally requires clean native evidence only for rendered artifacts.
    """

    from .projected_financial_integrity import (
        BALANCE_SHEET_ITEM_IDS,
        CASH_FLOW_ITEM_IDS,
        TOTAL_REVENUE_ITEM_ID,
    )

    if requires_native and (
        native is None or native.divergence_count or native.unexpected_error_count
    ):
        return False
    historical = model.time_structure.historical_periods
    projected = model.time_structure.projection_periods
    if not historical or not projected:
        return False
    anchors = (
        TOTAL_REVENUE_ITEM_ID,
        "tpl.fm.income_statement.net_income",
        *BALANCE_SHEET_ITEM_IDS,
        *CASH_FLOW_ITEM_IDS,
    )
    if not all(
        period in values.get(item_id, {})
        for item_id in anchors
        for period in (*historical, *projected)
    ):
        return False
    if not values.get("tpl.v.dcf.equity_value_per_share"):
        return False
    return not any(
        finding.severity == "error"
        and finding.code.startswith((
            "historical_balance_sheet_",
            "historical_bs_balance_",
            "projected_balance_sheet_",
            "historical_cash_",
            "projected_cash_",
        ))
        for finding in findings
    )


def _document_declinations(
    findings: Iterable[Finding], operations: Iterable[Mapping[str, Any]],
) -> tuple[Declination, ...]:
    """One ``Declination`` per (``decline`` op, finding code) the op names.

    ``operations`` are ``model_build.document.declinations(doc)``'s rows
    (``item_id``, ``periods``, ``reason``, ``edit_index``): a finding of the
    op's item whose periods lie within the op's (empty: every period) is named
    by its code. A declinable warning is answered; a finding no declination
    clears (an error, or a warning cleared by repair) stays, and the op is
    reported by ``apply_declinations`` as ``declination_unmatched`` naming it,
    so the analyst reads that the decline did not clear it. ``info`` findings
    record and are not named. ``apply_declinations`` attaches them beside the
    ModelSpec's own.
    """

    findings = tuple(findings)
    declared: list[Declination] = []
    for operation in operations:
        periods = tuple(int(period) for period in operation["periods"])
        codes = {
            finding.code
            for finding in findings
            if finding.severity != "info"
            and finding.item_id == operation["item_id"]
            and (not periods or set(finding.periods) <= set(periods))
        }
        declared.extend(
            Declination(
                declination_id=f"document-decline-{operation['edit_index']}-{code}",
                code=code,
                item_id=operation["item_id"],
                periods=periods,
                reason=operation["reason"],
                author=DeclinationAuthor(kind="agent"),
            )
            for code in sorted(codes)
        )
    return tuple(declared)


def validate_offline_build_artifacts(
    *,
    model: Any,
    render_plan: Any | None = None,
    workbook_path: str | Path | bytes | bytearray | None = None,
    source_bundle: SourceBundle,
    model_spec: ModelSpec,
    provenance: ProvenanceMap,
    annual_presentation: Any | None = None,
    validated_at: datetime,
    compile_stage: CompileStageEvidence,
    template: Any | None = None,
    provenance_diagnostics: tuple[BuildDiagnostic, ...] = (),
    bound_scenario_control_ids: tuple[str, ...] = (),
    unclaimed_filed_lines: tuple[Mapping[str, Any], ...] = (),
    native: NativeRecalculation | None = None,
    values: Mapping[str, Mapping[int, Any]] | None = None,
    controls_by_item: Mapping[str, tuple[str, ...]] | None = None,
    document_decline_ops: Sequence[Mapping[str, Any]] = (),
) -> OfflineValidation:
    """Evaluate every build finding: compile stage, readback, recalculation.

    The one evaluator for every caller: it returns the findings and the one
    verdict ``derive_build_verdict`` gives them. ``unclaimed_filed_lines`` are the build's
    ``schema.filed_line_unclaimed.unclaimed_line_record`` rows: each is one blocking
    ``filed_line_unclaimed`` finding, and a statement that has one makes the identity
    findings over it ``binding``-owned.

    With no workbook inputs, the same document/graph evaluators run on the
    model's semantic cells. Readback, annual-presentation and native evidence
    are evaluated only when those workbook inputs are supplied.
    ``values`` supplies an already computed typed document graph; publication
    computes it here when no graph result was supplied.

    ``native`` is the external recalculation of the rendered workbook, supplied
    by the publisher (``model_build.publish.compile_version`` runs LibreOffice).
    With it, every rendered formula cell is compared with the value the
    external engine computed and the report records the evaluation; without it
    (``check_model`` recalculates with the DependencyGraph only) the report's
    recalculation-evidence fields are ``None``: nothing evaluated the rendered
    formulas.

    ``document_decline_ops`` (``document.declinations(doc)``: the analyst's ``decline`` ops) answer
    the declinable warnings of the items they name beside ``model_spec.declinations``, so a check
    and a publication of one document clear identically.
    """

    from .dependency_graph import DependencyGraph, formula_owned_recompute_ids
    from .models import ItemType

    validated_source_bundle = (
        source_bundle
        if isinstance(source_bundle, SourceBundle)
        else decode_source_bundle_payload(source_bundle)
    )
    # A cell in a historical period the pin binds nothing in lacks its source
    # and its value because the window is unbound, not because a filing
    # withholds it: those cells report ``period_unbound_in_pin`` (a warning
    # naming the window and its door); the same facts on a bound period stay
    # ``missing_source`` / ``semantic_value_not_known`` errors.
    unbound_accessions = pin_unbound_accessions(validated_source_bundle)
    effective_provenance_diagnostics = (
        tuple(provenance_diagnostics)
        if provenance_diagnostics
        else provenance_diagnostics_for_map(
            provenance,
            render_plan=render_plan,
            model=model,
        )
    )
    bound_provenance_diagnostics, unbound_facts = split_unbound_period_diagnostics(
        effective_provenance_diagnostics, unbound_periods=unbound_accessions.keys()
    )
    missing_document_template = (
        render_plan is None and workbook_path is None
        and template is None and controls_by_item is not None
    )
    findings: list[Finding] = [
        # The pinned bundle's own gaps (and the compile-stage rows callers
        # append to it). ``ModelSpec.declared_gaps`` is the server's copy of
        # the same list and is not a second source.
        *source_diagnostic_findings(
            validated_source_bundle.diagnostics, basis="pinned"
        ),
        # Cell lineage with no source, projected from the rendered plan.
        *source_diagnostic_findings(bound_provenance_diagnostics, basis="rendered"),
        *compile_stage_findings(
            compile_stage, model=model, template=template, controls_by_item=controls_by_item,
        ),
        *judgment_info_findings(
            model_spec=model_spec,
            provenance=provenance,
            validated_at=validated_at,
        ),
    ]
    if missing_document_template:
        findings.append(make_finding(
            "working_document_template_unavailable",
            message="The working document has no persisted construction template; build_model binds it.",
            expected={"next_call": "build_model"},
        ))
    from .build_scenarios import _scenario_lowering_diagnostics
    from .model_kind_dispatch import (
        declared_capabilities,
        declared_scenario_controls,
        declared_semantic_unit_plane,
    )

    # Plan §3 row S4: the model kind's declared semantic-row unit plane, read
    # off the same declared input, so offline validation holds a compiled model
    # to exactly the unit contract the compile did.
    semantic_unit_plane = declared_semantic_unit_plane(validated_source_bundle)

    # Plan §3 row A1: the control-to-workbook binding this diagnostic resolves
    # against is the model kind's declared scenario plane, read off the same
    # declared input the compile reads it from. A kind that declares none
    # resolves controls only by their own item ids.
    lowering_scenario_controls = declared_scenario_controls(
        declared_capabilities(validated_source_bundle)
    )
    for lowering in _scenario_lowering_diagnostics(
        model,
        model_spec.scenarios,
        bound_control_ids=bound_scenario_control_ids,
        workbook_targets=(
            {}
            if lowering_scenario_controls is None
            else lowering_scenario_controls.workbook_target_map()
        ),
    ):
        lowering_period = (
            re.fullmatch(r"FY(\d{4})", lowering.period_id)
            if lowering.period_id is not None
            else None
        )
        findings.append(
            make_finding(
                "scenario_control_unbound",
                message=lowering.message,
                item_id=lowering.semantic_item_id,
                periods=(
                    (int(lowering_period.group(1)),)
                    if lowering_period is not None
                    else ()
                ),
            )
        )
    annual_status = getattr(annual_presentation, "status", None)
    if annual_presentation is not None and annual_status != "ready":
        annual_blockers = tuple(annual_presentation.blockers)
        grouped_blockers: dict[tuple[str, str | None, int | None], list[Any]] = {}
        for blocker in annual_blockers:
            grouped_blockers.setdefault(
                (
                    blocker.code,
                    None
                    if blocker.code == "template_residue_visible"
                    else blocker.item_id,
                    blocker.period,
                ),
                [],
            ).append(blocker)
        for (code, item_id, period), blockers in grouped_blockers.items():
            first = blockers[0]
            findings.append(
                make_finding(
                    code,
                    message=bounded_finding_message(
                        f"{len(blockers)} annual presentation blocker(s) "
                        f"{code}; first {first.reason}"
                        + (
                            f" (location: {first.location})"
                            if first.location is not None
                            else ""
                        )
                        + ": ",
                        first.detail,
                    ),
                    item_id=item_id,
                    periods=() if period is None else (period,),
                    observed={
                        "reason": first.reason,
                        "location": first.location,
                        "blocker_count": len(blockers),
                    },
                )
            )
    if not provenance.cells:
        findings.append(
            make_finding(
                "missing_source",
                basis="rendered",
                message="No semantic cell provenance was emitted.",
            )
        )

    final_writes = {
        (write.sheet, str(write.cell).upper()): write.value
        for write in getattr(render_plan, "writes", ())
    }
    readback_mismatches = (
        workbook_readback_mismatches(workbook_path, final_writes)
        if workbook_path is not None
        else ()
    )
    if readback_mismatches:
        sheet, cell, expected, observed = readback_mismatches[0]
        findings.append(
            make_finding(
                "readback_mismatch",
                message=(
                    f"{len(readback_mismatches)} RenderPlan formula/literal "
                    f"readback mismatch(es); first {sheet}!{cell}: "
                    f"expected {expected!r}, observed {observed!r}."
                ),
                observed={
                    "mismatch_count": len(readback_mismatches),
                    "first_cell": f"{sheet}!{cell}",
                },
            )
        )

    all_semantic_contracts = {
        (str(cell.sheet), str(cell.cell).upper()): cell
        for cell in getattr(render_plan, "semantic_cells", ())
    }
    semantic_contracts = {
        (str(cell.sheet), str(cell.cell).upper()): cell
        for cell in getattr(render_plan, "semantic_cells", ())
        if cell.cell_kind != "presentation"
    }
    caption_formula_contracts = {
        identity: cell
        for identity, cell in all_semantic_contracts.items()
        if cell.semantic_item_id in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
        and isinstance(final_writes.get(identity), str)
        and final_writes[identity].startswith("=")
    }
    recalculation_contracts = {
        **semantic_contracts,
        **caption_formula_contracts,
    }
    structural_item_ids = {
        item.id
        for item in model._index.values()
        if item.item_type in {ItemType.header, ItemType.spacer}
    }
    formula_contracts = [
        cell
        for cell in recalculation_contracts.values()
        if (
            cell.cell_kind in {"formula", "unavailable"}
            or cell.semantic_item_id in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
        )
        and (
            cell.semantic_item_id not in structural_item_ids
            or cell.semantic_item_id in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
        )
    ]
    broken_links = [
        (sheet, cell, formula)
        for (sheet, cell), formula in final_writes.items()
        if isinstance(formula, str)
        and formula.startswith("=")
        and ("#REF!" in formula.upper() or "[" in formula)
    ]
    if broken_links:
        sheet, cell, formula = broken_links[0]
        findings.append(
            make_finding(
                "formula_link_broken",
                message=(
                    f"{len(broken_links)} broken or external workbook link(s); "
                    f"first {sheet}!{cell}: {formula!r}."
                ),
                observed={
                    "link_count": len(broken_links),
                    "first_cell": f"{sheet}!{cell}",
                },
            )
        )

    # The rendered ``NA()`` formulas standing for a non-known value: authored
    # errors, counted from the render contract and never an evaluation.
    intentional_na_count = 0
    for contract in recalculation_contracts.values():
        if (
            contract.semantic_item_id in structural_item_ids
            and contract.semantic_item_id
            not in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
        ):
            continue
        if (
            contract.expected_state != "known"
            and isinstance(
                rendered_formula := final_writes.get(
                    (contract.sheet, contract.cell.upper())
                ),
                str,
            )
            and "NA()" in rendered_formula.upper()
        ):
            intentional_na_count += 1

    native_evidence: _NativeEvidence | None = None
    if render_plan is not None and native is not None:
        native_findings, native_evidence = _native_recalculation_evidence(
            native,
            contracts=recalculation_contracts,
            formula_contract_count=len(formula_contracts),
            structural_item_ids=structural_item_ids,
            final_writes=final_writes,
            all_semantic_contracts=all_semantic_contracts,
            model=model,
            render_plan=render_plan,
        )
        findings.extend(native_findings)

    # Document findings use authored formulas and the full semantic domain in
    # both check and publish. Rendered literals and omitted workbook writes
    # must not change the graph's inputs or hide a document invariant.
    graph = DependencyGraph()
    graph.build(model)
    graph_failed = False
    recalc_started = time.perf_counter()
    try:
        typed_computed = (
            values
            if values is not None
            else graph.compute_typed(
                {},
                recompute=formula_owned_recompute_ids(model),
            )
        )
        computed = {
            item_id: {
                period: value.value
                for period, value in by_period.items()
                if value.state == "known"
            }
            for item_id, by_period in typed_computed.items()
        }
    except (ArithmeticError, RuntimeError, ValueError) as exc:
        graph_failed = True
        findings.append(
            make_finding(
                "semantic_graph_failed",
                message=bounded_finding_message(
                    "Canonical Python dependency-graph compute failed: ", exc
                ),
            )
        )
        typed_computed = {}
        computed = {}
    recalc_wall_seconds = time.perf_counter() - recalc_started

    source_absence_evidenced = _source_absence_evidenced_identities(
        provenance, render_plan, model=model
    )
    unitless_semantic_values: list[tuple[str, int]] = []
    nonknown_semantic_values: list[tuple[str, int, str, str, tuple[str, ...]]] = []
    from .build_company_grounded_schedules import OWED_READOUT_REASON, owed_operand_roles

    # A typed reason names the door that repairs it: an owed readout (every operand cited
    # non-participating) is repaired by a correspondence that says where the operands are included.
    reason_repair: dict[str, dict[str, FindingScalar]] = {
        OWED_READOUT_REASON: {
            "next_call": "fms_persist_correspondence",
            "repair_state": "not_separately_presented",
            "repair_field": "included_in",
        },
    }
    for _sheet, _cell, item, period, spec in _semantic_cells(model, None):
        if spec is None:
            continue
        # Caption rows and formulas that resolve only to display text are
        # workbook presentation metadata, not financial semantic values.  No
        # numeric formula-bearing row may use an absent unit to escape typed
        # acceptance validation.
        if item.id in _WORKBOOK_METADATA_ITEM_IDS or _formula_resolves_to_text_metadata(
            model, item, period, spec
        ):
            continue
        if item.unit is None:
            unitless_semantic_values.append((item.id, period))
            continue
        value = typed_computed.get(item.id, {}).get(period)
        owed_roles = owed_operand_roles(
            model, spec, int(period),
            is_known=lambda member, period=period: (
                (operand := typed_computed.get(member, {}).get(period)) is not None
                and operand.state == "known"
            ),
        )
        # An evidenced absence (a declared-domain valuation outside its
        # domain, a cited absence leaf, and each dependent the absence flows
        # to) is an intentional NA(), not a missing value; the recalculation
        # check above counts its rendered #N/A as intentional. A formula that
        # reads only cited source-absence cells is the same class:
        # graph.compute_typed reports source_value_unavailable, but the
        # absence is already evidenced on the leaf. A readout rule (d) holds
        # owed is the exception: it is the readout that is unresolved,
        # whatever its operands' evidence.
        if (
            value is None
            or value.state == "known"
            or (value.state == "evidenced_absent" and not owed_roles)
            or (
                not owed_roles
                and _unavailable_only_from_source_absence(
                    value, source_absence_evidenced
                )
            )
        ):
            continue
        if int(period) in unbound_accessions:
            unbound_facts.append((item.id, int(period), "semantic_value_not_known"))
            continue
        nonknown_semantic_values.append((
            item.id, period,
            # Rule (d) holds the readout itself unresolved, so an owed
            # readout over evidenced-absent operands reports as missing,
            # exactly as its reason is remapped to the owed reason below.
            "missing" if owed_roles and value.state == "evidenced_absent" else value.state,
            OWED_READOUT_REASON if owed_roles else value.primary_reason, owed_roles,
        ))
    nonknown_by_item: dict[str, list[tuple[int, str, str, tuple[str, ...]]]] = {}
    for item_id, period, state, reason, roles in nonknown_semantic_values:
        nonknown_by_item.setdefault(item_id, []).append((period, state, reason, roles))
    for item_id, rows in nonknown_by_item.items():
        # The repair hint rides whichever row carries it: the first row may be a period whose
        # formula never held the annotation.
        _period, state, reason, roles = next((row for row in rows if row[3]), rows[0])
        findings.append(
            make_finding(
                "semantic_value_not_known",
                message=(
                    f"{len(rows)} required semantic value(s) of {item_id} are "
                    f"non-known; first is {state} ({reason})."
                ),
                item_id=item_id,
                periods=(period for period, _state, _reason, _roles in rows),
                observed={
                    "state": state,
                    "reason": reason,
                    **reason_repair.get(reason, {}),
                    **({"operand_roles": ", ".join(roles)} if roles else {}),
                },
            )
        )
    findings.extend(
        period_unbound_in_pin_findings(
            unbound_facts, unbound_accessions=unbound_accessions
        )
    )
    unitless_by_item: dict[str, list[int]] = {}
    for item_id, period in unitless_semantic_values:
        unitless_by_item.setdefault(item_id, []).append(period)
    for item_id, periods in unitless_by_item.items():
        findings.append(
            make_finding(
                "semantic_unit_missing",
                message=(
                    f"Formula-bearing semantic row {item_id!r} has no declared "
                    f"unit in {len(periods)} period(s)."
                ),
                item_id=item_id,
                periods=periods,
            )
        )

    from .build_check_projection import non_gaap_bridge_findings
    from .projected_financial_integrity import (
        BALANCE_SHEET_ITEM_IDS,
        evaluate_projected_financial_integrity,
        evaluate_rendered_historical_statement_integrity,
        evaluate_rendered_projected_statement_integrity,
        evaluate_row_definition_continuity,
        evaluate_statement_total_operand_placement,
    )
    from .templates import load_sia_generic_template

    findings.extend(unclaimed_line_findings(unclaimed_filed_lines))
    unclaimed_statements = frozenset(str(record["statement"]) for record in unclaimed_filed_lines)

    def integrity_finding(invariant: Any, *, basis: FindingBasis) -> Finding:
        observed: dict[str, FindingScalar] = {
            **invariant.observed,
            "repair_owner": invariant.repair_owner,
        }
        if invariant.statement is not None:
            observed["statement"] = invariant.statement
        if invariant.observed_delta is not None:
            observed["observed_delta"] = invariant.observed_delta
        return make_finding(
            invariant.code,
            basis=basis,
            message=bounded_finding_message(
                f"{'Recalculated' if basis == 'rendered' else 'Semantic'} "
                f"invariant {invariant.code!r} failed: ",
                invariant.detail,
            ),
            item_id=invariant.item_id,
            periods=invariant.periods,
            operands=invariant.related_item_ids,
            repair_periods=invariant.repair_periods,
            observed=observed,
            expected=(
                {} if invariant.bound is None else {"bound": float(invariant.bound)}
            ),
        )

    # Identities are evaluated on the DependencyGraph's computed values. When
    # the graph compute failed there are none, so this block is skipped and
    # ``semantic_graph_failed`` stands alone (spec §4.3); only evaluations that
    # read no computed values (row-definition continuity, below) still run.
    readback_invariants_evaluated = not graph_failed
    readback_invariant_findings: tuple[Any, ...] = ()
    recalculated_semantics: dict[str, dict[int, float]] = {}
    if readback_invariants_evaluated:
        for sheet, cell, item, period, _spec in _semantic_cells(model, None):
            value = computed.get(item.id, {}).get(period)
            if (
                isinstance(value, numbers.Real)
                and not isinstance(value, bool)
                and math.isfinite(float(value))
            ):
                recalculated_semantics.setdefault(item.id, {})[period] = float(value)
        readback_invariant_result = evaluate_projected_financial_integrity(
            model,
            computed_values=recalculated_semantics,
            projection_periods=list(model.time_structure.projection_periods),
            semantic_units=semantic_unit_plane,
            unclaimed_statements=unclaimed_statements,
        )
        if template is None and not missing_document_template:
            template = load_sia_generic_template()
        template_findings: tuple[Any, ...] = ()
        if template is not None:
            historical_readback_result = evaluate_rendered_historical_statement_integrity(
                model,
                recalculated_values=recalculated_semantics,
                template=template,
                printed_unit_dollars=printed_unit_dollars_by_period(
                    validated_source_bundle, total_assets_item_id=BALANCE_SHEET_ITEM_IDS[0],
                ),
            )
            projected_readback_result = evaluate_rendered_projected_statement_integrity(
                model,
                recalculated_values=recalculated_semantics,
                template=template,
            )
            operand_placement_result = evaluate_statement_total_operand_placement(
                model,
                template=template,
            )
            template_findings = (
                *historical_readback_result.findings,
                *projected_readback_result.findings,
                *operand_placement_result.findings,
            )
        readback_invariant_findings = (
            *readback_invariant_result.findings,
            *template_findings,
        )
        findings.extend(
            integrity_finding(invariant, basis="rendered")
            for invariant in readback_invariant_findings
        )
        # The issuer's filed adjusted-EBITDA bridge, compared stage by stage
        # with the recalculated rows; the pairing is check_model's.
        findings.extend(
            non_gaap_bridge_findings(
                validated_source_bundle.historical_basis_facts,
                recalculated_semantics,
            )
        )
        findings.extend(
            bear_case_above_spot_findings(recalculated_semantics, model_spec=model_spec)
        )
    # A row's definition at the forecast seam is read from its rendered specs,
    # so it is evaluated whether or not the recalculation succeeded.
    findings.extend(
        integrity_finding(invariant, basis="rendered")
        for invariant in evaluate_row_definition_continuity(model).findings
    )

    semantic_mismatches: list[tuple[str, int, float, float]] = []
    for sheet, cell, item, period, _spec in _semantic_cells(model, render_plan):
        expected = final_writes.get((sheet, cell))
        actual = computed.get(item.id, {}).get(period)
        if (
            item.column is None
            and isinstance(expected, (int, float))
            and not isinstance(expected, bool)
            and actual is not None
            and not math.isclose(
                float(expected),
                float(actual),
                rel_tol=_SEMANTIC_TOLERANCE,
                abs_tol=_SEMANTIC_TOLERANCE,
            )
        ):
            semantic_mismatches.append(
                (item.id, period, float(expected), float(actual))
            )
    if semantic_mismatches:
        item_id, period, expected, actual = semantic_mismatches[0]
        findings.append(
            make_finding(
                "semantic_parity_mismatch",
                message=(
                    f"{len(semantic_mismatches)} canonical Python semantic "
                    f"parity mismatch(es); first {item_id} FY{period} expected "
                    f"{expected!r}, computed {actual!r}."
                ),
                observed={
                    "mismatch_count": len(semantic_mismatches),
                    "first_item_id": item_id,
                    "first_period": period,
                },
            )
        )

    findings = fold_semantic_value_findings(
        collapse_findings(findings), typed_computed,
        root_cells=_cascade_root_cells(
            model, compile_stage.required_unavailable | compile_stage.render_unavailable,
        ),
    )
    requires_native = workbook_path is not None or render_plan is not None
    resolved = computed_model_findings(
        findings,
        computed_model=not graph_failed and _computed_working_model(
            model, recalculated_semantics, native_evidence, findings,
            requires_native=requires_native,
        ),
        basis="native" if requires_native else "graph",
    )
    collapsed = collapse_findings(
        apply_declinations(
            resolved,
            (*model_spec.declinations, *_document_declinations(resolved, document_decline_ops)),
        )
    )
    return OfflineValidation(
        report=ValidationReport(
            status=derive_build_verdict(collapsed),
            findings=collapsed,
            diagnostic_report=diagnostic_report_payload(
                compile_stage.diagnostic_report
            ),
            validated_at=validated_at,
            recalculation_engine=(
                None if native_evidence is None else native_evidence.engine_name
            ),
            recalculation_engine_version=(
                None if native_evidence is None else native_evidence.engine_version
            ),
            recalculated_formula_cell_count=(
                None if native_evidence is None else native_evidence.formula_cell_count
            ),
            intentional_na_count=intentional_na_count,
            unexpected_excel_error_count=(
                None
                if native_evidence is None
                else native_evidence.unexpected_error_count
            ),
            recalculation_divergence_count=(
                None if native_evidence is None else native_evidence.divergence_count
            ),
            readback_invariants_evaluated=readback_invariants_evaluated,
            readback_invariant_finding_count=len(readback_invariant_findings),
        ),
        recalculated_values=recalculated_semantics,
        recalculation_wall_seconds=recalc_wall_seconds,
        recalculation_input_bytes=0,
        recalculation_output_bytes=0,
    )
