"""Exact readback for immutable published ModelVersion bundles.

The generic workbook parser intentionally reconstructs a lightweight model from
visible XLSX rows.  Published ModelVersions also carry the compiler's canonical
semantic IDs in an adjacent, immutable ``semantic.json``.  This module overlays
only those canonical rows after checking their cached workbook cells byte-for-
cell semantically (within each persisted tolerance).  Arbitrary workbooks keep
the existing parser behavior.
"""

from __future__ import annotations

from functools import lru_cache
from io import BytesIO
import json
import math
from pathlib import Path
import re
from typing import Any, Sequence

from openpyxl import load_workbook
from openpyxl.compat import safe_string
from openpyxl.utils.cell import coordinate_to_tuple

from schema.model_ref import PUBLISHABLE_MODEL_STATUSES
from schema.build_model_product_contracts import (
    decode_validation_report_payload,
)
from schema.model_semantic_document import (
    ModelSemanticDocument,
    SemanticValueSnapshot,
    decode_model_semantic_document,
)
from schema.models import FinancialModel
from schema.tools_periods import _resolve_period_list
from schema.tools_summary import _formula_type


_VERSION_DIRECTORY = re.compile(r"^v([1-9][0-9]*)$")


class PublishedModelReadbackError(ValueError):
    """The adjacent published bundle failed exact readback verification."""

    error_code = "published_model_readback_mismatch"


@lru_cache(maxsize=2)
def _load_published_bundle_cached(
    semantic_path: str,
    semantic_mtime_ns: int,
    semantic_size: int,
    validation_path: str,
    validation_mtime_ns: int,
    validation_size: int,
) -> tuple[ModelSemanticDocument, dict[str, Any]]:
    del (
        semantic_mtime_ns,
        semantic_size,
        validation_mtime_ns,
        validation_size,
    )
    validation = json.loads(Path(validation_path).read_text(encoding="utf-8"))
    if not isinstance(validation, dict):
        raise PublishedModelReadbackError(
            "published ModelVersion validation report must be an object"
        )
    # A stored pre-2.2 report is rebuild-required, refused by its decoder
    # before any cell is compared.
    decode_validation_report_payload(validation)
    semantic = decode_model_semantic_document(Path(semantic_path).read_bytes())
    return semantic, validation


def _published_bundle(
    file_path: str,
) -> tuple[ModelSemanticDocument, dict[str, Any]] | None:
    workbook_path = Path(file_path)
    if workbook_path.name != "model.xlsx":
        return None
    version_match = _VERSION_DIRECTORY.fullmatch(workbook_path.parent.name)
    if version_match is None:
        return None
    semantic_path = workbook_path.with_name("semantic.json")
    validation_path = workbook_path.with_name("validation.json")
    if not semantic_path.is_file() or not validation_path.is_file():
        return None
    semantic_stat = semantic_path.stat()
    validation_stat = validation_path.stat()
    document, validation = _load_published_bundle_cached(
        str(semantic_path),
        semantic_stat.st_mtime_ns,
        semantic_stat.st_size,
        str(validation_path),
        validation_stat.st_mtime_ns,
        validation_stat.st_size,
    )
    version = int(version_match.group(1))
    model_id = workbook_path.parent.parent.name
    if document.model_id != model_id or document.model_version != version:
        raise PublishedModelReadbackError(
            "published semantic identity does not match the ModelVersion path"
        )
    if str(validation.get("status") or "") != document.validation_status:
        raise PublishedModelReadbackError(
            "published semantic and validation statuses disagree"
        )
    return document, validation


WRITER_NUMERIC_TOLERANCE: float = 1e-12


def writer_numeric_equal(
    expected: float, actual: float, tolerance: float = WRITER_NUMERIC_TOLERANCE
) -> bool:
    """Whether two numbers are the same number as the XLSX writer writes it."""
    return (
        math.isfinite(expected)
        and math.isfinite(actual)
        and (
            # Numeric XLSX writes use safe_string (%.16g), not a binary64
            # round-trip representation: the first arm is the actual stored
            # value. The second admits drift inside the one writer tolerance
            # (rel 1e-12, abs max(the persisted tolerance, 1e-12)) and
            # nothing past it, so a cell declaring 0.0 still admits
            # one-ULP writer rounding.
            actual == float(safe_string(expected))
            or math.isclose(
                actual,
                expected,
                rel_tol=WRITER_NUMERIC_TOLERANCE,
                abs_tol=max(float(tolerance), WRITER_NUMERIC_TOLERANCE),
            )
        )
    )


def values_match(expected: Any, actual: Any, tolerance: float) -> bool:
    if isinstance(expected, bool) or isinstance(actual, bool):
        return expected is actual
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return writer_numeric_equal(float(expected), float(actual), tolerance)
    return expected == actual


def _sheet_cells(
    workbook: Any, wanted: dict[str, dict[tuple[int, int], str]]
) -> dict[tuple[str, str], Any]:
    """Every wanted cell's value, each sheet streamed once.

    A ``read_only`` worksheet re-parses its XML from the top on every
    ``ws[cell]``, so addressing 3,000 snapshots one by one costs tens of
    seconds per read (PAYC ``:8464``: 25–30 s a door call); one pass over the
    bounding box of the wanted cells costs tens of milliseconds.
    """

    values: dict[tuple[str, str], Any] = {}
    for sheet, coordinates in wanted.items():
        max_row = max(row for row, _ in coordinates)
        max_col = max(column for _, column in coordinates)
        rows = workbook[sheet].iter_rows(
            min_row=1, max_row=max_row, max_col=max_col, values_only=True
        )
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                cell = coordinates.get((row_index, column_index))
                if cell is not None:
                    values[(sheet, cell)] = value
    return values


def _verify_snapshots(
    workbook_source: str | bytes,
    snapshots: Sequence[SemanticValueSnapshot],
    validation: dict[str, Any],
) -> None:
    workbook_input = (
        BytesIO(workbook_source) if isinstance(workbook_source, bytes) else workbook_source
    )
    workbook = load_workbook(workbook_input, read_only=True, data_only=True)
    try:
        wanted: dict[str, dict[tuple[int, int], str]] = {}
        for snapshot in snapshots:
            if snapshot.sheet not in workbook.sheetnames:
                raise PublishedModelReadbackError(
                    f"published semantic sheet is missing: {snapshot.sheet}"
                )
            wanted.setdefault(snapshot.sheet, {})[coordinate_to_tuple(snapshot.cell)] = snapshot.cell
        actuals = _sheet_cells(workbook, wanted)
    finally:
        workbook.close()
    unmatched = [
        snapshot
        for snapshot in snapshots
        if not values_match(snapshot.value, actuals.get((snapshot.sheet, snapshot.cell)), snapshot.tolerance)
    ]
    if not unmatched:
        return
    # A cell the cached workbook holds no value for is validated by its formula
    # when the report recalculated the workbook clean.
    formula_input = (
        BytesIO(workbook_source) if isinstance(workbook_source, bytes) else workbook_source
    )
    formulas = load_workbook(formula_input, read_only=True, data_only=False)
    try:
        wanted = {}
        for snapshot in unmatched:
            wanted.setdefault(snapshot.sheet, {})[coordinate_to_tuple(snapshot.cell)] = snapshot.cell
        formula_cells = _sheet_cells(formulas, wanted)
    finally:
        formulas.close()
    recalculated_clean = (
        validation.get("status") in PUBLISHABLE_MODEL_STATUSES
        and validation.get("readback_invariants_evaluated") is True
        and validation.get("recalculation_divergence_count") == 0
        and validation.get("unexpected_excel_error_count") == 0
    )
    for snapshot in unmatched:
        formula = formula_cells.get((snapshot.sheet, snapshot.cell))
        validated_formula = (
            actuals.get((snapshot.sheet, snapshot.cell)) is None
            and snapshot.value is not None
            and isinstance(formula, str)
            and formula.startswith("=")
            and recalculated_clean
        )
        if not validated_formula:
            raise PublishedModelReadbackError(
                "published semantic value does not match cached workbook cell: "
                f"{snapshot.semantic_item_id} {snapshot.period} "
                f"({snapshot.sheet}!{snapshot.cell})"
            )


def _file_cache_signature(path: Path) -> tuple[int, int]:
    try:
        stat = path.stat()
    except OSError:
        return (-1, -1)
    return (stat.st_mtime_ns, stat.st_size)


@lru_cache(maxsize=2)
def _verified_published_bundle_cached(
    file_path: str,
    workbook_mtime_ns: int,
    workbook_size: int,
    semantic_mtime_ns: int,
    semantic_size: int,
    validation_mtime_ns: int,
    validation_size: int,
) -> tuple[ModelSemanticDocument, dict[str, Any]] | None:
    del (
        workbook_mtime_ns,
        workbook_size,
        semantic_mtime_ns,
        semantic_size,
        validation_mtime_ns,
        validation_size,
    )
    bundle = _published_bundle(file_path)
    if bundle is None:
        return None
    document, validation = bundle
    _verify_snapshots(file_path, document.semantic_values, validation)
    return document, validation


def _verified_published_bundle(
    file_path: str,
) -> tuple[ModelSemanticDocument, dict[str, Any]] | None:
    workbook_path = Path(file_path)
    workbook_signature = _file_cache_signature(workbook_path)
    semantic_signature = _file_cache_signature(workbook_path.with_name("semantic.json"))
    validation_signature = _file_cache_signature(
        workbook_path.with_name("validation.json")
    )
    return _verified_published_bundle_cached(
        str(workbook_path),
        *workbook_signature,
        *semantic_signature,
        *validation_signature,
    )


def verified_published_model(file_path: str) -> FinancialModel | None:
    """Return the exact canonical model for an immutable published bundle.

    All persisted semantic snapshots are checked against the adjacent workbook
    and validation report before the model is allowed to replace the lossy XLSX
    reconstruction used for arbitrary workbooks.
    """

    bundle = _verified_published_bundle(file_path)
    return None if bundle is None else bundle[0].model


def verified_published_model_from_artifacts(
    workbook_bytes: bytes,
    semantic_document: ModelSemanticDocument,
    validation_report: Any,
) -> FinancialModel:
    """Verify exact workbook bytes against decoded immutable sidecars."""

    validation = (
        dict(validation_report)
        if isinstance(validation_report, dict)
        else validation_report.model_dump(mode="json")
    )
    if validation.get("status") != semantic_document.validation_status:
        raise PublishedModelReadbackError(
            "published semantic and validation statuses disagree"
        )
    _verify_snapshots(
        workbook_bytes,
        semantic_document.semantic_values,
        validation,
    )
    return semantic_document.model




def overlay_published_model_values(
    file_path: str,
    item_ids: Sequence[str],
    *,
    periods: str | list[int | str],
    historical_cutoff_year: int | None,
    fallback: dict[str, Any],
) -> dict[str, Any]:
    """Overlay exact canonical semantic rows when ``file_path`` is a bundle."""

    del historical_cutoff_year
    bundle = _verified_published_bundle(file_path)
    if bundle is None:
        return fallback
    document, validation = bundle
    period_list, period_label = _resolve_period_list(periods, document.model)
    selected_periods = set(period_list)
    requested = tuple(dict.fromkeys(str(item_id) for item_id in item_ids))
    snapshots_by_item: dict[str, list[SemanticValueSnapshot]] = {}
    for snapshot in document.semantic_values:
        if (
            snapshot.semantic_item_id in requested
            and snapshot.period in selected_periods
        ):
            snapshots_by_item.setdefault(snapshot.semantic_item_id, []).append(snapshot)
    if not snapshots_by_item:
        return fallback
    selected_snapshots = tuple(
        snapshot
        for item_id in requested
        for snapshot in sorted(
            snapshots_by_item.get(item_id, ()),
            key=lambda row: (row.period, row.sheet, row.cell),
        )
    )
    document.model.build_index()
    fallback_by_id = {
        str(row.get("id") or ""): dict(row)
        for row in fallback.get("items") or ()
        if isinstance(row, dict)
    }
    rows: list[dict[str, Any]] = []
    for item_id in requested:
        snapshots = snapshots_by_item.get(item_id)
        if not snapshots:
            rows.append(fallback_by_id.get(item_id, {"id": item_id}))
            continue
        item = document.model.get_item(item_id)
        rows.append(
            {
                "id": item_id,
                "label": item.label,
                "item_type": item.item_type.value,
                "formula_type": _formula_type(item),
                "values": {
                    snapshot.period: snapshot.value
                    for snapshot in sorted(snapshots, key=lambda row: row.period)
                },
                "readback_source": "published_model_semantic_sidecar",
            }
        )
    result = dict(fallback)
    result.update(
        {
            "items": rows,
            "periods_returned": period_label,
            "period_count": len(period_list),
            "published_model_readback": {
                "model_id": document.model_id,
                "model_version": document.model_version,
                "ticker": document.ticker,
                "validation_status": document.validation_status,
                "validation_report_schema_version": validation.get("schema_version"),
                "verified_item_count": len(snapshots_by_item),
                "verified_cell_count": len(selected_snapshots),
                "valuation_evidence_status": (
                    document.valuation_evidence_status.model_dump(mode="json")
                ),
            },
        }
    )
    return result




__all__ = [
    "PublishedModelReadbackError",
    "WRITER_NUMERIC_TOLERANCE",
    "overlay_published_model_values",
    "values_match",
    "verified_published_model",
    "verified_published_model_from_artifacts",
    "writer_numeric_equal",
]
