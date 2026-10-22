"""Annual formatting, fiscal-header, and residue presentation guard.

The guard evaluates; it does not refuse. A blocked result carries its blockers,
and ``schema/build_offline_validation.py`` turns each into its error Finding.
"""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Sequence
import datetime
import hashlib
from io import BytesIO
import logging
from pathlib import Path
import re
from typing import Any, Literal, Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from pydantic import BaseModel, ConfigDict, Field, field_validator

from .build_company_info import (
    _FINANCIAL_MODEL_YEAR_HEADER_ID,
    _YEAR_HEADER_ID,
    _fiscal_year_end_date,
    _gregorian_to_excel_serial,
)
from .formatter import ModelFormatter, SIA_FORMATTER
from .models import (
    FinancialModel,
    FormulaType,
    ItemType,
    LineItem,
    PERIOD_MODE_YEARLY,
)
from .presentation_captions import WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS
from industry_slice_core.industry_types import Unit
from .renderer_columns import AbsoluteColumnMapper
from .renderer_types import RenderPlan


AnnualEstimateStatus = Literal["actual", "estimate"]
AnnualPresentationStatus = Literal["ready", "blocked"]
AnnualPresentationBlockerCode = Literal[
    "unit_format_mismatch",
    "fiscal_axis_header_missing",
    "template_residue_visible",
]
AnnualPresentationReason = Literal[
    "annual_fiscal_axis_invalid",
    "fiscal_year_end_invalid",
    "header_row_missing",
    "model_header_shape_invalid",
    "render_plan_missing",
    "rendered_header_missing",
    "rendered_header_value_mismatch",
    "rendered_header_format_mismatch",
    "readback_missing",
    "readback_value_mismatch",
    "readback_format_mismatch",
    "explicit_format_invalid",
    "semantic_format_metadata_missing",
    "rendered_semantic_format_missing",
    "rendered_semantic_format_mismatch",
    "model_text_residue",
    "rendered_text_residue",
    "readback_text_residue",
]
PresentationScalar = str | int | float | bool | None


class AnnualYearHeaderReadback(BaseModel):
    """Canonical semantic expectation for one visible annual header."""

    model_config = ConfigDict(extra="forbid", frozen=True, strict=True)

    period: int
    display_label: str = Field(min_length=1)
    estimate_status: AnnualEstimateStatus


class PresentationCellReadback(BaseModel):
    """Typed workbook presentation readback for one exact visible cell."""

    model_config = ConfigDict(extra="forbid", frozen=True, strict=True)

    sheet: str = Field(min_length=1)
    cell: str = Field(pattern=r"^[A-Za-z]+[1-9][0-9]*$")
    value: PresentationScalar = None
    number_format: str | None = None


class AnnualPresentationBlocker(BaseModel):
    """One granular reason under the frozen BMA-012 blocker taxonomy."""

    model_config = ConfigDict(extra="forbid", frozen=True, strict=True)

    code: AnnualPresentationBlockerCode
    reason: AnnualPresentationReason
    detail: str = Field(min_length=1)
    period: int | None = None
    item_id: str | None = None
    location: str | None = None
    expected: PresentationScalar = None
    observed: PresentationScalar = None


class AnnualPresentationGuardResult(BaseModel):
    """Complete BMA-012 presentation-guard result."""

    model_config = ConfigDict(extra="forbid", frozen=True, strict=True)

    status: AnnualPresentationStatus
    workbook_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    render_state_hash: str = Field(pattern=r"^[0-9a-f]{64}$")
    expected_axis: tuple[int, ...]
    expected_headers: tuple[AnnualYearHeaderReadback, ...]
    blockers: tuple[AnnualPresentationBlocker, ...]

    @field_validator("expected_axis", "expected_headers", "blockers", mode="before")
    @classmethod
    def _restore_persisted_json_arrays(cls, value: object) -> object:
        """Restore immutable tuple fields from their canonical JSON arrays."""

        return tuple(value) if isinstance(value, list) else value


class AnnualPresentationBindingError(ValueError):
    """Raised when presentation evidence does not bind the exact workbook revision."""


def _model_ref_payload(model_ref: object) -> dict[str, Any]:
    if isinstance(model_ref, Mapping):
        return dict(model_ref)
    dump = getattr(model_ref, "model_dump", None)
    if callable(dump):
        payload = dump(mode="json")
        if isinstance(payload, dict):
            return payload
    raise AnnualPresentationBindingError("model_ref must be a typed object or mapping")


def _workbook_content_sha256(
    *,
    model_ref: Mapping[str, Any],
    workbook_path: str | Path | None,
    workbook_content: bytes | None,
) -> str:
    if workbook_path is not None and workbook_content is not None:
        raise AnnualPresentationBindingError(
            "provide workbook_path or workbook_content, not both"
        )
    if workbook_content is not None:
        return hashlib.sha256(workbook_content).hexdigest()
    raw_path = workbook_path or str(
        model_ref.get("file_path") or model_ref.get("workbook_path") or ""
    )
    if not str(raw_path).strip():
        raise AnnualPresentationBindingError("model_ref workbook path is required")
    resolved = Path(raw_path).expanduser()
    try:
        content = resolved.read_bytes()
    except OSError as exc:
        raise AnnualPresentationBindingError(
            f"bound annual-presentation workbook is unreadable: {resolved}"
        ) from exc
    return hashlib.sha256(content).hexdigest()


def require_bound_ready_annual_presentation(
    model_ref: object,
    *,
    workbook_path: str | Path | None = None,
    workbook_content: bytes | None = None,
) -> AnnualPresentationGuardResult:
    """Validate that model_ref contains well-formed presentation evidence.

    When workbook_path or workbook_content is explicitly provided (such as during
    publication provenance compilation), verifies that the evidence binds the exact
    workbook bytes being published.
    """
    payload = _model_ref_payload(model_ref)
    try:
        result = AnnualPresentationGuardResult.model_validate(
            payload.get("annual_presentation")
        )
    except (TypeError, ValueError) as exc:
        raise AnnualPresentationBindingError(
            "model_ref annual_presentation evidence is invalid"
        ) from exc
    if workbook_path is not None or workbook_content is not None:
        actual_workbook_hash = _workbook_content_sha256(
            model_ref=payload,
            workbook_path=workbook_path,
            workbook_content=workbook_content,
        )
        if result.workbook_sha256 != actual_workbook_hash:
            raise AnnualPresentationBindingError(
                "annual_presentation workbook_sha256 does not match workbook bytes"
            )
    return result


def _annual_presentation_readback_value(
    value: object,
    *,
    workbook_epoch: datetime.datetime,
) -> PresentationScalar:
    """Normalize one openpyxl value into the guard's strict scalar contract."""

    if isinstance(
        value,
        (datetime.datetime, datetime.date, datetime.time, datetime.timedelta),
    ):
        serial = to_excel(value, epoch=workbook_epoch)
        return float(serial) if serial is not None else None
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def read_annual_presentation_cells(
    workbook_path: str | Path | bytes | bytearray,
) -> tuple[PresentationCellReadback, ...] | None:
    """Read visible values and formatted cells from one persisted workbook."""

    try:
        workbook = load_workbook(
            (
                BytesIO(bytes(workbook_path))
                if isinstance(workbook_path, (bytes, bytearray))
                else str(workbook_path)
            ),
            data_only=False,
            read_only=False,
        )
    except Exception:
        logging.exception(
            "Annual presentation workbook readback failed for %s",
            workbook_path,
        )
        return None

    cells: list[PresentationCellReadback] = []
    try:
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            for row_index, row in enumerate(worksheet.iter_rows(), start=1):
                for column_index, cell in enumerate(row, start=1):
                    column = get_column_letter(column_index)
                    if (
                        worksheet.row_dimensions[row_index].hidden
                        or worksheet.column_dimensions[column].hidden
                    ):
                        continue
                    value = _annual_presentation_readback_value(
                        cell.value,
                        workbook_epoch=workbook.epoch,
                    )
                    number_format = cell.number_format
                    if value is None and number_format == "General":
                        continue
                    cells.append(
                        PresentationCellReadback(
                            sheet=worksheet.title,
                            cell=cell.coordinate.upper(),
                            value=value,
                            number_format=number_format,
                        )
                    )
    except Exception:
        logging.exception(
            "Annual presentation workbook readback failed for %s",
            workbook_path,
        )
        return None
    finally:
        workbook.close()
    return tuple(cells)


_BRACKETED_PLACEHOLDER_RE = re.compile(r"\[[^\[\]]+\]")
_NUMBERED_PLACEHOLDER_RE = re.compile(
    r"\b(?:current|long[- ]term)\s+(?:asset|liability)\s+\d+\b|"
    r"\b(?:business segment|volume driver|price driver|operating metric|comp)\s+\d+\b",
    re.IGNORECASE,
)
# A whole label that is a balance-sheet slot noun is template residue with or
# without its number. Matched against the full label so ordinary captions
# ("Total current assets", "Other current liabilities") never match.
_SLOT_NOUN_LABEL_RE = re.compile(
    r"(?:current|long[- ]term)\s+(?:asset|liability)(?:\s+\d+)?",
    re.IGNORECASE,
)
_INSTRUCTION_TEXT_RE = re.compile(
    r"\b(?:populate via|adjust link|replace with|template instruction|"
    r"software investment,? if applicable|if applicable,? adjust)\b",
    re.IGNORECASE,
)
_CELL_RE = re.compile(r"^\$?(?P<column>[A-Z]+)\$?(?P<row>[1-9][0-9]*)$")
_READBACK_REQUIRED_UNITS = frozenset(
    {
        Unit.percentage,
        Unit.ratio,
        Unit.multiple,
        Unit.per_share,
        Unit.share_count,
        Unit.date,
    }
)


def _annual_axis(model: FinancialModel) -> tuple[tuple[int, ...], frozenset[int]]:
    time_structure = model.time_structure
    if time_structure.period_mode != PERIOD_MODE_YEARLY:
        # Per SC5 ruling (2026-10-09/10): a non-yearly model renders its own cadence;
        # do not refuse or fabricate annual periods.
        historical = tuple(int(p) for p in time_structure.historical_periods if isinstance(p, int) or str(p).isdigit())
        projection = tuple(int(p) for p in time_structure.projection_periods if isinstance(p, int) or str(p).isdigit())
        if historical or projection:
            return (historical + projection), frozenset(projection)
        return (), frozenset()

    historical_periods = tuple(
        int(period) for period in time_structure.historical_periods
    )
    historical_years = tuple(int(period) for period in time_structure.historical_years)
    projection_periods = tuple(
        int(period) for period in time_structure.projection_periods
    )
    projection_years = tuple(int(period) for period in time_structure.projection_years)
    if (
        historical_periods
        and historical_years
        and historical_periods != historical_years
    ):
        raise ValueError("historical_periods and historical_years disagree")
    if (
        projection_periods
        and projection_years
        and projection_periods != projection_years
    ):
        raise ValueError("projection_periods and projection_years disagree")

    historical = historical_periods or historical_years
    projection = projection_periods or projection_years
    if not historical and not projection:
        raise ValueError("annual fiscal axis is empty")
    if tuple(sorted(set(historical))) != historical:
        raise ValueError("historical annual periods must be sorted and unique")
    if tuple(sorted(set(projection))) != projection:
        raise ValueError("projection annual periods must be sorted and unique")
    if set(historical).intersection(projection):
        raise ValueError("actual and estimate annual periods overlap")
    if historical and projection and historical[-1] >= projection[0]:
        raise ValueError("estimate periods must follow all actual periods")

    axis = historical + projection
    if axis != tuple(range(axis[0], axis[-1] + 1)):
        raise ValueError("annual fiscal axis must be contiguous")
    return axis, frozenset(projection)


def expected_annual_year_headers(
    model: FinancialModel,
) -> tuple[AnnualYearHeaderReadback, ...]:
    """Return the sole canonical visible header/status projection for the model."""

    axis, estimate_periods = _annual_axis(model)
    return tuple(
        AnnualYearHeaderReadback(
            period=period,
            display_label=f"{period}E" if period in estimate_periods else str(period),
            estimate_status="estimate" if period in estimate_periods else "actual",
        )
        for period in axis
    )


def _model_items(
    model: FinancialModel,
) -> tuple[tuple[str, LineItem], ...]:
    return tuple(
        (sheet_name, item)
        for sheet_name, sheet in model.sheets.items()
        for section in sheet.sections
        for item in section.line_items
    )


def _item_by_id(
    model: FinancialModel,
    item_id: str,
) -> tuple[str, LineItem] | None:
    return next(
        (
            (sheet_name, item)
            for sheet_name, item in _model_items(model)
            if item.id == item_id
        ),
        None,
    )


def _mapper_for_sheet(model: FinancialModel, sheet_name: str) -> AbsoluteColumnMapper:
    sheet = model.sheets[sheet_name]
    layout = sheet.layout
    first_data_column = (
        layout.first_data_column if layout and layout.first_data_column else "B"
    )
    period_scope = layout.period_scope if layout else "all"
    return AbsoluteColumnMapper(
        model.time_structure,
        str(first_data_column).upper(),
        period_scope=period_scope,
    )


def _period_cell(
    model: FinancialModel,
    *,
    sheet_name: str,
    item: LineItem,
    period: int,
) -> str:
    mapper = _mapper_for_sheet(model, sheet_name)
    return f"{mapper.col_for_period(period)}{item.row}".upper()


def _target_cells(
    model: FinancialModel,
    *,
    sheet_name: str,
    item: LineItem,
) -> tuple[str, ...]:
    if item.column is not None:
        return (f"{item.column}{item.row}".upper(),)
    mapper = _mapper_for_sheet(model, sheet_name)
    return tuple(
        f"{mapper.col_for_period(period)}{item.row}".upper()
        for period in mapper.all_periods()
    )


def _col_index(column: str) -> int:
    result = 0
    for character in column:
        result = result * 26 + ord(character) - ord("A") + 1
    return result


def _cell_parts(cell: str) -> tuple[int, int]:
    match = _CELL_RE.fullmatch(str(cell).upper())
    if match is None:
        raise ValueError(f"invalid A1 cell reference: {cell!r}")
    return _col_index(match.group("column")), int(match.group("row"))


NumberFormatBounds = tuple[int, int, int, int, str]
NumberFormatIndex = Mapping[str, tuple[NumberFormatBounds, ...]]


def _number_format_index(render_plan: RenderPlan) -> NumberFormatIndex:
    """Parse format ranges once instead of once per model cell."""
    grouped: dict[str, list[NumberFormatBounds]] = defaultdict(list)
    for cell_format in render_plan.formats:
        if cell_format.number_format is None:
            continue
        start_text, separator, end_text = str(cell_format.range).upper().partition(":")
        end_text = end_text if separator else start_text
        try:
            start_col, start_row = _cell_parts(start_text)
            end_col, end_row = _cell_parts(end_text)
        except ValueError:
            continue
        grouped[cell_format.sheet].append(
            (
                min(start_col, end_col),
                max(start_col, end_col),
                min(start_row, end_row),
                max(start_row, end_row),
                cell_format.number_format,
            )
        )
    return {sheet: tuple(formats) for sheet, formats in grouped.items()}


def _effective_number_format(
    format_index: NumberFormatIndex,
    *,
    sheet: str,
    cell: str,
) -> str | None:
    result: str | None = None
    try:
        cell_col, cell_row = _cell_parts(cell)
    except ValueError:
        return None
    for min_col, max_col, min_row, max_row, number_format in format_index.get(sheet, ()):
        if min_col <= cell_col <= max_col and min_row <= cell_row <= max_row:
            result = number_format
    return result


def _writes_by_cell(
    render_plan: RenderPlan,
) -> dict[tuple[str, str], tuple[PresentationScalar, ...]]:
    grouped: dict[tuple[str, str], list[PresentationScalar]] = defaultdict(list)
    for write in render_plan.writes:
        grouped[(write.sheet, write.cell.upper())].append(write.value)
    return {key: tuple(values) for key, values in grouped.items()}


def _readback_by_cell(
    readback: Sequence[PresentationCellReadback] | None,
) -> dict[tuple[str, str], tuple[PresentationCellReadback, ...]]:
    grouped: dict[tuple[str, str], list[PresentationCellReadback]] = defaultdict(list)
    for cell in readback or ():
        grouped[(cell.sheet, cell.cell.upper())].append(cell)
    return {key: tuple(values) for key, values in grouped.items()}


def _header_model_blockers(
    model: FinancialModel,
    *,
    axis: tuple[int, ...],
    estimate_periods: frozenset[int],
    formatter: ModelFormatter,
) -> list[AnnualPresentationBlocker]:
    blockers: list[AnnualPresentationBlocker] = []
    assumptions_entry = _item_by_id(model, _YEAR_HEADER_ID)
    financial_entry = _item_by_id(model, _FINANCIAL_MODEL_YEAR_HEADER_ID)
    for item_id, entry in (
        (_YEAR_HEADER_ID, assumptions_entry),
        (_FINANCIAL_MODEL_YEAR_HEADER_ID, financial_entry),
    ):
        if entry is None:
            blockers.append(
                AnnualPresentationBlocker(
                    code="fiscal_axis_header_missing",
                    reason="header_row_missing",
                    detail=f"required annual header row {item_id!r} is missing",
                    item_id=item_id,
                )
            )

    if assumptions_entry is not None:
        _sheet_name, item = assumptions_entry
        values = item.values.values if item.values is not None else {}
        if (
            tuple(values) != axis
            or item.historical is not None
            or item.projected is not None
            or item.overrides is not None
            or list(item.formula_periods or [])
            or item.unit is not Unit.date
            or formatter.number_format_for(item) != formatter.date_format
        ):
            blockers.append(
                AnnualPresentationBlocker(
                    code="fiscal_axis_header_missing",
                    reason="model_header_shape_invalid",
                    detail="Assumptions annual dates are not an exact typed-date value axis",
                    item_id=_YEAR_HEADER_ID,
                )
            )
        for period in axis:
            value_cell = values.get(period)
            expected_status = (
                "estimate" if period in estimate_periods else "actual"
            )
            expected_serial = float(
                _gregorian_to_excel_serial(
                    _fiscal_year_end_date(
                        model.time_structure.fiscal_year_end,
                        period,
                    )
                )
            )
            if (
                value_cell is None
                or value_cell.value != expected_serial
                or value_cell.note != f"annual_header_status={expected_status}"
            ):
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="model_header_shape_invalid",
                        detail=(
                            f"Assumptions header {period} must bind the exact fiscal "
                            f"date and {expected_status} status"
                        ),
                        period=period,
                        item_id=_YEAR_HEADER_ID,
                        expected=expected_serial,
                        observed=value_cell.value if value_cell is not None else None,
                    )
                )

    if financial_entry is not None:
        _sheet_name, item = financial_entry
        overrides = item.overrides or {}
        if (
            tuple(overrides) != axis
            or tuple(int(period) for period in item.formula_periods or []) != axis
            or item.format != "General"
            or item.unit is not None
            or item.historical is not None
            or item.projected is not None
            or item.values is not None
        ):
            blockers.append(
                AnnualPresentationBlocker(
                    code="fiscal_axis_header_missing",
                    reason="model_header_shape_invalid",
                    detail="Financial_model annual labels are not one exact constant axis",
                    item_id=_FINANCIAL_MODEL_YEAR_HEADER_ID,
                )
            )
        for expected in expected_annual_year_headers(model):
            spec = overrides.get(expected.period)
            if (
                spec is None
                or spec.type is not FormulaType.constant
                or set(spec.params) != {"value"}
                or spec.params.get("value") != expected.display_label
            ):
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="model_header_shape_invalid",
                        detail=(
                            f"Financial_model header {expected.period} must be the "
                            f"exact label {expected.display_label!r}"
                        ),
                        period=expected.period,
                        item_id=_FINANCIAL_MODEL_YEAR_HEADER_ID,
                        expected=expected.display_label,
                        observed=(
                            spec.params.get("value") if spec is not None else None
                        ),
                    )
                )
    return blockers


def _header_plan_and_readback_blockers(
    model: FinancialModel,
    *,
    axis: tuple[int, ...],
    render_plan: RenderPlan | None,
    readback: Sequence[PresentationCellReadback] | None,
    formatter: ModelFormatter,
    format_index: NumberFormatIndex,
) -> list[AnnualPresentationBlocker]:
    if render_plan is None:
        return [
            AnnualPresentationBlocker(
                code="fiscal_axis_header_missing",
                reason="render_plan_missing",
                detail="the exact annual render plan is required",
            )
        ]

    assumptions_entry = _item_by_id(model, _YEAR_HEADER_ID)
    financial_entry = _item_by_id(model, _FINANCIAL_MODEL_YEAR_HEADER_ID)
    if assumptions_entry is None or financial_entry is None:
        return []

    writes = _writes_by_cell(render_plan)
    readbacks = _readback_by_cell(readback)
    blockers: list[AnnualPresentationBlocker] = []
    assumptions_sheet, assumptions_item = assumptions_entry
    financial_sheet, financial_item = financial_entry
    expected_formats = {
        _YEAR_HEADER_ID: formatter.number_format_for(assumptions_item),
        _FINANCIAL_MODEL_YEAR_HEADER_ID: formatter.number_format_for(financial_item),
    }
    headers = {header.period: header for header in expected_annual_year_headers(model)}

    for period in axis:
        expected_serial = float(
            _gregorian_to_excel_serial(
                _fiscal_year_end_date(model.time_structure.fiscal_year_end, period)
            )
        )
        expected_cells = (
            (
                assumptions_sheet,
                _period_cell(
                    model,
                    sheet_name=assumptions_sheet,
                    item=assumptions_item,
                    period=period,
                ),
                _YEAR_HEADER_ID,
                expected_serial,
            ),
            (
                financial_sheet,
                _period_cell(
                    model,
                    sheet_name=financial_sheet,
                    item=financial_item,
                    period=period,
                ),
                _FINANCIAL_MODEL_YEAR_HEADER_ID,
                headers[period].display_label,
            ),
        )
        for sheet, cell, item_id, expected_value in expected_cells:
            location = f"{sheet}!{cell}"
            rendered_values = writes.get((sheet, cell), ())
            if not rendered_values:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="rendered_header_missing",
                        detail=f"render plan omits annual header {location}",
                        period=period,
                        item_id=item_id,
                        location=location,
                        expected=expected_value,
                    )
                )
            elif len(rendered_values) != 1 or rendered_values[0] != expected_value:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="rendered_header_value_mismatch",
                        detail=f"rendered annual header value differs at {location}",
                        period=period,
                        item_id=item_id,
                        location=location,
                        expected=expected_value,
                        observed=(
                            rendered_values[0]
                            if len(rendered_values) == 1
                            else repr(rendered_values)
                        ),
                    )
                )

            expected_format = expected_formats[item_id]
            rendered_format = _effective_number_format(
                format_index,
                sheet=sheet,
                cell=cell,
            )
            if rendered_format != expected_format:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="rendered_header_format_mismatch",
                        detail=f"rendered annual header format differs at {location}",
                        period=period,
                        item_id=item_id,
                        location=location,
                        expected=expected_format,
                        observed=rendered_format,
                    )
                )

            observed_cells = readbacks.get((sheet, cell), ())
            if len(observed_cells) != 1:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="readback_missing",
                        detail=(
                            f"exactly one workbook header readback is required for {location}"
                        ),
                        period=period,
                        item_id=item_id,
                        location=location,
                        expected=expected_value,
                        observed=len(observed_cells),
                    )
                )
                continue
            observed_cell = observed_cells[0]
            if observed_cell.value != expected_value:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="readback_value_mismatch",
                        detail=f"workbook header readback differs at {location}",
                        period=period,
                        item_id=item_id,
                        location=location,
                        expected=expected_value,
                        observed=observed_cell.value,
                    )
                )
            if observed_cell.number_format != expected_format:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="fiscal_axis_header_missing",
                        reason="readback_format_mismatch",
                        detail=f"workbook header number format differs at {location}",
                        period=period,
                        item_id=item_id,
                        location=location,
                        expected=expected_format,
                        observed=observed_cell.number_format,
                    )
                )
    return blockers


def _semantic_format_blockers(
    model: FinancialModel,
    *,
    render_plan: RenderPlan | None,
    readback: Sequence[PresentationCellReadback] | None,
    formatter: ModelFormatter,
    format_index: NumberFormatIndex,
) -> list[AnnualPresentationBlocker]:
    blockers: list[AnnualPresentationBlocker] = []
    readbacks = _readback_by_cell(readback)
    for sheet_name, item in _model_items(model):
        # A presentation caption is not a value row: it owes no unit or number format
        # (``presentation_captions`` owns which rows are captions).
        if item.item_type in {ItemType.header, ItemType.spacer} or item.id in WORKBOOK_PRESENTATION_CAPTION_ITEM_IDS:
            continue
        formula_bearing = bool(
            item.historical is not None
            or item.projected is not None
            or item.overrides
        )
        style_format = (
            item.style.number_format
            if item.style is not None
            else None
        )
        if (
            formula_bearing
            and item.unit is None
            and not str(item.format or "").strip()
            and not str(style_format or "").strip()
        ):
            blockers.append(
                AnnualPresentationBlocker(
                    code="unit_format_mismatch",
                    reason="semantic_format_metadata_missing",
                    detail=(
                        f"formula-bearing annual item {item.label!r} lacks unit "
                        "or explicit number-format metadata; heuristic formatting "
                        "is not authoritative"
                    ),
                    item_id=item.id,
                    location=f"model:{sheet_name}/{item.row}",
                )
            )
        try:
            expected_format = formatter.number_format_for(item)
        except ValueError as exc:
            blockers.append(
                AnnualPresentationBlocker(
                    code="unit_format_mismatch",
                    reason="explicit_format_invalid",
                    detail=str(exc),
                    item_id=item.id,
                )
            )
            continue
        if expected_format is None:
            continue
        cells = _target_cells(model, sheet_name=sheet_name, item=item)
        if render_plan is None:
            blockers.append(
                AnnualPresentationBlocker(
                    code="unit_format_mismatch",
                    reason="rendered_semantic_format_missing",
                    detail=f"render plan is required to verify {item.id!r}",
                    item_id=item.id,
                    expected=expected_format,
                )
            )
            continue

        for cell in cells:
            location = f"{sheet_name}!{cell}"
            rendered_format = _effective_number_format(
                format_index,
                sheet=sheet_name,
                cell=cell,
            )
            if rendered_format is None:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="unit_format_mismatch",
                        reason="rendered_semantic_format_missing",
                        detail=f"render plan omits the number format for {location}",
                        item_id=item.id,
                        location=location,
                        expected=expected_format,
                    )
                )
                break
            if rendered_format != expected_format:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="unit_format_mismatch",
                        reason="rendered_semantic_format_mismatch",
                        detail=f"render plan uses the wrong number format at {location}",
                        item_id=item.id,
                        location=location,
                        expected=expected_format,
                        observed=rendered_format,
                    )
                )
                break
            if item.unit not in _READBACK_REQUIRED_UNITS:
                continue
            observed_cells = readbacks.get((sheet_name, cell), ())
            if len(observed_cells) != 1:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="unit_format_mismatch",
                        reason="readback_missing",
                        detail=(
                            f"exactly one semantic-format readback is required for {location}"
                        ),
                        item_id=item.id,
                        location=location,
                        expected=expected_format,
                        observed=len(observed_cells),
                    )
                )
                break
            if observed_cells[0].number_format != expected_format:
                blockers.append(
                    AnnualPresentationBlocker(
                        code="unit_format_mismatch",
                        reason="readback_format_mismatch",
                        detail=f"workbook number-format readback differs at {location}",
                        item_id=item.id,
                        location=location,
                        expected=expected_format,
                        observed=observed_cells[0].number_format,
                    )
                )
                break
    return blockers


def _is_visible_template_residue(value: PresentationScalar) -> bool:
    if not isinstance(value, str):
        return False
    normalized = value.strip()
    if not normalized or normalized.startswith("="):
        return False
    return bool(
        _BRACKETED_PLACEHOLDER_RE.search(normalized)
        or _NUMBERED_PLACEHOLDER_RE.search(normalized)
        or _SLOT_NOUN_LABEL_RE.fullmatch(normalized)
        or _INSTRUCTION_TEXT_RE.search(normalized)
    )


def _template_text(model: FinancialModel) -> tuple[str, ...]:
    """Every model text through which a template slot reaches the workbook.

    The model is the template's instantiation, and ``excel_writer`` builds
    each workbook from an empty ``Workbook()``: a template placeholder can be
    visible only as a sheet name, a section or line label, or a string value
    of a line. Text a door writes into the render plan from elsewhere — the
    Sources & Assumptions index, its declination citations and assumption
    rationale — is not template text, whatever brackets it carries.
    """

    texts: list[object] = []
    for sheet_name, sheet in model.sheets.items():
        texts.append(sheet_name)
        texts.extend(section.label for section in sheet.sections)
    for _sheet_name, item in _model_items(model):
        texts.append(item.label)
        for spec in (
            item.historical,
            item.projected,
            *(item.overrides or {}).values(),
        ):
            if spec is not None and spec.type is FormulaType.constant:
                texts.append(spec.params.get("value"))
        if item.values is not None:
            texts.extend(cell.value for cell in item.values.values.values())
    return tuple(text for text in texts if isinstance(text, str))


def _residue_blockers(
    model: FinancialModel,
    *,
    render_plan: RenderPlan | None,
    readback: Sequence[PresentationCellReadback] | None,
) -> list[AnnualPresentationBlocker]:
    blockers: list[AnnualPresentationBlocker] = []
    # A rendered or read-back cell is residue when it shows a template slot
    # the model still leaves unfilled; the slots are the owner of the set.
    unfilled_slots = frozenset(
        text.strip()
        for text in _template_text(model)
        if _is_visible_template_residue(text)
    )

    def _shows_unfilled_slot(value: PresentationScalar) -> bool:
        return isinstance(value, str) and value.strip() in unfilled_slots

    for sheet_name, sheet in model.sheets.items():
        for location, text in (
            (f"model:{sheet_name}", sheet_name),
            *(
                (f"model:{sheet_name}/{section.id}", section.label)
                for section in sheet.sections
            ),
        ):
            if _is_visible_template_residue(text):
                blockers.append(
                    AnnualPresentationBlocker(
                        code="template_residue_visible",
                        reason="model_text_residue",
                        detail=f"model-visible text retains template residue: {text!r}",
                        location=location,
                        observed=text,
                    )
                )
    for _sheet_name, item in _model_items(model):
        if _is_visible_template_residue(item.label):
            blockers.append(
                AnnualPresentationBlocker(
                    code="template_residue_visible",
                    reason="model_text_residue",
                    detail=f"model label retains template residue: {item.label!r}",
                    location=f"model:{item.id}",
                    observed=item.label,
                )
            )
    for write in render_plan.writes if render_plan is not None else ():
        if _shows_unfilled_slot(write.value):
            blockers.append(
                AnnualPresentationBlocker(
                    code="template_residue_visible",
                    reason="rendered_text_residue",
                    detail=f"render plan retains visible template residue: {write.value!r}",
                    location=f"{write.sheet}!{write.cell.upper()}",
                    observed=write.value,
                )
            )
    for cell in readback or ():
        if _shows_unfilled_slot(cell.value):
            blockers.append(
                AnnualPresentationBlocker(
                    code="template_residue_visible",
                    reason="readback_text_residue",
                    detail=f"workbook readback retains template residue: {cell.value!r}",
                    location=f"{cell.sheet}!{cell.cell.upper()}",
                    observed=cell.value,
                )
            )
    return blockers


def evaluate_annual_presentation(
    model: FinancialModel,
    *,
    render_plan: RenderPlan | None,
    cell_readback: Sequence[PresentationCellReadback] | None,
    workbook_sha256: str,
    render_state_hash: str,
    formatter: ModelFormatter = SIA_FORMATTER,
) -> AnnualPresentationGuardResult:
    """Evaluate the sole BMA-012 model → render-plan → readback guard path."""

    blockers: list[AnnualPresentationBlocker] = []
    format_index = _number_format_index(render_plan) if render_plan is not None else {}
    axis: tuple[int, ...] = ()
    estimate_periods: frozenset[int] = frozenset()
    expected_headers: tuple[AnnualYearHeaderReadback, ...] = ()
    try:
        axis, estimate_periods = _annual_axis(model)
        expected_headers = expected_annual_year_headers(model)
    except ValueError as exc:
        blockers.append(
            AnnualPresentationBlocker(
                code="fiscal_axis_header_missing",
                reason="annual_fiscal_axis_invalid",
                detail=str(exc),
            )
        )

    fiscal_year_end_valid = False
    if axis:
        try:
            for period in axis:
                _fiscal_year_end_date(model.time_structure.fiscal_year_end, period)
            fiscal_year_end_valid = True
        except ValueError as exc:
            blockers.append(
                AnnualPresentationBlocker(
                    code="fiscal_axis_header_missing",
                    reason="fiscal_year_end_invalid",
                    detail=str(exc),
                )
            )

    if axis and fiscal_year_end_valid:
        blockers.extend(
            _header_model_blockers(
                model,
                axis=axis,
                estimate_periods=estimate_periods,
                formatter=formatter,
            )
        )
        blockers.extend(
            _header_plan_and_readback_blockers(
                model,
                axis=axis,
                render_plan=render_plan,
                readback=cell_readback,
                formatter=formatter,
                format_index=format_index,
            )
        )
    blockers.extend(
        _semantic_format_blockers(
            model,
            render_plan=render_plan,
            readback=cell_readback,
            formatter=formatter,
            format_index=format_index,
        )
    )
    blockers.extend(
        _residue_blockers(
            model,
            render_plan=render_plan,
            readback=cell_readback,
        )
    )
    blockers.sort(
        key=lambda blocker: (
            blocker.code,
            blocker.reason,
            blocker.period if blocker.period is not None else -1,
            blocker.item_id or "",
            blocker.location or "",
            blocker.detail,
        )
    )
    return AnnualPresentationGuardResult(
        status="blocked" if blockers else "ready",
        workbook_sha256=workbook_sha256,
        render_state_hash=render_state_hash,
        expected_axis=axis,
        expected_headers=expected_headers,
        blockers=tuple(blockers),
    )


def evaluate_persisted_annual_presentation(
    model: FinancialModel,
    *,
    render_plan: RenderPlan,
    workbook_path: str | Path,
    render_state: object,
    formatter: ModelFormatter = SIA_FORMATTER,
) -> AnnualPresentationGuardResult:
    """Freshly evaluate exact final workbook bytes and their render state.

    This is the only producer-side path that may create ready presentation
    evidence. It performs a new persisted-XLSX readback and independently
    derives both content bindings; it never upgrades or rebinds a prior result.
    """

    from .render_state import render_state_hash

    resolved = Path(workbook_path).expanduser()
    try:
        workbook_sha256 = hashlib.sha256(resolved.read_bytes()).hexdigest()
    except OSError as exc:
        raise AnnualPresentationBindingError(
            f"annual-presentation workbook is unreadable: {resolved}"
        ) from exc
    actual_render_state_hash = render_state_hash(render_state)
    if actual_render_state_hash == "absent":
        raise AnnualPresentationBindingError(
            "annual-presentation evaluation requires a complete render_state"
        )
    return evaluate_annual_presentation(
        model,
        render_plan=render_plan,
        cell_readback=read_annual_presentation_cells(resolved),
        workbook_sha256=workbook_sha256,
        render_state_hash=actual_render_state_hash,
        formatter=formatter,
    )


def evaluate_ready_annual_presentation_model_ref(
    model_ref: object,
    *,
    model: FinancialModel,
    render_plan: RenderPlan,
    workbook_path: str | Path,
    formatter: ModelFormatter = SIA_FORMATTER,
) -> tuple[dict[str, Any], AnnualPresentationGuardResult]:
    """Freshly evaluate and stamp one final producer candidate.

    Unlike a rebind helper, this always re-runs the full model → render-plan →
    persisted-XLSX guard over the final bytes. The stamped result carries its
    status and blockers as evaluated.
    """

    payload = _model_ref_payload(model_ref)
    result = evaluate_persisted_annual_presentation(
        model,
        render_plan=render_plan,
        workbook_path=workbook_path,
        render_state=payload.get("render_state"),
        formatter=formatter,
    )
    payload["annual_presentation"] = result.model_dump(mode="json")
    require_bound_ready_annual_presentation(payload, workbook_path=workbook_path)
    return payload, result


__all__ = [
    "AnnualEstimateStatus",
    "AnnualPresentationBlocker",
    "AnnualPresentationBlockerCode",
    "AnnualPresentationBindingError",
    "AnnualPresentationGuardResult",
    "AnnualPresentationReason",
    "AnnualPresentationStatus",
    "AnnualYearHeaderReadback",
    "PresentationCellReadback",
    "evaluate_annual_presentation",
    "evaluate_persisted_annual_presentation",
    "evaluate_ready_annual_presentation_model_ref",
    "expected_annual_year_headers",
    "read_annual_presentation_cells",
    "require_bound_ready_annual_presentation",
]
