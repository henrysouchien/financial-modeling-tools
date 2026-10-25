"""The live semantic sheet: the persisted ``FinancialModel`` as an authored ``ModelSheet``.

A ``ModelSheet`` is the ``model`` of a committed ModelVersion's ``ModelSemanticDocument``, taken as the
authority for structure and formulas. Edits go through ``schema.modify_core``'s dispatcher;
formulas are authored as readable infix text over semantic item ids and parsed here into the
nested-expression grammar ``schema/dependency_graph_eval.py`` evaluates — anything outside it
(``IF``, comparisons, lookups, other functions, text) is stored as a ``raw`` formula over the
same ids, which our graph reports as ``raw_excel_formula_semantics_unavailable`` and the native
engine evaluates. ``recalc`` is ours (``DependencyGraph.compute_typed``, no stored value standing
in for a formula; required cited absences and uncited blanks propagate as missing;
only catalog-optional cited SUM terms may be omitted, retaining their citations);
``render`` is the product renderer and XLSX writer; ``native_recalc`` is
headless LibreOffice; ``compare`` joins the two through the render plan's semantic cells.
``inspect`` reads lines back: formulas as infix text (``formula_text``, the inverse of
``parse_formula``), operands, dependents, and how each period's cell is produced.
``upstream_cells`` / ``downstream_cells`` walk the graph cell by cell (period offsets resolved).
``apply_ops`` is the agent's edit batch: each op's ``rationale`` and ``source`` ride on its edit
and the columns of the cells it writes. ``run_journeys`` is the scripted J1/J3/J5/J6/J8.

Formula text: ``item.id`` references the same period, ``item.id[-1]`` the prior one; operators
``+ - * / ^`` with spreadsheet precedence; functions ``SUM``, ``AVG``/``AVERAGE``, ``MAX``,
``MIN``; numeric literals. A raw formula keeps the ids; they resolve to cells per
period at render, so the sheet never holds a coordinate.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Callable, Collection, Iterable, Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import UTC, datetime
import copy
import hashlib
import json
import math
import os
from pathlib import Path
import pickle
import re
import shutil
import subprocess
import tempfile
import time
from typing import TYPE_CHECKING, Any, Literal, NoReturn, get_args

from industry_slice_core.industry_types import Unit
from schema.build_company_grounded_schedules import (
  apply_not_separately_presented_periods,
  not_separately_presented_formula,
)
from schema.build_value_writers import CellSourceStamp, stamp_cell_source
from schema.dependency_graph import (
  DependencyGraph,
  DependencyGraphKnownNumberValue,
  DependencyGraphNumberValue,
  formula_owned_recompute_ids,
)
from schema.driver_formula import driver_formula_params
from schema.excel_writer import render_xlsx_bytes
from schema.finding_vocabulary import FINDING_MESSAGE_MAX_LENGTH
from schema.models import (
  canonical_strategy,
  FinancialModel,
  FormulaSpec,
  FormulaType,
  ItemType,
  LineItem,
  LineItemRef,
  ModelScale,
  PeriodAnchor,
  Section,
  SheetLayout,
  ValueCell,
  ValueProvenance,
  ValueSeries,
)
from schema.coverage_profile import AbsenceCitation
from schema.model_number_value import (
  AbsenceDisposition,
  AbsentCause,
  CellUnavailableReason,
  EvidencedAbsentNumberValue,
  is_absence_disposition,
  MissingReasonCode,
  ZeroDisposition,
)
from schema.model_queries import ANALYST_NOTE_ARTIFACT
from schema.modify_core import OperationResult, _dispatch_op, validate_operation
from schema.published_model_readback import writer_numeric_equal
from schema.renderer import (
  AbsoluteColumnMapper,
  _fixed_cell_anchor_period,
  _spec_for_period as _render_spec_for_period,
  render_model,
)
from schema.refs import line_item_ref_from_obj
from schema.renderer_formula import ExcelFormulaCompiler
from schema.renderer_types import AnnualRenderContext
from schema.stable_control_identity import CellKey, FiscalPeriod
from schema.statement_bindings_v2 import NOT_DISCLOSED_IN_VINTAGE

if TYPE_CHECKING:
  from schema.build_model_product_contracts import ModelSpec
  from schema.coverage_profile import AnnualSourceRequirement
  from schema.model_semantic_document import ModelSemanticDocument, SemanticValueSnapshot


Value = DependencyGraphNumberValue
Values = dict[tuple[str, int], Value]
ApplyTo = Literal["projected", "historical", "both"]

#: Recorded gaps and inputs: stamps of a coordinate the row owns (``item_owns_cell``), never a claim on one.
STAMPED_CELL_KINDS: frozenset[str] = frozenset({"unavailable", "not_separately_presented", "input"})

#: The workbook timestamp every bench render carries; renders of one document are byte-stable.
_RENDER_TIME = datetime(1980, 1, 1, tzinfo=UTC)
_SOFFICE_CANDIDATES = (
  "/opt/homebrew/bin/soffice",
  "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


# --------------------------------------------------------------------------
# the document
# --------------------------------------------------------------------------


@dataclass
class ModelSheet:
  """One authored model: the ``FinancialModel`` plus what it was loaded beside.

  ``semantic`` is the document's semantic column per cell, ``(item_id, period)`` → the record
  the writer that last wrote the cell was given (``source``: a ``FilingLineRef`` key, a
  judgment row, a derivation; ``origin``; for a cell made unavailable, ``state:
  "unavailable"`` and its ``reason``). Each writer also carries the record on its
  ``doc.edits`` row, so the edit history is the column's lineage. ``values_snapshot`` is
  the last recalc saved with the document (``value_row`` per cell), restored by ``reopen``.
  """

  model: FinancialModel
  model_spec: ModelSpec | None
  semantic_values: tuple[SemanticValueSnapshot, ...]
  source: str
  edits: list[dict[str, Any]] = field(default_factory=list)
  semantic: dict[tuple[str, int], dict[str, Any]] = field(default_factory=dict)
  values_snapshot: dict[tuple[str, int], dict[str, Any]] = field(default_factory=dict)
  #: How many ``transaction`` blocks are open on this document (0 outside one).
  _open: int = field(default=0, init=False, repr=False, compare=False)

  def __post_init__(self) -> None:
    self.semantic = {key: stamp_cell_source(row) for key, row in self.semantic.items()}
    for item in self.model._index.values():
      if item.values is None:
        continue
      for period, cell_obj in item.values.values.items():
        lineage = getattr(cell_obj, "lineage_reference", None)
        key = (item.id, int(period))
        if isinstance(lineage, Mapping) and key not in self.semantic:
          self.semantic[key] = stamp_cell_source(lineage)


def copy_model(model: FinancialModel) -> FinancialModel:
  """An independent deep copy of ``model`` (a pickle round trip: the object graph's sharing is
  kept, as ``model_copy(deep=True)`` keeps it, at about a quarter of its cost on a full model).
  The copy's ``_index`` is the original's; callers ``build_index`` it as before."""
  return pickle.loads(pickle.dumps(model, protocol=pickle.HIGHEST_PROTOCOL))


def from_semantic_document(semantic: ModelSemanticDocument, model_spec: ModelSpec | None) -> ModelSheet:
  """A committed ModelVersion's semantic-document model as a detached ``ModelSheet``. The adapter reads
  the version (the bench's ``context.load_document``) and passes its two artifacts."""
  model = semantic.model.model_copy(deep=True)
  model.build_index()
  return ModelSheet(
    model=model,
    model_spec=model_spec,
    semantic_values=tuple(semantic.semantic_values),
    source=f"{semantic.ticker} {semantic.model_id}@{semantic.model_version}",
  )


def fork(doc: ModelSheet) -> ModelSheet:
  """An independent copy: edits to it never reach ``doc`` (try / compare / keep)."""
  model = copy_model(doc.model)
  model.build_index()
  return ModelSheet(
    model=model,
    model_spec=doc.model_spec,
    semantic_values=doc.semantic_values,
    source=doc.source,
    edits=[dict(edit) for edit in doc.edits],
    semantic={key: stamp_cell_source(row) for key, row in doc.semantic.items()},
    values_snapshot={key: dict(row) for key, row in doc.values_snapshot.items()},
  )


def _package_path(path: Path, part: str) -> Path:
  return path.with_name(f"{path.stem}.{part}.json")


def _edits_path(path: Path) -> Path:
  return _package_path(path, "edits")


def _cell_rows(cells: Mapping[tuple[str, int], Mapping[str, Any]]) -> list[dict[str, Any]]:
  return [{"item_id": item, "period": period, **dict(row)} for (item, period), row in sorted(cells.items())]


def _cell_map(rows: Iterable[Mapping[str, Any]]) -> dict[tuple[str, int], Any]:
  return {
    (str(row["item_id"]), int(row["period"])): stamp_cell_source({
      key: value for key, value in row.items() if key not in ("item_id", "period")
    })
    for row in rows
  }


def save(doc: ModelSheet, path: Path, *, values: Values | None = None) -> Path:
  """The document package: the ``FinancialModel`` JSON at ``path`` (the file ``dump values``
  reads) and beside it ``<stem>.edits.json`` (the edit history, written when non-empty),
  ``<stem>.semantic.json`` (every cell's semantic column) and ``<stem>.values.json`` (the
  value snapshot: ``value_row`` per cell and its ``values_digest``). ``values`` is the recalc
  to snapshot; default ``recalc(doc)``. ``reopen`` restores all of it.
  """
  target = Path(path).expanduser()
  target.parent.mkdir(parents=True, exist_ok=True)
  target.write_text(doc.model.model_dump_json(indent=1))
  edits = _edits_path(target)
  if doc.edits:
    edits.write_text(json.dumps(doc.edits, indent=1, default=str))
  elif edits.exists():
    edits.unlink()
  computed = recalc(doc) if values is None else values
  doc.values_snapshot = {key: value_row(value) for key, value in computed.items()}
  _package_path(target, "semantic").write_text(json.dumps(_cell_rows(doc.semantic), indent=1, default=str))
  _package_path(target, "values").write_text(json.dumps(
    {"values_sha256": values_digest(computed), "cells": _cell_rows(doc.values_snapshot)}, indent=1, default=str,
  ))
  return target


def reopen(path: Path) -> ModelSheet:
  """A saved document package; a ``ModelSemanticDocument`` JSON is read through its ``model``.

  The edits, semantic columns and value snapshot are restored when their files exist, so a
  recalc after reopen equals the saved snapshot (``values_digest``)."""
  source = Path(path).expanduser()
  payload = json.loads(source.read_text())
  if isinstance(payload, Mapping) and "model" in payload and "company" not in payload:
    payload = payload["model"]
  model = FinancialModel.model_validate(payload)
  model.build_index()
  edits_path = _edits_path(source)
  edits = json.loads(edits_path.read_text()) if edits_path.is_file() else []
  semantic_path = _package_path(source, "semantic")
  semantic = _cell_map(json.loads(semantic_path.read_text())) if semantic_path.is_file() else {}
  values_path = _package_path(source, "values")
  snapshot = _cell_map(json.loads(values_path.read_text())["cells"]) if values_path.is_file() else {}
  return ModelSheet(
    model=model, model_spec=None, semantic_values=(), source=str(source), edits=edits,
    semantic=semantic, values_snapshot=snapshot,
  )


def saved_values_digest(path: Path) -> str | None:
  """The ``values_sha256`` a saved package recorded (``None`` without a value snapshot)."""
  values_path = _package_path(Path(path).expanduser(), "values")
  return json.loads(values_path.read_text())["values_sha256"] if values_path.is_file() else None


# --------------------------------------------------------------------------
# edits
# --------------------------------------------------------------------------


class EditRefused(Exception):
  """``modify_core`` returned an error result; the document is unchanged."""

  def __init__(self, result: OperationResult, *, op_index: int | None = None) -> None:
    self.result = result
    self.op_index = op_index
    self.op_errors: list[dict[str, Any]] = []
    self.valid_op_indices: list[int] = []
    super().__init__(f"{result.op_type.value} {result.item_id}: {result.reason} {result.details or ''}".strip())


def _apply(
  doc: ModelSheet,
  payload: Mapping[str, Any],
  *,
  then: Callable[[FinancialModel], None] | None = None,
  declared: Mapping[str, Any] | None = None,
  source: Mapping[str, Any] | None = None,
  agent: Mapping[str, Any] | None = None,
) -> OperationResult:
  """One ``modify_core`` operation on the document's working model, swapped in only when it
  succeeds (``transaction``: outside one, that is a copy of its own); ``declared``
  rides beside the op in ``doc.edits`` (what the agent declared that the payload cannot carry),
  and so does ``source`` (the semantic column the writer records for the cells it writes) or,
  for an agent op (``apply_ops``), ``agent`` (its ``source``, ``rationale`` and ``edit_index``)."""
  operation = validate_operation(dict(payload))
  with transaction(doc) as trial:
    result = _dispatch_op(trial, operation, file_path=doc.source)
    record: dict[str, Any] = {"op": dict(payload)}
    if declared is not None:
      record["declared"] = dict(declared)
    if source is not None:
      record["source"] = dict(source)
    if agent is not None:
      record["agent"] = dict(agent)
    doc.edits.append({**record, "result": result.model_dump(mode="json")})
    if result.status == "error":
      raise EditRefused(result)
    if then is not None:
      then(trial)
  return result


@contextmanager
def transaction(doc: ModelSheet) -> Iterator[FinancialModel]:
  """One ``copy_model`` + ``build_index`` on enter; every write inside (``_apply`` and the other
  writers that copy the model) works on that one copy, which is ``doc.model`` while the block
  runs; on exit without error the copy is re-indexed and stays, on error ``doc.model`` and
  ``doc.semantic`` are as they were on enter (``doc.edits`` keeps what was appended: ``_apply``
  records its edit before it can refuse). All-or-nothing: the caller of a block that raises gets
  the document it passed in.

  A transaction opened inside another one joins it: it copies nothing, an error leaves the
  rollback to the outermost block, and a block that catches an error raised inside a nested one
  and carries on writes on a model the failed write may have left half-changed (no caller does)."""
  if doc._open:
    doc._open += 1
    try:
      yield doc.model
    finally:
      doc._open -= 1
    doc.model.build_index()
    return
  original, semantic = doc.model, dict(doc.semantic)
  trial = copy_model(original)
  trial.build_index()
  doc.model = trial
  doc._open = 1
  try:
    yield trial
  except BaseException:
    doc.model, doc.semantic = original, semantic
    raise
  finally:
    doc._open = 0
  doc.model.build_index()


def record_edit(doc: ModelSheet, op: Mapping[str, Any], *, source: Mapping[str, Any], result: Mapping[str, Any]) -> None:
  """An edit applied by a typed library rather than ``modify_core`` (a structural lowering),
  recorded in ``doc.edits`` with the artifact it was authored from."""
  doc.edits.append({"op": dict(op), "source": dict(source), "result": dict(result)})


def set_company(doc: ModelSheet, ticker: str, name: str | None = None, cik: str | None = None) -> None:
  """Stamp the company a document built from the template models: ``model.company``'s ticker and
  name (the template's own are ``TPL`` / its template name; ``name`` defaults to the ticker),
  recorded as one ``set_company`` edit. ``CompanyInfo`` has no CIK field, so ``cik`` is kept on
  that edit (``company(doc)`` reads all three back)."""
  previous = doc.model.company.model_dump(mode="json")
  doc.model.company = doc.model.company.model_copy(update={"ticker": str(ticker), "name": str(name or ticker)})
  record_edit(
    doc, {"type": "set_company", "ticker": str(ticker), "name": str(name or ticker), "cik": cik},
    source={"origin": "research_file"}, result={"status": "ok", "previous": previous},
  )


def company(doc: ModelSheet) -> dict[str, Any]:
  """The document's company: ``ticker`` and ``name`` (``model.company``) and the ``cik`` the
  latest ``set_company`` edit recorded (``None`` without one)."""
  stamped = next((edit["op"] for edit in reversed(doc.edits) if (edit.get("op") or {}).get("type") == "set_company"), {})
  return {"ticker": doc.model.company.ticker, "name": doc.model.company.name, "cik": stamped.get("cik")}


class OpUsage(ValueError):
  """An op the bench cannot read: a missing or unknown key, an ambiguous unit, a bad ``$ref``.

  The CLI reports it as a usage error (exit 2) whose message names the expected keys or forms.
  """

  def __init__(
    self,
    message: str,
    *,
    op_index: int | None = None,
    reason: str = "op_usage",
    next_call: Mapping[str, Any] | None = None,
  ):
    super().__init__(message)
    self.op_index = op_index
    self.reason = reason
    self.next_call = None if next_call is None else dict(next_call)
    self.op_errors: list[dict[str, Any]] = []
    self.valid_op_indices: list[int] = []

_SCALES: dict[str, ModelScale] = {
  "millions": "millions", "million": "millions", "mm": "millions", "m": "millions", "units": "units",
}
_CURRENCY_WORDS = frozenset({"usd", "$", "dollars", "dollar"})
_UNIT_ALIASES = {"%": Unit.percentage, "percent": Unit.percentage, "pct": Unit.percentage}


def unit_grammar() -> str:
  """The accepted unit forms, as every unit refusal states them."""
  units = ", ".join(member.value for member in Unit if member is not Unit.dollars)
  return (
    "'USD millions' or 'USD units' (currency USD|$|dollars|dollar, then its scale — a currency "
    f"without a scale is ambiguous); or one of {units} (percentage also %, percent, pct), "
    "optionally followed by a scale millions|million|mm|m|units (default units)"
  )


def parse_unit(text: str) -> tuple[Unit, ModelScale]:
  """The agent-declared unit of a new line → ``(LineItem.unit, LineItem.model_scale)``.

  Nothing is inferred: an unknown word, or a currency with no scale, is an ``OpUsage``.
  ``LineItem`` has no currency field, so only USD is representable, and ``Unit`` has no
  per-employee or boolean member (type gaps the bench reports, not hides).
  """
  words = text.strip().lower().replace("_", " ").split()
  if not words:
    raise OpUsage(f"unit: declare a unit; accepted: {unit_grammar()}")
  by_value = {member.value: member for member in Unit}
  head, rest = words[0], words[1:]
  if head in _CURRENCY_WORDS:
    if not rest:
      raise OpUsage(f"unit: {text!r} names a currency without a scale; accepted: {unit_grammar()}")
    unit = Unit.dollars
  elif head in _UNIT_ALIASES:
    unit = _UNIT_ALIASES[head]
  elif len(words) >= 2 and "_".join(words[:2]) in by_value:
    unit, rest = by_value["_".join(words[:2])], words[2:]
  elif head in by_value:
    unit = by_value[head]
  else:
    raise OpUsage(f"unit: {text!r} is not an accepted unit; accepted: {unit_grammar()}")
  if not rest:
    return unit, "units"
  if len(rest) != 1 or rest[0] not in _SCALES:
    raise OpUsage(f"unit: scale {' '.join(rest)!r} in {text!r} is not accepted; accepted: {unit_grammar()}")
  return unit, _SCALES[rest[0]]


def _slug(label: str) -> str:
  slug = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
  if not slug:
    raise ValueError(f"label {label!r} has no identifier characters")
  return slug


def _sections(model: FinancialModel) -> Iterator[tuple[str, Section]]:
  for sheet_name, sheet in model.sheets.items():
    for section in sheet.sections:
      yield sheet_name, section


def _locate_section(model: FinancialModel, section_id: str, after_item_id: str | None) -> tuple[str, Section]:
  if after_item_id is not None:
    for sheet_name, section in _sections(model):
      if any(item.id == after_item_id for item in section.line_items):
        if section.id != section_id:
          raise ValueError(f"{after_item_id!r} is in section {section.id!r}, not {section_id!r}")
        return sheet_name, section
    raise ValueError(f"after_item_id {after_item_id!r} is not a line of the document")
  matches = [(sheet_name, section) for sheet_name, section in _sections(model) if section.id == section_id]
  if len(matches) != 1:
    known = sorted({section.id for _sheet, section in _sections(model)})
    raise ValueError(f"section {section_id!r} names {len(matches)} sections (pass after_item_id); sections: {known}")
  return matches[0]


def add_line(
  doc: ModelSheet,
  *,
  section: str,
  label: str,
  unit: str,
  after_item_id: str | None = None,
  formula: str | None = None,
  inputs: Mapping[int, float] | None = None,
  agent: Mapping[str, Any] | None = None,
) -> str:
  """A new line in ``section`` (a section id as ``doc_load`` lists it); returns the minted id.

  The id is ``doc.<section>.<slug(label)>`` — ``slug`` lowercases and joins runs of anything
  but ``a-z0-9`` with ``_`` — with ``_2``, ``_3``… appended on collision. The agent declares
  ``label`` and ``unit``; ``formula`` (applied to every period) makes the line derived,
  ``inputs`` are per-period input values. The id is minted before ``formula`` is parsed, so
  the formula may read the line itself (``<id>[-1]``, a roll-forward). Row and type are the
  door's.
  """
  unit_value, scale = parse_unit(unit)
  sheet_name, located = _locate_section(doc.model, section, after_item_id)
  base = f"doc.{located.id}.{_slug(label)}"
  item_id, suffix = base, 2
  while item_id in doc.model._index:
    item_id, suffix = f"{base}_{suffix}", suffix + 1
  payload: dict[str, Any] = {
    "type": "add_item",
    "item_id": item_id,
    "sheet_name": sheet_name,
    "section_id": located.id,
    "label": label,
    "after_item_id": after_item_id,
    "item_type": "derived" if formula is not None else "input",
    "unit": unit_value.value,
  }
  if inputs:
    payload["values"] = {int(period): float(value) for period, value in inputs.items()}
  if formula is not None:
    payload["formula"] = parse_formula(formula, doc.model, self_id=item_id).model_dump(mode="json")
    payload["apply_to"] = "both"

  def _scale(model: FinancialModel) -> None:
    # AddItemOperation carries no scale; the declared one is set on the new row (type gap).
    line = model.get_item(item_id)
    line.model_scale = scale
    if formula is None and inputs:
      line.formula_periods = sorted(int(period) for period in inputs)

  _apply(doc, payload, then=_scale, declared={"unit": unit, "model_scale": scale}, agent=agent)
  return item_id


def set_formula(
  doc: ModelSheet,
  item_id: str,
  formula_expr: str,
  *,
  apply_to: ApplyTo = "both",
  source: Mapping[str, Any] | None = None,
  agent: Mapping[str, Any] | None = None,
) -> FormulaSpec:
  """Author ``item_id``'s formula and stamp any supplied source on the cells it produces."""
  spec, _, _, _ = _apply_set_formula_with_columns(
    doc, item_id, formula_expr, apply_to=apply_to, source=source, agent=agent, graphs=_BatchGraphs(doc),
  )
  return spec


def _agent_inputs_in_scope(doc: ModelSheet, item_id: str, apply_to: ApplyTo) -> list[dict[str, Any]]:
  """The ``_cell_state`` of every cell of ``item_id`` on the ``apply_to`` side(s) that holds an agent
  input (``cell: input`` column with an ``edit_index``) or a projected override whose per-period
  spec would outrank the formula written there (``_spec_for_period`` selects the override first);
  a fixed-cell line's one cell is its anchor. Historical filing constants, historical formula
  overrides, and ``not_separately_presented`` markers stay."""
  item = doc.model.get_item(item_id)
  time = doc.model.time_structure
  projection_periods = {
    int(period) for period in (time.projection_periods or time.projection_years or ())
  }
  periods: set[int] = set()
  if apply_to in ("projected", "both"):
    periods.update(projection_periods)
  if apply_to in ("historical", "both"):
    periods.update(int(period) for period in time.historical_periods or time.historical_years or ())
  if item.column is not None:
    anchor = _fixed_cell_anchor_period(doc.model, item)
    periods &= set() if anchor is None else {int(anchor)}
  rows: list[dict[str, Any]] = []
  for period in sorted(periods):
    override = (item.overrides or {}).get(period)
    column = doc.semantic.get((item_id, period)) or {}
    if override is None or override.subtype == "not_separately_presented":
      continue
    if period in projection_periods or (
      override.type is FormulaType.constant
      and column.get("cell") == "input"
      and isinstance(column.get("edit_index"), int)
    ):
      rows.append(_cell_state(doc, (item_id, period)))
  return rows


def _apply_set_formula_with_columns(
  doc: ModelSheet,
  item_id: str,
  formula_expr: str,
  *,
  apply_to: ApplyTo,
  source: Mapping[str, Any] | None = None,
  agent: Mapping[str, Any] | None = None,
  graphs: _BatchGraphs,
) -> tuple[FormulaSpec, set[CellKey], list[list[Any]], list[dict[str, Any]]]:
  """``set_formula`` and the columns it stamps: the spec, the cells written, the columns replaced
  and, for an agent op, the agent inputs it cleared (``_agent_inputs_in_scope``: their override and
  ``ValueCell`` leave with the write so the formula is what the cell evaluates; each cell's prior
  column, override and value ride on the edit as ``replaced``, the shape ``set_inputs`` records)."""
  spec = parse_formula(formula_expr, doc.model)
  # the specs the agent's formula displaces, read before the write; a line the model does not hold is
  # ``_apply``'s refusal (``_dispatch_op``), recorded as a refused edit, never an index error here
  prior = doc.model._index.get(item_id)
  cleared = _agent_inputs_in_scope(doc, item_id, apply_to) if prior is not None else []
  replaced_formula = None if prior is None else {
    "historical": None if prior.historical is None else prior.historical.model_dump(mode="json"),
    "projected": None if prior.projected is None else prior.projected.model_dump(mode="json"),
    "formula_periods": None if prior.formula_periods is None else [int(period) for period in prior.formula_periods],
  }

  def _clear_inputs(model: FinancialModel) -> None:
    line = model.get_item(item_id)
    for row in cleared:
      (line.overrides or {}).pop(row["period"], None)
      if line.values is not None:
        line.values.values.pop(row["period"], None)

  _apply(
    doc,
    {"type": "set_formula", "item_id": item_id, "formula": spec.model_dump(mode="json"), "apply_to": apply_to},
    then=_clear_inputs,
    source=source,
    agent=agent,
  )
  if cleared and agent is not None:
    doc.edits[-1]["replaced"] = cleared
  if agent is not None and replaced_formula is not None:
    # the line's specs before the agent's formula: what a construction re-run reads in its place
    # (``_formula_baseline_before_agent``), the way ``replaced`` records a displaced cell
    doc.edits[-1]["replaced_formula"] = replaced_formula
  graphs.written(item_id)
  meta = agent if agent is not None else source
  if meta is None:
    return spec, set(), [], cleared
  written, replaced = _formula_columns(doc, item_id, meta, apply_to, is_agent=agent is not None, graphs=graphs)
  return spec, written, replaced, cleared


#: A written input's ``ValueCell`` provenance by its semantic column's ``origin`` (what the
#: product's population records beside the override: ``build_population._set_imported_value``).
_VALUE_PROVENANCE = {
  "observed": ValueProvenance.imported_edgar,
  "derived": ValueProvenance.derived,
  "manual_input": ValueProvenance.input,
  "cited_input": ValueProvenance.input,
}

#: The ``source`` object kinds that cite evidence (the correspondence door's ``_SourceKey`` kinds
#: and a corpus capture): an input written with one carries ``origin: cited_input``.
CITED_SOURCE_KINDS = frozenset({"filing_line", "kpi", "line_identity", "capture"})


def _input_origin(source: Any) -> str:
  """``cited_input`` for a parsed ``CellSource`` of a cited kind (``filing_line``, ``kpi``,
  ``line_identity``, ``capture``), else ``manual_input`` (free text, ``{kind: agent}``,
  partial/legacy opaque, absent)."""
  parsed = parse_stored_source(source)
  return "cited_input" if parsed is not None and parsed.origin == "cited_input" else "manual_input"

#: The typed ``source`` object kinds with which an agent gives an input as its own judgment
#: (``{kind: analyst_note | assumption, note: <text>}``); ``ANALYST_NOTE_ARTIFACT`` names the
#: judgment artifact the minted row belongs to (``analyst_note_judgment``).
ANALYST_NOTE_KINDS = frozenset({"analyst_note", "assumption"})
_ANALYST_NOTE_PROSE_HANDLE_RE = re.compile(
  r"\[S_?([1-9]\d*)\]|(?<![A-Za-z0-9_])(src_[1-9]\d*)(?![A-Za-z0-9_])"
)
_SINGLE_SOURCE_HANDLE_RE = re.compile(r"^(?:\[S_?([1-9]\d*)\]|S_?([1-9]\d*)|(src_[1-9]\d*))$")


def _canonical_source_handle(raw: str) -> str | None:
  match = _SINGLE_SOURCE_HANDLE_RE.fullmatch(raw.strip())
  if match is None:
    return None
  num = match.group(1) or match.group(2)
  return f"src_{int(num)}" if num is not None else match.group(3)

def _parse_analyst_note_source_refs(raw_refs: Any, *, where: str) -> list[str]:
  if not isinstance(raw_refs, Sequence) or isinstance(raw_refs, (str, bytes)):
    raise OpUsage(
      f"{where}: source.source_refs must be a list of '[Sn]' or 'src_<n>' handles, not {raw_refs!r}"
    )
  normalized: list[str] = []
  for item in raw_refs:
    if not isinstance(item, str):
      raise OpUsage(
        f"{where}: source.source_refs entries must be '[Sn]' or 'src_<n>' strings, not {item!r}"
      )
    canon = _canonical_source_handle(item)
    if canon is None:
      raise OpUsage(
        f"{where}: source.source_refs entries must match '[Sn]' or 'src_<n>', not {item!r}"
      )
    normalized.append(canon)
  return list(dict.fromkeys(normalized))


def extract_analyst_note_handles(
  note: str | None, source_refs: Sequence[str] | None = None,
) -> tuple[str, ...]:
  """The canonical ``src_<n>`` ``SourceId`` handles an analyst note cites (from ``[Sn]``, ``[S_n]``,
  or ``src_<n>`` in ``note``, or ``[Sn]`` / ``Sn`` / ``src_<n>`` in ``source_refs``)."""
  handles: list[str] = []
  if isinstance(note, str):
    for match in _ANALYST_NOTE_PROSE_HANDLE_RE.finditer(note):
      handles.append(f"src_{int(match.group(1))}" if match.group(1) is not None else match.group(2))
  if isinstance(source_refs, Sequence) and not isinstance(source_refs, (str, bytes)):
    for ref in source_refs:
      if isinstance(ref, str):
        canon = _canonical_source_handle(ref)
        if canon is not None:
          handles.append(canon)
  return tuple(dict.fromkeys(handles))


def is_historical_statement_cell(model: FinancialModel, item_id: str, period: int) -> bool:
  """Whether ``(item_id, period)`` is a historical ``tpl.fm.*`` financial-statement cell."""
  if not str(item_id).startswith("tpl.fm."):
    return False
  time = model.time_structure
  historical = {int(p) for p in (time.historical_periods or time.historical_years or ())}
  return int(period) in historical
def analyst_note_judgment(
  doc: ModelSheet, source: Any, item_id: str, period: int, edit_index: Any,
) -> dict[str, Any] | None:
  """The ``judgment_row`` source a typed ``{kind: analyst_note | assumption, note: <non-empty text>}``
  object becomes on the cell ``(item_id, period)``, else ``None`` (free text, ``{kind: agent}``, a
  note-less object, a historical ``tpl.fm.*`` statement cell and every other kind stay as they are).
  It is the shape a judgment artifact's row gives a cell, so publish records the cell's assumption
  (``document_judgment``) and the cell publishes ``complete`` by the analyst's own judgment:
  ``{kind: judgment_row, artifact: analyst_note, ref: <assumption_id>, row: {assumption_id,
  rationale: <note>, kind: <given kind>}}``. The ``assumption_id`` is minted
  ``analyst_note:<item_id>:<period>:<edit_index>`` — one per written cell, using the anchor period
  for a fixed-cell line, and the same on every replay of the log (``edit_index`` is the edit's
  position in it)."""
  if not isinstance(source, Mapping) or source.get("kind") not in ANALYST_NOTE_KINDS:
    return None
  note = source.get("note")
  if not isinstance(note, str) or not note.strip():
    return None
  cell_item_id, cell_period = _input_cell(doc, item_id, period)
  if is_historical_statement_cell(doc.model, cell_item_id, cell_period):
    return None
  assumption_id = f"{ANALYST_NOTE_ARTIFACT}:{item_id}:{cell_period}:{edit_index}"
  row: dict[str, Any] = {
    "assumption_id": assumption_id,
    "rationale": note.strip(),
    "kind": source["kind"],
  }
  raw_refs = source.get("source_refs")
  if isinstance(raw_refs, Sequence) and not isinstance(raw_refs, (str, bytes)):
    explicit_refs = [
      canon
      for ref in raw_refs
      if isinstance(ref, str) and (canon := _canonical_source_handle(ref)) is not None
    ]
    if explicit_refs:
      row["source_refs"] = list(dict.fromkeys(explicit_refs))
  return {
    "kind": "judgment_row",
    "artifact": ANALYST_NOTE_ARTIFACT,
    "ref": assumption_id,
    "row": row,
  }


def _cell_state(doc: ModelSheet, cell: CellKey) -> dict[str, Any]:
  """What an agent input at ``cell`` replaces: ``{item_id, period, column, override, value}`` — the
  semantic column, the constant override and the ``ValueCell``, the last two as
  ``model_dump(mode="json")`` (the shape ``clear_input`` puts in ``removed``)."""
  item = doc.model.get_item(cell[0])
  override = (item.overrides or {}).get(cell[1])
  value = item.values.values.get(cell[1]) if item.values is not None else None
  return {
    "item_id": cell[0], "period": cell[1],
    "column": copy.deepcopy(dict(doc.semantic[cell])) if cell in doc.semantic else None,
    "override": None if override is None else override.model_dump(mode="json"),
    "value": None if value is None else value.model_dump(mode="json"),
  }


def _baseline_before_agent(doc: ModelSheet, cell: CellKey) -> dict[str, Any] | None:
  """The constructed state of ``cell`` under the agent inputs it now holds: the column's
  ``edit_index`` names the edit that wrote it, whose ``replaced`` row for the cell is what that edit
  displaced; while that row's column is itself an agent edit's (an int ``edit_index``) the walk
  continues. An edit that displaced a column without a ``replaced`` row (a ``set_formula`` over a
  formula column: no input to clear) names it in its ``columns_superseded``; the walk continues
  from that column, so the chain reaches the earlier edit's ``replaced`` row and the state of a cell
  is recorded once. The last row reached is the baseline; ``None`` where no recorded edit leads back
  (the column is not an agent input's, or its edit predates ``replaced``)."""
  column, baseline, before = doc.semantic.get(cell), None, len(doc.edits)
  # each step reads an earlier edit than the last, so the walk ends
  while isinstance(column, Mapping) and isinstance(column.get("edit_index"), int) and column["edit_index"] < before:
    before = column["edit_index"]
    edit = doc.edits[before]
    row = next((row for row in edit.get("replaced") or () if (row["item_id"], row["period"]) == cell), None)
    if row is not None:
      baseline, column = row, row["column"]
      continue
    column = next(
      (row["column"] for row in edit.get("columns_superseded") or () if (row["item_id"], row["period"]) == cell),
      None,
    )
  return baseline


def _formula_baseline_before_agent(doc: ModelSheet, item_id: str) -> dict[str, Any] | None:
  """The line's ``historical`` / ``projected`` specs and ``formula_periods`` as construction left
  them, before the agent's first ``set_formula`` on it: that edit's ``replaced_formula`` row (an
  agent ``set_formula`` records the specs it displaced). ``None`` for a line the agent never set a
  formula on, or whose edits predate ``replaced_formula``."""
  for edit in doc.edits:
    op = edit.get("op") or {}
    if op.get("type") == "set_formula" and str(op.get("item_id")) == item_id and "agent" in edit:
      row = edit.get("replaced_formula")
      return None if row is None else dict(row)
  return None


def _restore_baseline(doc: ModelSheet, baseline: Mapping[str, Any]) -> None:
  """Put a ``_cell_state`` back: the override (or none), the ``ValueCell`` and the column (or none)."""
  cell = (baseline["item_id"], baseline["period"])
  with transaction(doc) as model:
    item = model.get_item(cell[0])
    if baseline["override"] is None:
      (item.overrides or {}).pop(cell[1], None)
    else:
      if item.overrides is None:
        item.overrides = {}
      item.overrides[cell[1]] = FormulaSpec.model_validate(baseline["override"])
    if baseline["value"] is not None:
      if item.values is None:
        item.values = ValueSeries()
      item.values.values[cell[1]] = ValueCell.model_validate(baseline["value"])
  if baseline["column"] is None:
    doc.semantic.pop(cell, None)
  else:
    doc.semantic[cell] = stamp_cell_source(copy.deepcopy(baseline["column"]))


def set_inputs(
  doc: ModelSheet,
  item_id: str,
  values: Mapping[int, float],
  *,
  sources: Mapping[int, Mapping[str, Any]] | None = None,
  agent: Mapping[str, Any] | None = None,
) -> OperationResult:
  """Input values of one line in one ``set_value``; ``sources`` is each period's semantic
  column, stored on the edit and as ``doc.semantic[(item_id, period)]`` (``cell: input``); a
  forecast the cell held in an actual year rides on it as ``forecast_record``
  (``supersede_forecast``). An agent op passes ``agent`` instead (``source``, ``rationale``,
  ``edit_index``): stored on the edit as ``agent`` and as each cell's column ``{cell: input,
  origin, **agent}`` — ``origin`` is ``cited_input`` when the agent's ``source`` is an object of a
  ``CITED_SOURCE_KINDS`` kind (a filing line, KPI, line identity or capture citation), else
  ``manual_input`` (free text, ``{kind: agent}``, any other object); a typed ``{kind: analyst_note |
  assumption, note}`` source is the analyst's judgment and the column's ``source`` is its
  ``analyst_note_judgment`` row (``origin`` stays ``manual_input``, the edit keeps the source as
  given). A ``sources`` column with no
  ``origin`` of its own whose ``kind`` is cited takes ``cited_input`` the same way. A fixed-cell
  line takes exactly one value, stored at its anchor
  period. Beside each constant the line's value series holds the value (``ValueCell``, provenance
  from the column's ``origin``), as the product's population writes it: the compile pieces that
  read a line's reported years (``_bound_years``) read that series.

  An agent op records what it replaced on its edit (``replaced``: each cell's column, override and
  ``ValueCell``, ``_cell_state``). An agent value equal to the constructed one, written over a cell
  that already holds an agent input (``_baseline_before_agent`` follows the chain of ``replaced``
  rows back to the constructed cell), restores that cell — its column, override note and
  ``ValueCell`` — and marks the edit ``restored_baseline``; the same value on a cell no agent has
  touched is an input like any other (an agent column), and a fixed-cell line is never restored."""
  if agent is not None:
    origin = _input_origin(agent.get("source"))
    columns: dict[int, dict[str, Any]] = {}
    for period in values:
      judgment = analyst_note_judgment(doc, agent.get("source"), item_id, int(period), agent.get("edit_index"))
      columns[int(period)] = {"origin": origin, **dict(agent), **({} if judgment is None else {"source": judgment})}
    sources = columns
  elif sources is not None:
    sources = {
      int(period): {"origin": "cited_input", **row}
      if "origin" not in row and _input_origin(row) == "cited_input" else row
      for period, row in sources.items()
    }
  item = doc.model.get_item(item_id)
  payload: dict[str, Any] = {"type": "set_value", "item_id": item_id}
  if item.column is not None:
    if len(values) != 1:
      raise ValueError(f"{item_id} is a fixed-cell line: one value, not {len(values)}")
    ((period, value),) = values.items()
    anchor = _fixed_cell_anchor_period(doc.model, item)
    payload["value"] = float(value)
    keyed = {int(period): int(anchor if anchor is not None else period)}
  else:
    payload["values"] = {int(period): float(value) for period, value in values.items()}
    keyed = {int(period): int(period) for period in values}
  edit_source = None if sources is None or agent is not None else {str(period): dict(row) for period, row in sources.items()}
  replaced = [_cell_state(doc, (item_id, cell_period)) for cell_period in keyed.values()] if agent is not None else []
  baselines = {
    cell_period: baseline
    for cell_period in keyed.values()
    if item.column is None
    and (baseline := _baseline_before_agent(doc, (item_id, cell_period))) is not None
    and baseline["value"] is not None
  }

  def _value_cells(model: FinancialModel) -> None:
    line = model.get_item(item_id)
    if line.values is None:
      line.values = ValueSeries()
    if line.formula_periods is not None and line.column is None:
      line.formula_periods = sorted({*(int(p) for p in line.formula_periods), *keyed.values()})
    for period, cell_period in keyed.items():
      origin = None if sources is None else (sources.get(period) or {}).get("origin")
      line.values.values[cell_period] = ValueCell(
        period=cell_period, value=float(values[period]),
        provenance=_VALUE_PROVENANCE.get(str(origin), ValueProvenance.input),
        lineage_reference=(
          stamp_cell_source({"cell": "input", **dict(sources[period])})
          if sources is not None and sources.get(period) is not None else None
        ),
      )

  result = _apply(doc, payload, then=_value_cells, source=edit_source, agent=agent)
  for period, cell_period in keyed.items():
    row = None if sources is None else sources.get(period)
    if row is None:
      doc.semantic.pop((item_id, cell_period), None)
    elif agent is not None:
      held = doc.semantic.get((item_id, cell_period)) or {}
      record = {FORECAST_RECORD: held[FORECAST_RECORD]} if FORECAST_RECORD in held else {}
      doc.semantic[(item_id, cell_period)] = stamp_cell_source({"cell": "input", **dict(row), **record})
    else:
      doc.semantic[(item_id, cell_period)] = supersede_forecast(doc, (item_id, cell_period), {"cell": "input", **dict(row)})
  if agent is not None:
    doc.edits[-1]["replaced"] = replaced
    restored = [
      baselines[cell_period] for period, cell_period in keyed.items()
      if cell_period in baselines and baselines[cell_period]["value"]["value"] == float(values[period])
    ]
    for baseline in restored:
      _restore_baseline(doc, baseline)
    if restored:
      doc.edits[-1]["restored_baseline"] = True
  return result


def set_input(
  doc: ModelSheet,
  item_id: str,
  period: int,
  value: float | None,
  *,
  source: Mapping[str, Any] | None = None,
  agent: Mapping[str, Any] | None = None,
) -> OperationResult | None:
  """One input value (``set_value``); ``None`` clears it (``clear_input``). ``source`` is the
  cell's semantic column (``set_inputs``); ``agent`` an agent op's note."""
  if value is None:
    clear_input(doc, item_id, period, source=source, agent=agent)
    return None
  return set_inputs(
    doc, item_id, {int(period): float(value)}, sources=None if source is None else {int(period): source}, agent=agent,
  )


def clear_input(
  doc: ModelSheet,
  item_id: str,
  period: int,
  *,
  source: Mapping[str, Any] | None = None,
  agent: Mapping[str, Any] | None = None,
) -> None:
  """Remove one period's input (value snapshot and constant override).

  ``modify_core`` has no operation that makes a value absent — ``set_value`` writes a constant
  — so this is the bench's own edit, recorded in ``doc.edits`` as such; ``cleared_by`` is ``agent``
  for an agent op and ``writer`` for a bench writer (``mark_unavailable``, the advance's gap marks):
  the recalc reason (``source_value_unavailable``) is the same as an unbound filing source's.
  ``source`` is the cleared cell's semantic column (``cell: cleared``); an agent op's ``agent``
  note is stored on the edit and as the column ``{cell: cleared, origin: manual_input, **agent}``;
  without either the column is dropped. An agent op records the cell's prior column, override and
  ``ValueCell`` on its edit as ``replaced`` (``_cell_state``), as ``set_inputs`` does. A fixed-cell
  line (``LineItem.column``) is one cell: every value and override it stores is removed, whatever
  period it is keyed to (the renderer writes any of them there), and its column is the anchor's
  (``_input_cell``), as ``set_inputs`` writes it.
  """
  cell = _input_cell(doc, item_id, period)
  replaced = [_cell_state(doc, cell)] if agent is not None else None
  item = doc.model.get_item(item_id)
  periods = (
    {int(key) for key in (item.overrides or {})} | {int(key) for key in (item.values.values if item.values else {})}
    if item.column is not None
    else {int(period)}
  )
  holds = any(
    (item.overrides and key in item.overrides) or (item.values is not None and key in item.values.values)
    for key in periods
  )
  removed: dict[str, Any] = {}
  if holds:
    with transaction(doc) as trial:
      item = trial.get_item(item_id)
      for key in sorted(periods):
        if item.overrides and key in item.overrides:
          removed.setdefault("override", {})[key] = item.overrides.pop(key).model_dump(mode="json")
        if item.values is not None and key in item.values.values:
          removed.setdefault("value", {})[key] = item.values.values.pop(key).model_dump(mode="json")
  record: dict[str, Any] = {"op": {"type": "clear_input", "item_id": item_id, "period": int(period)}}
  if source is not None:
    record["source"] = dict(source)
  if agent is not None:
    record["agent"] = dict(agent)
    record["replaced"] = replaced
    source = {"origin": "manual_input", **dict(agent)}
  doc.edits.append({
    **record,
    "result": {
      "status": "ok" if removed else "noop", "owner": "bench",
      "cleared_by": "agent" if agent is not None else "writer", "removed": removed,
    },
  })
  if source is None:
    doc.semantic.pop(cell, None)
  elif agent is not None:
    held = doc.semantic.get(cell) or {}
    history = {FORECAST_RECORD: held[FORECAST_RECORD]} if FORECAST_RECORD in held else {}
    doc.semantic[cell] = stamp_cell_source({"cell": "cleared", **dict(source), **history})
  else:
    # A writer's absence (``mark_unavailable``) supersedes a forecast the cell held, as an input does.
    doc.semantic[cell] = supersede_forecast(doc, cell, {"cell": "cleared", **dict(source)})


def mark_unavailable(
  doc: ModelSheet,
  item_id: str,
  period: int,
  *,
  reason: CellUnavailableReason,
  source: Mapping[str, Any] | None = None,
) -> None:
  """An input cell the document holds as ``#N/A(reason)``: any input there is cleared and the
  semantic column records ``cell: unavailable`` with ``reason``; ``recalc`` reports the cell —
  and every dependent whose cause it is — missing with that reason, ``render`` writes ``=NA()``."""
  clear_input(doc, item_id, period, source={**dict(source or {}), "reason": reason})
  doc.semantic[_input_cell(doc, item_id, period)]["cell"] = "unavailable"


def mark_not_separately_presented(
  doc: ModelSheet,
  item_id: str,
  period: int,
  *,
  included_in: str,
  absence_citation: Mapping[str, str],
  source: Mapping[str, Any] | None = None,
) -> None:
  """An input cell C holds as a cited inclusion: any input there is cleared (a forecast it held in
  an actual year rides on the column as ``forecast_record``), the compiler's shared marker
  (``not_separately_presented_formula``: ``=NA()`` with ``included_in`` and ``absence_citation``)
  becomes the period's override, and the column records ``cell: not_separately_presented``. Recalc
  reads the marker's own unsupported reason, never a missing input or an evidenced-absence zero;
  ``render`` writes ``=NA()``; save/reopen keeps the override. Operand sets are the compiler
  lowering's (``apply_not_separately_presented_periods``)."""
  key = _input_cell(doc, item_id, period)
  original = (doc.model.get_item(item_id).overrides or {}).get(key[1])
  spec = not_separately_presented_formula(included_in, absence_citation, original_override=original)
  column = {
    **dict(source or {}),
    "reason": "not_separately_presented",
    "source": {"kind": "absence_citation", "citation": dict(absence_citation), "included_in": included_in},
  }
  clear_input(doc, item_id, period, source=column)
  item = doc.model.get_item(item_id)
  item.overrides = {**(item.overrides or {}), key[1]: spec}
  doc.semantic[key]["cell"] = "not_separately_presented"
  record_edit(
    doc,
    {"type": "not_separately_presented", "item_id": item_id, "period": key[1]},
    source=column,
    result={"status": "ok", "formula": spec.model_dump(mode="json")},
  )


def _inclusion_role_items(doc: ModelSheet) -> dict[str, set[str]]:
  """All document lines for a role, including C identities not stamped on the template."""
  roles: dict[str, set[str]] = {}
  for item in doc.model._index.values():
    if item.required_role is not None:
      roles.setdefault(item.required_role, set()).add(item.id)
  for (item_id, _period), column in doc.semantic.items():
    role = column.get("required_role")
    if role is not None:
      roles.setdefault(role, set()).add(item_id)
  return roles


@dataclass(frozen=True)
class Protected:
  """What construction and source lowering never change: whole lines (``items``: the agent's
  added and formula-set lines) and single cells (``cells``: set or cleared cells, judgment rows)."""

  items: frozenset[str] = frozenset()
  cells: frozenset[tuple[str, int]] = frozenset()


def _agent_authored(doc: ModelSheet) -> tuple[set[str], set[tuple[str, int]]]:
  """What the agent authored through ``document.apply_ops``, read off ``doc.edits``: the lines it
  added or set a formula on (every period), and the cells it set or cleared — less the cells the
  year's actual superseded (``superseded_by_actual``: a forecast, history now). An agent op's edit
  carries ``agent`` and never a top-level ``source`` (the C, judgment and construction writers
  always record one)."""
  items: set[str] = set()
  cells: set[tuple[str, int]] = set()
  for edit in doc.edits:
    op = edit.get("op") or {}
    kind = op.get("type")
    if kind in ("add_item", "set_formula"):
      items.add(str(op["item_id"]))
    elif kind == "set_value" and "source" not in edit:
      if "values" in op:
        cells.update((str(op["item_id"]), int(period)) for period in op["values"])
        cells -= superseded_cells(edit)
      else:
        items.add(str(op["item_id"]))
    elif kind == "clear_input" and "source" not in edit:
      cells.add((str(op["item_id"]), int(op["period"])))
  return items, cells


def _protected(doc: ModelSheet) -> Protected:
  """What construction and source lowering preserve: the agent's lines and cells
  (``apply_ops``) and every judgment-row cell."""
  agent_items, agent_cells = _agent_authored(doc)
  judgment_cells = {key for key, column in doc.semantic.items() if judgment_row_ref(column.get("source")) is not None}
  return Protected(items=frozenset(agent_items), cells=frozenset(agent_cells | judgment_cells))


def reconcile_inclusions(
  doc: ModelSheet,
  requirements: Iterable[AnnualSourceRequirement] = (),
  *,
  role_items: Mapping[str, Collection[str]] | None = None,
  cleared_role_periods: Collection[tuple[str, int]] = (),
  previous_markers: Mapping[tuple[str, int], FormulaSpec] | None = None,
  protected_cells: Collection[tuple[str, int]] = (),
) -> dict[str, Any]:
  """Run the shared inclusion owner with the document's complete role identities.

  Source writes supply only their actual pins; an authored overwrite supplies an explicit
  role-period clear. Unmentioned C cells keep their columns and citations. Restored or lowered
  formulas take a formula column without losing an actual's forecast history.
  Existing agent and judgment protection retains authored values, formulas and
  operand membership in the shared owner; these are not source observations
  to lower. ``protected_cells`` supplements that document-owned protection.
  """
  protect = _protected(doc)
  periods = (*doc.model.time_structure.historical_periods, *doc.model.time_structure.projection_periods)
  protected_cells = frozenset(protected_cells) | protect.cells | frozenset(
    (item_id, int(period)) for item_id in protect.items for period in periods
  )
  identities = _inclusion_role_items(doc)
  for role, item_ids in (role_items or {}).items():
    identities.setdefault(role, set()).update(item_ids)

  def overrides(model: FinancialModel) -> dict[tuple[str, int], Any]:
    return {
      (item.id, int(period)): spec.model_dump(mode="json")
      for item in model._index.values()
      for period, spec in (item.overrides or {}).items()
    }

  before = overrides(doc.model)
  with transaction(doc) as trial:
    for (item_id, period), marker in (previous_markers or {}).items():
      item = trial.get_item(item_id)
      if (item.overrides or {}).get(period) is None:
        item.overrides = {**(item.overrides or {}), period: marker.model_copy(deep=True)}
    apply_not_separately_presented_periods(
      trial, requirements, role_items=identities, cleared_role_periods=cleared_role_periods,
      protected_cells=protected_cells,
    )
    after = overrides(trial)
    changed = sorted(key for key in before.keys() | after.keys() if before.get(key) != after.get(key))
    for item_id, period in sorted(set(changed) | set(previous_markers or ())):
      spec = _render_spec_for_period(trial, trial.get_item(item_id), period)
      if spec is None or spec.type is FormulaType.constant or spec.subtype == "not_separately_presented":
        continue
      held = doc.semantic.get((item_id, period)) or {}
      record = {FORECAST_RECORD: held[FORECAST_RECORD]} if FORECAST_RECORD in held else {}
      identity = {"required_role": held["required_role"]} if "required_role" in held else {}
      doc.semantic[(item_id, period)] = stamp_cell_source({**formula_column(spec), **identity, **record})
  cells = [{"item_id": item_id, "period": period} for item_id, period in changed]
  return {"cells_changed": cells}


_UNAVAILABLE_REASON_CODES: frozenset[str] = frozenset(get_args(MissingReasonCode)) | frozenset(get_args(AbsenceDisposition))


def unavailable_reasons(doc: ModelSheet) -> dict[tuple[str, int], CellUnavailableReason]:
  """``(item_id, period)`` → reason of every cell the document marks unavailable.

  ``mark_unavailable`` writes a ``CellUnavailableReason``; a saved package reads it back as text,
  so the read narrows it against the value algebra's own code set and refuses anything else —
  a package carrying an unknown reason is corrupt, not a new state.
  """
  out: dict[tuple[str, int], CellUnavailableReason] = {}
  for key, row in doc.semantic.items():
    if row.get("cell") != "unavailable":
      continue
    reason = str(row["reason"])
    if reason not in _UNAVAILABLE_REASON_CODES:
      raise ValueError(f"cell {key!r} carries an unknown unavailable reason {reason!r}")
    out[key] = cast_reason(reason)
  return out


def cast_reason(reason: str) -> CellUnavailableReason:
  """The typed code for a reason text already checked against ``CellUnavailableReason``."""
  for code in (*get_args(MissingReasonCode), *get_args(AbsenceDisposition)):
    if code == reason:
      return code
  raise ValueError(reason)


def remove_line(doc: ModelSheet, item_id: str, *, agent: Mapping[str, Any] | None = None) -> OperationResult:
  """Remove a line (``remove_item``); the semantic columns of its cells go with it."""
  result = _apply(doc, {"type": "remove_item", "item_id": item_id}, agent=agent)
  for key in [key for key in doc.semantic if key[0] == item_id]:
    del doc.semantic[key]
  return result


def record_declination(
  doc: ModelSheet, item_id: str, periods: Sequence[int], reason: str, *, agent: Mapping[str, Any],
) -> None:
  """The analyst's answer to a declinable warning on ``item_id`` (``periods`` empty: every period),
  recorded on ``doc.edits`` as ``{op: {type: decline, item_id, periods, reason}, agent, result}``
  with the agent's note; no cell is written. ``declinations`` reads them back."""
  doc.model.get_item(item_id)
  doc.edits.append({
    "op": {"type": "decline", "item_id": item_id, "periods": sorted({int(period) for period in periods}), "reason": reason},
    "agent": dict(agent),
    "result": {"status": "ok", "owner": "bench"},
  })


def record_note(doc: ModelSheet, text: str, *, agent: Mapping[str, Any]) -> None:
  """An analyst limitation: ``text`` recorded on ``doc.edits`` as ``{op: {type: note, text}, agent,
  result}``; no cell is written."""
  doc.edits.append({
    "op": {"type": "note", "text": text}, "agent": dict(agent), "result": {"status": "ok", "owner": "bench"},
  })


def declinations(doc: ModelSheet) -> tuple[dict[str, Any], ...]:
  """The ``decline`` ops recorded on ``doc.edits``, in order: ``{item_id, periods, reason,
  edit_index}`` each (``periods`` empty: every period), for ``publish.validate_document_graph``."""
  return tuple(
    {**{key: edit["op"][key] for key in ("item_id", "periods", "reason")}, "edit_index": index}
    for index, edit in enumerate(doc.edits)
    if (edit.get("op") or {}).get("type") == "decline"
  )


# --------------------------------------------------------------------------
# period advance
# --------------------------------------------------------------------------


#: The semantic-column key of a cell the advance made historical that still holds its forecast
#: (``advance_period``) until the year's actual supersedes it.
PROVISIONAL = "provisional_projection"
#: The semantic-column key of a cell whose actual superseded the forecast it held: that forecast.
FORECAST_RECORD = "forecast_record"


from model_build.keyed_dump import _FILING_LINE_PREFIX, parse_filing_line_id

#: Coordinate keys owned here for each structured ``CellSource`` kind at live write admission.
#: For ``filing_line``: ``ref`` and ``capture_id`` (``FilingLineSource``), ``accession`` and ``tag``
#: (``FilingLineRef``), ``observation_id`` and ``fact_id`` (``SourceFactRef``), or ``line_qname``
#: (statement-binding line identity).
JUDGMENT_ROW_REF_KEYS = ("artifact", "ref", "row")
FILING_LINE_REF_KEYS = ("ref", "accession", "capture_id", "observation_id", "line_qname", "tag", "fact_id")
KPI_REF_KEYS = ("capture_id", "ref", "row", "metric_key", "observation_id", "accession")
LINE_IDENTITY_REF_KEYS = ("ref", "line_qname", "tag", "capture_id", "statement")
CAPTURE_REF_KEYS = ("capture_id", "ref")
ABSENCE_CITATION_REF_KEYS = ("citation", "closure", "line")
CONSTRUCTION_HELPER_REF_KEYS = ("helper", "owner")


@dataclass(frozen=True)
class TextSource:
  """Free-text rationale or citation (``origin: manual_input``)."""

  text: str
  kind: Literal["text"] = "text"
  origin: Literal["manual_input"] = "manual_input"

  def to_raw(self) -> str:
    return self.text


@dataclass(frozen=True)
class JudgmentRowSource:
  """Complete judgment-artifact row citation ``{kind: judgment_row, artifact, ref, row}``."""

  artifact: Any
  ref: Any
  row: Any
  raw: dict[str, Any] = field(default_factory=dict)
  kind: Literal["judgment_row"] = "judgment_row"
  origin: Literal["judgment_row"] = "judgment_row"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw) if self.raw else {"kind": "judgment_row", "artifact": self.artifact, "ref": self.ref, "row": self.row}


@dataclass(frozen=True)
class FilingLineSource:
  """Cited filing-line reference (``origin: cited_input``)."""

  raw: dict[str, Any]
  ref: Mapping[str, Any] | None = None
  capture_id: str | None = None
  kind: Literal["filing_line"] = "filing_line"
  origin: Literal["cited_input"] = "cited_input"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw)


@dataclass(frozen=True)
class KpiSource:
  """Cited KPI reference (``origin: cited_input``)."""

  raw: dict[str, Any]
  capture_id: str | None = None
  ref: Mapping[str, Any] | None = None
  kind: Literal["kpi"] = "kpi"
  origin: Literal["cited_input"] = "cited_input"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw)


@dataclass(frozen=True)
class LineIdentitySource:
  """Cited statement line-identity reference (``origin: cited_input``)."""

  raw: dict[str, Any]
  ref: Mapping[str, Any] | None = None
  kind: Literal["line_identity"] = "line_identity"
  origin: Literal["cited_input"] = "cited_input"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw)


@dataclass(frozen=True)
class CaptureSource:
  """Cited runtime-corpus capture reference (``origin: cited_input``)."""

  raw: dict[str, Any]
  capture_id: str | None = None
  ref: Mapping[str, Any] | None = None
  kind: Literal["capture"] = "capture"
  origin: Literal["cited_input"] = "cited_input"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw)


@dataclass(frozen=True)
class AbsenceCitationSource:
  """Cited absence evidence reference (``{kind: absence_citation, citation|closure|line}``)."""

  raw: dict[str, Any]
  citation: Mapping[str, Any] | None = None
  kind: Literal["absence_citation"] = "absence_citation"
  origin: Literal["evidenced_absent"] = "evidenced_absent"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw)


@dataclass(frozen=True)
class ConstructionHelperSource:
  """Construction-helper derivation stamp (``{kind: construction_helper, helper, owner}``)."""

  helper: str
  owner: str
  raw: dict[str, Any] = field(default_factory=dict)
  kind: Literal["construction_helper"] = "construction_helper"
  origin: Literal["derived"] = "derived"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw) if self.raw else {"kind": "construction_helper", "helper": self.helper, "owner": self.owner}


@dataclass(frozen=True)
class LegacyOpaqueSource:
  """Open-extension or stored legacy/partial source object (``origin: manual_input``)."""

  raw: dict[str, Any]
  reason: str = "opaque"
  kind: Literal["legacy_opaque"] = "legacy_opaque"
  origin: Literal["manual_input"] = "manual_input"

  def to_raw(self) -> dict[str, Any]:
    return dict(self.raw)


CellSource = (
  TextSource
  | JudgmentRowSource
  | FilingLineSource
  | KpiSource
  | LineIdentitySource
  | CaptureSource
  | AbsenceCitationSource
  | ConstructionHelperSource
  | LegacyOpaqueSource
)
_CELL_SOURCE_CLASSES: tuple[type[Any], ...] = get_args(CellSource)


def _has_source_coordinate(data: Mapping[str, Any], key: str) -> bool:
  if key not in data or data[key] is None:
    return False
  val = data[key]
  if isinstance(val, (str, Mapping, Sequence)) and len(val) == 0:
    return False
  return True


def _resolve_filing_line_source(
  data: dict[str, Any], *, where: str = "source", stored: bool = False,
) -> CellSource:
  """Single resolution for ``kind == 'filing_line'``: decodes any canonical ``filing-line:`` id via
  ``keyed_dump.parse_filing_line_id`` onto its one declared key ``observation_id`` (and ``ref``)."""
  raw_obs = data.get("observation_id")
  raw_id = data.get("id")
  candidate = raw_obs if isinstance(raw_obs, str) and raw_obs.strip() else raw_id
  if isinstance(candidate, str) and (candidate.strip().startswith(_FILING_LINE_PREFIX) or "id" in data):
    text = candidate.strip()
    parsed = parse_filing_line_id(text)
    if parsed is None:
      if stored:
        return LegacyOpaqueSource(raw=data, reason="incomplete_filing_line")
      raise OpUsage(
        f"{where}: filing_line id {candidate!r} is not a canonical '{_FILING_LINE_PREFIX}<accession>:FY<year>:<statement_role>:<tag>:<dimensions>' observation_id; "
        f"pass a canonical '{_FILING_LINE_PREFIX}...' id or a filing coordinate ({', '.join(FILING_LINE_REF_KEYS)})"
      )
    data = {k: v for k, v in data.items() if k != "id"}
    data["observation_id"] = text
    if not isinstance(data.get("ref"), Mapping):
      data["ref"] = {k: v for k, v in parsed.items() if k != "fiscal_year"}
  if not any(_has_source_coordinate(data, key) for key in FILING_LINE_REF_KEYS):
    if stored:
      return LegacyOpaqueSource(raw=data, reason="incomplete_filing_line")
    raise OpUsage(
      f"{where}: a filing_line source requires a filing coordinate ({', '.join(FILING_LINE_REF_KEYS)}); "
      f"cite the filing line or pass source as text"
    )
  ref = data.get("ref")
  cap = data.get("capture_id")
  return FilingLineSource(
    raw=data,
    ref=ref if isinstance(ref, Mapping) else None,
    capture_id=str(cap) if cap is not None else None,
  )


def _parse_source(raw: Any, *, where: str = "source", stored: bool = False) -> CellSource | None:
  if raw is None:
    return None
  if isinstance(raw, _CELL_SOURCE_CLASSES):
    return raw
  if isinstance(raw, str):
    return TextSource(text=raw)
  if not isinstance(raw, Mapping):
    if stored:
      return LegacyOpaqueSource(raw={"value": raw}, reason="non_mapping_source")
    raise OpUsage(f"{where}: source must be text or an object (a judgment-row ref, evidence refs), not {raw!r}")
  data = dict(raw)
  kind = data.get("kind")
  if kind == "judgment_row":
    missing = [key for key in JUDGMENT_ROW_REF_KEYS if key not in data or data[key] is None]
    if missing:
      if stored:
        return LegacyOpaqueSource(raw=data, reason="incomplete_judgment_row")
      raise OpUsage(
        f"{where}: a judgment_row source is {{kind: judgment_row, artifact, ref, row}} (the judgment "
        f"artifact, its ref and the cited row); missing {missing}; cite the row or pass source as text"
      )
    return JudgmentRowSource(artifact=data["artifact"], ref=data["ref"], row=data["row"], raw=data)
  if kind == "filing_line":
    return _resolve_filing_line_source(data, where=where, stored=stored)
  if kind == "kpi":
    if not any(_has_source_coordinate(data, key) for key in KPI_REF_KEYS):
      if stored:
        return LegacyOpaqueSource(raw=data, reason="incomplete_kpi")
      raise OpUsage(
        f"{where}: a kpi source requires a KPI coordinate ({', '.join(KPI_REF_KEYS)}); "
        f"cite the KPI row or pass source as text"
      )
    ref = data.get("ref")
    cap = data.get("capture_id")
    return KpiSource(
      raw=data,
      capture_id=str(cap) if cap is not None else None,
      ref=ref if isinstance(ref, Mapping) else None,
    )
  if kind == "line_identity":
    if not any(_has_source_coordinate(data, key) for key in LINE_IDENTITY_REF_KEYS):
      if stored:
        return LegacyOpaqueSource(raw=data, reason="incomplete_line_identity")
      raise OpUsage(
        f"{where}: a line_identity source requires a line coordinate ({', '.join(LINE_IDENTITY_REF_KEYS)}); "
        f"cite the line or pass source as text"
      )
    ref = data.get("ref")
    return LineIdentitySource(raw=data, ref=ref if isinstance(ref, Mapping) else None)
  if kind == "capture":
    if not any(_has_source_coordinate(data, key) for key in CAPTURE_REF_KEYS):
      if stored:
        return LegacyOpaqueSource(raw=data, reason="incomplete_capture")
      raise OpUsage(
        f"{where}: a capture source requires a capture coordinate ({', '.join(CAPTURE_REF_KEYS)}); "
        f"cite the capture or pass source as text"
      )
    ref = data.get("ref")
    cap = data.get("capture_id")
    return CaptureSource(
      raw=data,
      capture_id=str(cap) if cap is not None else None,
      ref=ref if isinstance(ref, Mapping) else None,
    )
  if kind == "absence_citation":
    if not any(_has_source_coordinate(data, key) for key in ABSENCE_CITATION_REF_KEYS):
      if stored:
        return LegacyOpaqueSource(raw=data, reason="incomplete_absence_citation")
      raise OpUsage(
        f"{where}: an absence_citation source requires a citation coordinate ({', '.join(ABSENCE_CITATION_REF_KEYS)}); "
        f"cite the absence or pass source as text"
      )
    cit = data.get("citation")
    return AbsenceCitationSource(raw=data, citation=cit if isinstance(cit, Mapping) else None)
  if kind == "construction_helper":
    missing = [
      key for key in CONSTRUCTION_HELPER_REF_KEYS
      if not isinstance(data.get(key), str) or not str(data[key]).strip()
    ]
    if missing:
      if stored:
        return LegacyOpaqueSource(raw=data, reason="incomplete_construction_helper")
      raise OpUsage(
        f"{where}: a construction_helper source requires ['helper', 'owner']; missing {missing}"
      )
    return ConstructionHelperSource(helper=str(data["helper"]), owner=str(data["owner"]), raw=data)
  if kind in ANALYST_NOTE_KINDS and "source_refs" in data and data["source_refs"] is not None:
    if not stored:
      data["source_refs"] = _parse_analyst_note_source_refs(data["source_refs"], where=where)
  return LegacyOpaqueSource(raw=data, reason="agent_note" if kind == "agent" else "opaque_kind")


def parse_new_source(raw: Any, *, where: str = "source") -> CellSource | None:
  """Strict write-admission parser for ``source`` (``stored=False``)."""
  return _parse_source(raw, where=where, stored=False)


def parse_stored_source(raw: Any, *, where: str = "source") -> CellSource | None:
  """Tolerant replay/reader parser for ``source`` (``stored=True``); never raises ``OpUsage``."""
  return _parse_source(raw, where=where, stored=True)


def judgment_row_ref(source: Any) -> dict[str, Any] | None:
  """``{artifact, ref, row}`` when ``source`` is a whole ``JudgmentRowSource``, else ``None``: a
  ``{kind: judgment_row}`` without them (an op the door admitted before it refused the shape) is a
  manual input with a note, and reads as one everywhere a cell's lineage is read."""
  parsed = parse_stored_source(source)
  if not isinstance(parsed, JudgmentRowSource):
    return None
  return {"artifact": parsed.artifact, "ref": parsed.ref, "row": parsed.row}


def forecast_origin(column: Mapping[str, Any] | None) -> str:
  """What a forecast cell's value came from: ``judgment_row`` (a judgment artifact's row),
  ``manual_input`` (the agent's op), ``construction`` (a construction helper's seed) or
  ``projection`` (the line's projected formula)."""
  column = column or {}
  raw_source = column.get("source")
  parsed = parse_stored_source(raw_source)
  if isinstance(parsed, JudgmentRowSource):
    return "judgment_row"
  if column.get("origin") in ("manual_input", "cited_input"):
    return "manual_input"
  kind = raw_source.get("kind") if isinstance(raw_source, Mapping) else None
  if isinstance(parsed, ConstructionHelperSource) or kind == "construction_helper" or "constructed_by" in column:
    return "construction"
  return "projection"


def _forecast_edit_index(doc: ModelSheet, key: CellKey, column: Mapping[str, Any] | None) -> int | None:
  """The edit in ``doc.edits`` that wrote the forecast ``column`` holds at ``key``: the agent's
  (its note's ``edit_index``), else the latest construction edit of the helper that wrote it,
  else the latest ``set_value`` covering the cell; ``None`` for a projection (no edit wrote it)."""
  if column is None:
    return None
  if isinstance(column.get("edit_index"), int):
    return int(column["edit_index"])
  helper = column.get("constructed_by")
  if helper is None and (column.get("source") or {}).get("kind") == "construction_helper":
    helper = column["source"].get("helper")
  for index in range(len(doc.edits) - 1, -1, -1):
    op = doc.edits[index].get("op") or {}
    if helper is not None:
      if op.get("type") == "construction" and op.get("helper") == helper:
        return index
    elif op.get("type") == "set_value" and op.get("item_id") == key[0] and (
      "value" in op or str(key[1]) in {str(period) for period in op.get("values") or {}}
    ):
      return index
  return None


def _forecast_record(
  doc: ModelSheet, key: CellKey, column: Mapping[str, Any] | None, value: float | None, last_actual: int,
) -> dict[str, Any]:
  """The forecast a cell holds as the advance makes its year actual: ``value``, ``origin``
  (``forecast_origin``), the column's ``rationale`` and ``source``, the ``edit_index`` of the edit
  that wrote it, and the axis it was made under (``recorded_at_axis: {last_actual}``)."""
  record: dict[str, Any] = {"value": value, "origin": forecast_origin(column)}
  for name in ("rationale", "source"):
    if column is not None and column.get(name) is not None:
      record[name] = column[name]
  edit_index = _forecast_edit_index(doc, key, column)
  if edit_index is not None:
    record["edit_index"] = edit_index
  record["recorded_at_axis"] = {"last_actual": int(last_actual)}
  return record


def supersede_forecast(doc: ModelSheet, key: CellKey, column: Mapping[str, Any]) -> CellSourceStamp:
  """The column the writer of an actual (a sourced input, an evidenced absence, the historical
  formula) puts on ``key``. A forecast the cell still holds (``PROVISIONAL``) rides on it as
  ``forecast_record``, and the edit that wrote that forecast stays in ``doc.edits`` marked
  ``superseded_by_actual`` (the cell and the advance's edit index); a record the cell already
  carries is kept."""
  held = doc.semantic.get(key) or {}
  out = stamp_cell_source(column)
  mark = held.get(PROVISIONAL)
  record = None if mark is None else mark.get("forecast")
  if mark is not None and record is not None:
    out[FORECAST_RECORD] = record
    edit_index = record.get("edit_index")
    if isinstance(edit_index, int) and 0 <= edit_index < len(doc.edits):
      entry = {"item_id": key[0], "period": int(key[1]), "advance_edit_index": mark["edit_index"]}
      marks = doc.edits[edit_index].setdefault("superseded_by_actual", [])
      if entry not in marks:
        marks.append(entry)
  elif mark is None and FORECAST_RECORD in held:
    out.setdefault(FORECAST_RECORD, held[FORECAST_RECORD])
  return out


def superseded_cells(edit: Mapping[str, Any]) -> set[CellKey]:
  """The cells of one edit an actual superseded (``supersede_forecast``)."""
  return {(str(row["item_id"]), int(row["period"])) for row in edit.get("superseded_by_actual") or ()}


def formula_column(spec: FormulaSpec) -> CellSourceStamp:
  """The semantic column of a cell its own formula computes (``origin: derived``)."""
  operands = sorted({ref.id for ref in DependencyGraph()._extract_refs(spec)})
  return stamp_cell_source({
    "cell": "formula", "origin": "derived",
    "source": {"kind": "formula", "formula": formula_text(spec), "operands": operands},
  })


def axis(model: FinancialModel) -> dict[str, list[int]]:
  """The model's time axis: ``historical`` and ``projection`` years."""
  ts = model.time_structure
  return {
    "historical": [int(period) for period in (ts.historical_periods or ts.historical_years)],
    "projection": [int(period) for period in (ts.projection_periods or ts.projection_years)],
  }


def _same_spec(left: FormulaSpec | None, right: FormulaSpec | None) -> bool:
  if left is None or right is None:
    return left is right
  # Rule (d)'s owed operands describe missing-value repair, not the formula a cell runs.
  exclude = {"params": {"_owed_operands"}}
  return left.model_dump(mode="json", exclude=exclude) == right.model_dump(mode="json", exclude=exclude)


def _shift_fixed_cell(item: LineItem) -> None:
  """A fixed-cell line keeps its place against the actual/forecast boundary (as the product's
  positional remap places it): every period it keys moves one year later."""
  if item.formula_periods is not None:
    item.formula_periods = [int(period) + 1 for period in item.formula_periods]
  if item.overrides:
    item.overrides = {int(period) + 1: spec for period, spec in item.overrides.items()}
  if item.values is not None and item.values.values:
    item.values.values = {
      int(period) + 1: cell.model_copy(update={"period": int(period) + 1}) for period, cell in item.values.values.items()
    }


def advance_period(
  doc: ModelSheet,
  new_actual_year: int,
  *,
  source: Mapping[str, Any] | None = None,
  agent: Mapping[str, Any] | None = None,
  values: Values | None = None,
  graphs: _BatchGraphs | None = None,
) -> dict[str, Any]:
  """Move the actual/forecast boundary one year: ``new_actual_year`` (the first projection year)
  becomes historical and the projection window gains the year after its last.

  The axis is moved by the product's own remap (``schema.build_time_structure.remap_time_structure``)
  with the year map a build of the later axis uses — every period one year later, a historical
  window one year longer — so the product decides what follows the boundary: each line's formula
  periods from the old last actual year on (a projection-only formula leaves the new actual year,
  a roll-forward that starts at the last actual year starts at the new one, a historical formula
  valid in the new actual year takes it, a line whose formula periods span the projection reaches
  the new projection year, as do a sparse projection sentinel, a self roll-forward and a
  period-relative offset scenario). The earlier actual years keep the document's formula periods
  (the old last actual year takes the one before it's), and every input (a constant override, a
  stored value) and semantic column stays at its ``(item, year)`` — the agent's edits included.
  What the product places by position against the boundary keeps its position: a formula override
  in a projection year (a template's first-projection-year bridge, a seeded step) moves one year
  later, yielding to an input already held there (``formula_overrides_yielded``), and a fixed-cell
  line (``LineItem.column``) moves every period it keys one year later.

  Every forecast the new actual year holds stays live until the year's actual supersedes it, and
  is marked ``provisional_projection: {edit_index, held, pinned, historical, forecast}`` where
  ``forecast`` is the record the actual will carry (``_forecast_record``: value, origin, rationale,
  source, the edit that wrote it, ``recorded_at_axis``). The agent's input and a judgment row's
  keep their column and value, the mark added. A cell whose selected spec the move changes keeps
  running the spec it ran before (an override, ``held: formula``; a stored value as a constant,
  ``held: input``), and so does a construction helper's projection seed there (a constant whose
  column's ``source.kind`` is ``construction_helper``), with the semantic column ``{cell, origin:
  provisional_projection, provisional_projection, column_before?}``, so nothing goes blank. The
  actual wins: a sourced input or an evidenced absence written there, or — ``release_provisional``
  — the line's historical formula once nothing provisional feeds it, supersedes the forecast
  (``supersede_forecast``). ``values`` is the document's recalc before the move (computed when a
  held formula's value is needed and none is given). One ``advance_period`` edit records the axes
  before and after (``source`` for a bench writer, ``agent`` for an agent op). Returns the edit's
  result."""
  from schema.build_time_structure import remap_time_structure

  before_axis = axis(doc.model)
  historical, projection = before_axis["historical"], before_axis["projection"]
  year = int(new_actual_year)
  if not projection or year != projection[0]:
    raise OpUsage(f"advance_period: new_actual_year must be the first projection year {projection[:1]}, not {year}")
  if len(projection) < 2:
    raise OpUsage("advance_period: the projection window must hold at least two years to move")
  span = [*historical, *projection]
  if span != list(range(span[0], span[0] + len(span))):
    raise OpUsage(f"advance_period: the axis {span} is not one run of consecutive years")
  added = projection[-1] + 1
  graph = (_BatchGraphs(doc) if graphs is None else graphs).cells()
  before = {item.id: graph._spec_for_period(item, year) for item in doc.model._index.values() if item.column is None}
  trial = copy_model(doc.model)
  trial.build_index()
  time = trial.time_structure
  # The year map of a build of the later axis: a historical window one year longer (a year before
  # the first stands in for the one the build adds at the front) and every period one year later.
  time.historical_periods, time.historical_years = [historical[0] - 1, *historical], [historical[0] - 1, *historical]
  time.projection_periods, time.projection_years = list(projection), list(projection)
  remap_time_structure(trial, most_recent_fy=year, n_historical=len(historical) + 1, n_projection=len(projection))
  last = historical[-1]
  extended: list[str] = []
  shifted: list[str] = []
  moved_overrides = 0
  yielded: list[dict[str, Any]] = []
  projection_years = set(projection)
  for item in trial._index.values():
    old = doc.model._index[item.id]
    # The remap's year map is the identity, but it drops an override keyed outside the axis: the
    # overrides are the document's, placed here.
    if item.column is not None:
      item.overrides = None if old.overrides is None else {
        int(period): spec.model_copy(deep=True) for period, spec in old.overrides.items()
      }
      item.formula_periods = None if old.formula_periods is None else list(old.formula_periods)
      _shift_fixed_cell(item)
      shifted.append(item.id)
      continue
    inputs: dict[int, FormulaSpec] = {}
    formulas: dict[int, FormulaSpec] = {}
    for period, spec in (old.overrides or {}).items():
      if int(period) in projection_years and spec.type is not FormulaType.constant:
        formulas[int(period) + 1] = spec.model_copy(deep=True)
      else:
        inputs[int(period)] = spec.model_copy(deep=True)
    moved_overrides += len(formulas)
    yielded.extend({"item_id": item.id, "period": period} for period in sorted(formulas) if period in inputs)
    item.overrides = {**formulas, **inputs} or None
    if old.formula_periods is None or item.formula_periods is None:
      continue
    held_periods = {int(period) for period in old.formula_periods}
    # Earlier actual years keep the document's periods; the old last actual year now plays the role
    # the year before it did; from the new actual year on, the remap's positional decision.
    earlier = {period for period in held_periods if period < last}
    role = last - 1 if len(historical) > 1 else last
    old_last = {last} if role in held_periods else set()
    item.formula_periods = sorted(earlier | old_last | {int(period) for period in item.formula_periods if int(period) >= year})
    if added in item.formula_periods:
      extended.append(item.id)
  trial.build_index()
  after_graph = _cell_graph(trial)
  edit_index = len(doc.edits)
  provisional: list[dict[str, Any]] = []
  forecasts = values

  def forecast(key: CellKey, pin: FormulaSpec, column: Mapping[str, Any] | None) -> dict[str, Any]:
    nonlocal forecasts
    if pin.type is FormulaType.constant:
      value = pin.params.get("value")
    else:
      forecasts = recalc(doc) if forecasts is None else forecasts
      held = forecasts.get(key)
      value = held.value if held is not None and held.state == "known" else None
    return _forecast_record(doc, key, column, None if value is None else float(value), last)

  for item in trial._index.values():
    if item.column is not None:
      continue
    key = (item.id, year)
    held_spec = before[item.id]
    now = after_graph._spec_for_period(item, year)
    prior = doc.semantic.get(key)
    if (
      prior is not None and prior.get("cell") == "input" and PROVISIONAL not in prior
      and forecast_origin(prior) in ("manual_input", "judgment_row")
      and held_spec is not None and held_spec.type is FormulaType.constant
    ):
      # The agent's or a judgment row's forecast: its column and value stay, marked; what the
      # line computes there without it is what the actual year selects.
      own_input = item.overrides.pop(year) if item.overrides and year in item.overrides else None
      own = after_graph._spec_for_period(item, year)
      if own_input is not None:
        item.overrides = {**(item.overrides or {}), year: own_input}
      mark = {
        "edit_index": edit_index, "held": "input", "pinned": formula_text(held_spec), "historical": formula_text(own),
        "forecast": forecast(key, held_spec, prior),
      }
      doc.semantic[key] = stamp_cell_source({**prior, PROVISIONAL: mark})
      provisional.append({"item_id": item.id, "period": year, **mark})
      continue
    # A construction helper's projection seed (``source.kind: construction_helper``) is the
    # projection's value too: what the line computes there without it is what the year now selects.
    seeded = (
      held_spec is not None and held_spec.type is FormulaType.constant
      and ((prior or {}).get("source") or {}).get("kind") == "construction_helper"
    )
    if seeded:
      seed = item.overrides.pop(year) if item.overrides and year in item.overrides else None
      now = after_graph._spec_for_period(item, year)
      if seed is not None:
        item.overrides = {**(item.overrides or {}), year: seed}
    elif _same_spec(held_spec, now):
      continue
    if held_spec is None:
      stored = None if item.values is None else item.values.values.get(year)
      if stored is None or stored.value is None:
        continue
      pin = FormulaSpec(type=FormulaType.constant, params={"value": float(stored.value)})
    else:
      pin = held_spec.model_copy(deep=True)
    held = "input" if pin.type is FormulaType.constant else "formula"
    item.overrides = {**(item.overrides or {}), year: pin}
    mark = {
      "edit_index": edit_index, "held": held, "pinned": formula_text(pin), "historical": formula_text(now),
      "forecast": forecast(key, pin, prior),
    }
    # ``origin: provisional_projection``: the projection's value, never a filing's or a derivation's.
    column: dict[str, Any] = {"cell": held, "origin": PROVISIONAL, PROVISIONAL: mark}
    if held == "formula":
      column["operands"] = sorted({ref.id for ref in after_graph._extract_refs(pin)})
    if prior is not None:
      column["column_before"] = prior
    doc.semantic[key] = stamp_cell_source(column)
    provisional.append({"item_id": item.id, "period": year, **mark})
  moved = set(shifted)
  doc.semantic = {
    ((item_id, period + 1) if item_id in moved else (item_id, period)): column
    for (item_id, period), column in doc.semantic.items()
  }
  doc.model = trial
  result = {
    "status": "ok",
    "axis_before": before_axis,
    "axis_after": axis(trial),
    "new_actual_year": year,
    "new_projection_year": added,
    "provisional_projection": provisional,
    "formula_periods_extended": sorted(extended),
    "fixed_cells_shifted": sorted(shifted),
    "formula_overrides_shifted": moved_overrides,
    "formula_overrides_yielded": yielded,
  }
  record: dict[str, Any] = {"op": {"type": "advance_period", "new_actual_year": year}}
  if source is not None:
    record["source"] = dict(source)
  if agent is not None:
    record["agent"] = dict(agent)
  doc.edits.append({**record, "result": result})
  return result


def provisional_cells(doc: ModelSheet) -> dict[CellKey, dict[str, Any]]:
  """Every cell still holding its forecast in an actual year (``advance_period``) → its mark."""
  return {key: column[PROVISIONAL] for key, column in doc.semantic.items() if PROVISIONAL in column}


def release_provisional(doc: ModelSheet) -> list[dict[str, Any]]:
  """Hand each provisional cell whose own spec there is a formula (the line's historical formula)
  to that formula once nothing it reads, at any depth, is still provisional; passes repeat until
  none releases a cell. The released cell's column is the formula's (``formula_column``) carrying
  the forecast it held as ``forecast_record`` (``supersede_forecast``). A provisional cell of a
  line with no historical formula stays until a sourced input replaces it. One
  ``release_provisional`` edit lists the released cells."""
  released: list[dict[str, Any]] = []
  pending = set(provisional_cells(doc))
  while pending:
    model = doc.model
    graph = _cell_graph(model)
    # Every cell that reads a provisional one, at any depth: what a released formula must not read.
    tainted = pending | _reachable(_cell_readers(graph, model), pending)
    freed: list[tuple[CellKey, FormulaSpec]] = []
    for key in sorted(pending):
      item = model.get_item(key[0])
      overrides = dict(item.overrides or {})
      pin = overrides.pop(key[1], None)
      item.overrides = overrides or None
      own = graph._spec_for_period(item, key[1])
      reads = set() if own is None else graph.cell_operands(item, key[1])
      if own is None or reads & tainted:
        if pin is not None:
          item.overrides = {**overrides, key[1]: pin}
        continue
      freed.append((key, own))
    if not freed:
      break
    for key, own in freed:
      item = model.get_item(key[0])
      # The cell is the formula's now: a stored value beside the pin (a seed's ``ValueCell``) goes too.
      if item.values is not None:
        item.values.values.pop(key[1], None)
      column = supersede_forecast(doc, key, formula_column(own))
      mark = doc.semantic.pop(key)[PROVISIONAL]
      if FORECAST_RECORD in column:
        doc.semantic[key] = column
      released.append({"item_id": key[0], "period": key[1], "formula": formula_text(own), "pinned": mark["pinned"]})
      pending.discard(key)
  if released:
    record_edit(doc, {"type": "release_provisional"}, source={"origin": "period_advance"}, result={"status": "ok", "released": released})
  return released


# --------------------------------------------------------------------------
# ops (the doc_edit batch)
# --------------------------------------------------------------------------


def _source_schema() -> dict[str, Any]:
  """Describe the source parser's known coordinates without closing its extension objects."""
  variants = []
  for cls, coordinates, required in (
    (JudgmentRowSource, JUDGMENT_ROW_REF_KEYS, True),
    (FilingLineSource, FILING_LINE_REF_KEYS, False),
    (KpiSource, KPI_REF_KEYS, False),
    (LineIdentitySource, LINE_IDENTITY_REF_KEYS, False),
    (CaptureSource, CAPTURE_REF_KEYS, False),
    (AbsenceCitationSource, ABSENCE_CITATION_REF_KEYS, False),
    (ConstructionHelperSource, CONSTRUCTION_HELPER_REF_KEYS, True),
  ):
    coordinate_schema = (
      {"type": "string", "pattern": r"\S"} if cls is ConstructionHelperSource else {"not": {"type": "null"}}
    ) if required else {}
    variant: dict[str, Any] = {
      "type": "object",
      "properties": {"kind": {"const": cls.kind}, **{key: coordinate_schema for key in coordinates}},
      "required": ["kind", *coordinates] if required else ["kind"],
      "additionalProperties": True,
    }
    if not required:
      any_of_keys = (*coordinates, "id") if cls is FilingLineSource else coordinates
      variant["anyOf"] = [{
        "required": [key],
        "properties": {key: {"not": {"enum": [None, "", {}, []]}}},
      } for key in any_of_keys]
    variants.append(variant)
  return {"anyOf": [
    {"type": "string"}, {"type": "null"}, *variants,
    {
      "type": "object",
      "properties": {"kind": {"not": {"enum": [variant["properties"]["kind"]["const"] for variant in variants]}}},
      "additionalProperties": True,
      "description": (
        "Other source kinds are open analyst extensions. {kind: analyst_note | assumption, note: <text>} "
        "records the value as the analyst's own judgment (the note is its rationale)."
      ),
    },
  ]}


@dataclass(frozen=True)
class _OpSpec:
  required: Mapping[str, dict[str, Any]]
  optional: Mapping[str, dict[str, Any]] = field(default_factory=dict)
  forms: tuple[tuple[str, ...], ...] = ()
  aliases: Mapping[str, str] = field(default_factory=dict)

  @property
  def properties(self) -> dict[str, dict[str, Any]]:
    return {**self.required, **self.optional}


_PERIOD_SCHEMA = {"anyOf": [{"type": "integer"}, {"type": "string"}]}
_VALUE_SCHEMA = {"anyOf": [{"type": "number"}, {"type": "string"}, {"type": "null"}]}
_ITEM_REF_SCHEMA = {"anyOf": [
  {"type": "string"},
  {"type": "object", "properties": {"$ref": {"type": "integer", "minimum": 0}}, "required": ["$ref"], "additionalProperties": False},
]}
_NOTE_FIELDS = {"rationale": {"type": ["string", "null"]}, "source": _source_schema()}

#: The edit grammar's keys, alternative forms, field schemas and aliases. Both the parser and
#: the model-facing tool schema read this authority; source coordinates remain source-owned.
OP_SPECS: dict[str, _OpSpec] = {
  "add_line": _OpSpec(
    {"section": {"type": "string"}, "label": {"type": "string"}, "unit": {"type": "string"}},
    {
      "after_item_id": {"anyOf": [*_ITEM_REF_SCHEMA["anyOf"], {"type": "null"}]},
      "formula": {"type": ["string", "null"]},
      "inputs": {"anyOf": [
        {"type": "object", "additionalProperties": {"anyOf": [{"type": "number"}, {"type": "string"}]}},
        {"type": "array", "maxItems": 0}, {"type": "null"},
      ]},
      **_NOTE_FIELDS,
    },
    aliases={"statement": "section"},
  ),
  "set_formula": _OpSpec(
    {"item_id": _ITEM_REF_SCHEMA, "formula_expr": {"type": "string"}},
    {"apply_to": {"anyOf": [
      {"type": "string", "enum": ["projected", "historical", "both", "all"]},
      {"type": "array", "items": _PERIOD_SCHEMA, "minItems": 1},
    ]}, **_NOTE_FIELDS},
  ),
  "set_input": _OpSpec(
    {"item_id": _ITEM_REF_SCHEMA},
    {
      "period": _PERIOD_SCHEMA, "value": _VALUE_SCHEMA,
      "values": {"type": "object", "additionalProperties": _VALUE_SCHEMA},
      "periods": {"type": "array", "items": _PERIOD_SCHEMA}, **_NOTE_FIELDS,
    },
    forms=(("period", "value"), ("values",), ("periods", "value")),
  ),
  "clear_input": _OpSpec({"item_id": _ITEM_REF_SCHEMA, "period": _PERIOD_SCHEMA}, _NOTE_FIELDS),
  "remove_line": _OpSpec({"item_id": _ITEM_REF_SCHEMA}, _NOTE_FIELDS),
  "advance_period": _OpSpec({"new_actual_year": _PERIOD_SCHEMA}, _NOTE_FIELDS),
  "decline": _OpSpec({
    "item_id": _ITEM_REF_SCHEMA, "periods": {"type": "array", "items": _PERIOD_SCHEMA},
    "reason": {"type": "string", "minLength": 1, "maxLength": FINDING_MESSAGE_MAX_LENGTH},
  }),
  "note": _OpSpec({"text": {"type": "string", "minLength": 1}}),
}


def edit_op_schema() -> dict[str, Any]:
  """A fresh JSON Schema for one edit op, derived from ``OP_SPECS`` (including aliases)."""
  variants: list[dict[str, Any]] = []
  for kind, spec in OP_SPECS.items():
    properties = spec.properties
    branch: dict[str, Any] = {
      "type": "object", "properties": {"op": {"const": kind}, **properties},
      "required": ["op", *(key for key in spec.required if key not in spec.aliases.values())],
      "additionalProperties": False,
    }
    if spec.forms:
      form_keys = {key for form in spec.forms for key in form}
      branch["oneOf"] = [{
        "required": list(form),
        "not": {"anyOf": [{"required": [key]} for key in sorted(form_keys - set(form))]},
      } for form in spec.forms]
    if spec.aliases:
      branch["allOf"] = []
      for alias, key in spec.aliases.items():
        branch["properties"][alias] = properties[key]
        branch["allOf"].append({"oneOf": [{"required": [key]}, {"required": [alias]}]})
    variants.append(branch)
  return copy.deepcopy({"oneOf": variants})


def _expected_keys(kind: str) -> str:
  spec = OP_SPECS[kind]
  expected = ", ".join([*spec.required, *(f"{key}?" for key in spec.optional)])
  if spec.forms:
    expected += "; accepted shapes: " + " or ".join(
      "(" + ", ".join([*spec.required, *form]) + ")" for form in spec.forms
    )
  return expected


_FY_PERIOD_RE = re.compile(r"^(?:FY)?(\d{4})[AE]?$", re.IGNORECASE)


def _op_int(where: str, key: str, value: Any) -> int:
  """An integer period: an int, an integral float (``2022.0``), or an integer / FY period string."""
  if isinstance(value, int) and not isinstance(value, bool):
    return value
  if isinstance(value, float) and value.is_integer():
    return int(value)
  if isinstance(value, str):
    stripped = value.strip()
    try:
      return int(stripped)
    except ValueError:
      match = _FY_PERIOD_RE.match(stripped)
      if match is not None:
        return int(match.group(1))
  raise OpUsage(f"{where}: {key} must be an integer period, not {value!r}")


def _op_float(where: str, key: str, value: Any) -> float:
  if isinstance(value, bool) or not isinstance(value, (int, float, str)):
    raise OpUsage(f"{where}: {key} must be a number, not {value!r}")
  try:
    return float(value)
  except ValueError:
    raise OpUsage(f"{where}: {key} must be a number, not {value!r}") from None


def _resolve_item_ref(where: str, key: str, value: Any, index: int, results: Sequence[Mapping[str, Any]]) -> Any:
  """``{"$ref": N}`` → the id ``add_line`` op ``N`` (an earlier op of this batch) minted."""
  if not isinstance(value, Mapping):
    return value
  target = value.get("$ref") if set(value) == {"$ref"} else None
  if (
    isinstance(target, bool) or not isinstance(target, int) or not 0 <= target < index
    or results[target]["op"] != "add_line"
  ):
    raise OpUsage(f"{where}: {key} {{'$ref': N}} must name an earlier add_line op by its index (< {index})")
  return results[target]["item_id"]


@dataclass(frozen=True)
class _EditOpBase:
  _raw_dict: dict[str, Any] = field(default_factory=dict, repr=False, compare=False)

  def to_op_dict(self) -> dict[str, Any]:
    return dict(self._raw_dict)


@dataclass(frozen=True)
class AddLineOp(_EditOpBase):
  section: str = ""
  label: str = ""
  unit: str = ""
  after_item_id: str | None = None
  formula: str | None = None
  inputs: dict[int, float] | None = None
  rationale: str | None = None
  source: CellSource | None = None
  op: Literal["add_line"] = "add_line"


@dataclass(frozen=True)
class SetFormulaOp(_EditOpBase):
  item_id: str = ""
  formula_expr: str = ""
  apply_to: ApplyTo | list[int] = "both"
  rationale: str | None = None
  source: CellSource | None = None
  op: Literal["set_formula"] = "set_formula"


@dataclass(frozen=True)
class SetInputOp(_EditOpBase):
  item_id: str = ""
  period: int | None = None
  value: float | None = None
  values: dict[int, float | None] | None = None
  periods: list[int] | None = None
  rationale: str | None = None
  source: CellSource | None = None
  op: Literal["set_input"] = "set_input"


@dataclass(frozen=True)
class ClearInputOp(_EditOpBase):
  item_id: str = ""
  period: int = 0
  rationale: str | None = None
  source: CellSource | None = None
  op: Literal["clear_input"] = "clear_input"


@dataclass(frozen=True)
class RemoveLineOp(_EditOpBase):
  item_id: str = ""
  rationale: str | None = None
  source: CellSource | None = None
  op: Literal["remove_line"] = "remove_line"


@dataclass(frozen=True)
class AdvancePeriodOp(_EditOpBase):
  new_actual_year: int = 0
  rationale: str | None = None
  source: CellSource | None = None
  op: Literal["advance_period"] = "advance_period"


@dataclass(frozen=True)
class DeclineOp(_EditOpBase):
  item_id: str = ""
  periods: list[int] = field(default_factory=list)
  reason: str = ""
  op: Literal["decline"] = "decline"


@dataclass(frozen=True)
class NoteOp(_EditOpBase):
  text: str = ""
  op: Literal["note"] = "note"


EditOp = (
  AddLineOp
  | SetFormulaOp
  | SetInputOp
  | ClearInputOp
  | RemoveLineOp
  | AdvancePeriodOp
  | DeclineOp
  | NoteOp
)


def _parse_op(
  index: int, raw: Any, results: Sequence[Mapping[str, Any]] = (), *, stored: bool = False,
) -> EditOp:
  """Parse one op into its typed ``EditOp`` variant (``stored=False`` strict writer;
  ``stored=True`` tolerant reader)."""
  if not isinstance(raw, Mapping):
    raise OpUsage(f"ops[{index}] must be an object with 'op' one of {', '.join(OP_SPECS)}")
  op = dict(raw)
  kind = op.pop("op", None)
  if kind not in OP_SPECS:
    raise OpUsage(f"ops[{index}]: op must be one of {', '.join(OP_SPECS)}, not {kind!r}")
  where = f"ops[{index}] ({kind})"
  spec = OP_SPECS[kind]
  for alias, key in spec.aliases.items():
    if alias in op:
      if key in op:
        raise OpUsage(f"{where}: pass {key!r} or its alias {alias!r}, not both")
      op[key] = op.pop(alias)
  missing = [key for key in spec.required if key not in op or op[key] is None]
  unknown = sorted(set(op) - set(spec.properties))
  form_keys = {key for form in spec.forms for key in form}
  invalid_form = bool(spec.forms) and not any(
    set(op) & form_keys == set(form)
    and all(op[key] is not None or key == "value" for key in form)
    for form in spec.forms
  )
  if missing or unknown or invalid_form:
    problems = [f"missing {missing}"] if missing else []
    problems += [f"unknown {unknown}"] if unknown else []
    problems += ["pass exactly one input form"] if invalid_form else []
    raise OpUsage(f"{where}: {'; '.join(problems)}; expected keys: {_expected_keys(kind)}")
  for key in ("item_id", "after_item_id"):
    if key in op:
      op[key] = _resolve_item_ref(where, key, op[key], index, results)
  for key in ("item_id", "after_item_id", "section", "label", "unit", "formula", "formula_expr", "rationale", "reason", "text"):
    if op.get(key) is not None and not isinstance(op[key], str):
      raise OpUsage(f"{where}: {key} must be a string, not {op[key]!r}")
  for key in ("reason", "text"):
    if key in op and not op[key].strip():
      raise OpUsage(f"{where}: {key} must be non-empty text")
  if not stored and kind == "decline" and len(op["reason"].strip()) > FINDING_MESSAGE_MAX_LENGTH:
    raise OpUsage(f"{where}: reason must be at most {FINDING_MESSAGE_MAX_LENGTH} characters, not {len(op['reason'].strip())}")
  if "periods" in op:
    if not isinstance(op["periods"], Sequence) or isinstance(op["periods"], str):
      raise OpUsage(f"{where}: periods must be a list of integer periods (empty: every period), not {op['periods']!r}")
    op["periods"] = sorted({_op_int(where, "periods", period) for period in op["periods"]})
  raw_source = op.get("source")
  parsed_source = (
    parse_stored_source(raw_source, where=where)
    if stored else parse_new_source(raw_source, where=where)
  )
  if parsed_source is not None:
    op["source"] = parsed_source.to_raw()
  for key in ("period", "new_actual_year"):
    if key in op:
      op[key] = _op_int(where, key, op[key])
  if op.get("value") is not None:
    op["value"] = _op_float(where, "value", op["value"])
  if "values" in op:
    if not isinstance(op["values"], Mapping):
      raise OpUsage(f"{where}: values must be an object {{period: value}}; {_expected_keys(kind)}")
    op["values"] = {
      _op_int(where, "values period", period): None if value is None else _op_float(where, f"values[{period}]", value)
      for period, value in op["values"].items()
    }
  if op.get("inputs") is not None:
    if isinstance(op["inputs"], Sequence) and not isinstance(op["inputs"], (str, bytes)) and len(op["inputs"]) == 0:
      op["inputs"] = {}
    elif not isinstance(op["inputs"], Mapping):
      raise OpUsage(f"{where}: inputs must be an object {{period: value}}")
    else:
      op["inputs"] = {
        _op_int(where, "inputs period", period): _op_float(where, f"inputs[{period}]", value)
        for period, value in op["inputs"].items()
      }
  raw_apply_to = op.get("apply_to", "both")
  apply_to_values = OP_SPECS["set_formula"].optional["apply_to"]["anyOf"][0]["enum"]
  if raw_apply_to == "all":
    op["apply_to"] = "both"
  elif isinstance(raw_apply_to, Sequence) and not isinstance(raw_apply_to, (str, bytes)):
    if not raw_apply_to:
      raise OpUsage(f"{where}: apply_to must be one of {', '.join(apply_to_values)}, not {raw_apply_to!r}")
    op["apply_to"] = sorted({_op_int(where, "apply_to", period) for period in raw_apply_to})
  elif raw_apply_to not in apply_to_values:
    raise OpUsage(f"{where}: apply_to must be one of {', '.join(apply_to_values)}, not {raw_apply_to!r}")
  if kind == "add_line":
    try:
      parse_unit(op["unit"])
    except OpUsage as exc:
      raise OpUsage(f"{where}: {exc}") from None
    return AddLineOp(
      _raw_dict=op, section=op["section"], label=op["label"], unit=op["unit"],
      after_item_id=op.get("after_item_id"), formula=op.get("formula"),
      inputs=op.get("inputs"), rationale=op.get("rationale"), source=parsed_source,
    )
  if kind == "set_formula":
    return SetFormulaOp(
      _raw_dict=op, item_id=op["item_id"], formula_expr=op["formula_expr"],
      apply_to=op.get("apply_to", "both"), rationale=op.get("rationale"), source=parsed_source,
    )
  if kind == "set_input":
    return SetInputOp(
      _raw_dict=op, item_id=op["item_id"], period=op.get("period"), value=op.get("value"),
      values=op.get("values"), periods=op.get("periods"),
      rationale=op.get("rationale"), source=parsed_source,
    )
  if kind == "clear_input":
    return ClearInputOp(
      _raw_dict=op, item_id=op["item_id"], period=op["period"],
      rationale=op.get("rationale"), source=parsed_source,
    )
  if kind == "remove_line":
    return RemoveLineOp(
      _raw_dict=op, item_id=op["item_id"],
      rationale=op.get("rationale"), source=parsed_source,
    )
  if kind == "advance_period":
    return AdvancePeriodOp(
      _raw_dict=op, new_actual_year=op["new_actual_year"],
      rationale=op.get("rationale"), source=parsed_source,
    )
  if kind == "decline":
    return DeclineOp(
      _raw_dict=op, item_id=op["item_id"], periods=op["periods"], reason=op["reason"],
    )
  return NoteOp(_raw_dict=op, text=op["text"])


def parse_new_op(
  index: int, raw: Any, results: Sequence[Mapping[str, Any]] = (),
) -> EditOp:
  """Strict write-admission parser for one ``model_document`` op (``stored=False``)."""
  return _parse_op(index, raw, results, stored=False)


def parse_stored_op(
  index: int, raw: Any, results: Sequence[Mapping[str, Any]] = (),
) -> EditOp:
  """Tolerant replay parser for one stored ``model_document`` op (``stored=True``)."""
  return _parse_op(index, raw, results, stored=True)


def _read_op(
  index: int, raw: Any, results: Sequence[Mapping[str, Any]], *, stored: bool = False,
) -> tuple[str, dict[str, Any]]:
  """One op checked against ``OP_SPECS``: kind and its keys, ``$ref``s resolved, values typed. A
  ``stored`` op is one the log already holds: what was admitted then is not re-judged here."""
  parsed = parse_stored_op(index, raw, results) if stored else parse_new_op(index, raw, results)
  return parsed.op, parsed.to_op_dict()


def dependents_map(model: FinancialModel) -> dict[str, set[str]]:
  """Every item → the items whose formulas (any period, any offset) read it; self-reads excluded."""
  return _dependents(_cell_graph(model))


def _dependents(graph: DependencyGraph) -> dict[str, set[str]]:
  out = {node: set(readers) for node, readers in graph.adj.items()}
  for reader, refs in graph.time_edges.items():
    for ref in refs:
      out.setdefault(ref.id, set()).add(reader)
  for node, readers in out.items():
    readers.discard(node)
  return out


# --------------------------------------------------------------------------
# cells: the period-resolved dependency graph
# --------------------------------------------------------------------------


def _cell_graph(model: FinancialModel) -> DependencyGraph:
  graph = DependencyGraph()
  graph.build(model)
  return graph


class _BatchGraphs:
  """The cell graph of one ``apply_ops`` batch, and the cell → readers index over it, built when
  first asked for and again only after something they were built from changed, never once per op.

  ``DependencyGraph.build`` records the structure (sheets, rows, columns) and an item-level
  adjacency over every formula spec a line holds; ``cell_operands`` and ``_spec_for_period`` then
  read each line's specs, overrides and ``formula_periods`` live, and the recorded maps a cell
  read consults (sheet, row and column of each line) change only with the line set or the axis.
  So ``cells()`` stays current across a write to one line the model already holds — an input, an
  override, a formula authored there: that write changes only that line's specs, which
  ``written`` follows by re-reading the line's cells into the index and refreshing its recorded
  adjacency edges in place — and is rebuilt after ``structural()`` (a line added or removed,
  the axis moved), or when the document's model is a different object than the graph was
  built on.

  ``readers()`` is ``_cell_readers`` of the current model. A write to one line changes the
  reads of that line's cells alone (the specs it selects, the periods it owns), so ``written``
  re-reads that line's cells into the index; a write whose reach is wider (a fixed-cell line,
  whose anchor its readers resolve to, or an inclusion lowering) drops the index, and one
  whose lines are not known also leaves the item-level adjacency to a rebuild."""

  def __init__(self, doc: ModelSheet) -> None:
    self.doc = doc
    self._graph: DependencyGraph | None = None
    self._cells_stale = True
    self._adjacency_stale = True
    self._readers: dict[CellKey, set[CellKey]] | None = None
    self._reads: dict[CellKey, set[CellKey]] = {}
    self._cells_of: dict[str, set[CellKey]] = {}

  def cells(self) -> DependencyGraph:
    graph = self._graph
    if graph is None or self._cells_stale or graph.model is not self.doc.model:
      return self._build()
    return graph

  def adjacency(self) -> DependencyGraph:
    graph = self._graph
    if graph is None or self._cells_stale or self._adjacency_stale or graph.model is not self.doc.model:
      return self._build()
    return graph

  def readers(self) -> dict[CellKey, set[CellKey]]:
    """Every cell → the cells that read it directly, of the model as it stands."""
    graph = self.cells()
    if self._readers is None:
      self._readers, self._reads, self._cells_of = {}, {}, {}
      for item in graph._loaded_model._index.values():
        self._read_cells(graph, item)
    return self._readers

  def structural(self) -> None:
    """The line set or the axis changed: every reading of the graph is stale."""
    self._cells_stale = True

  def written(self, item_id: str | None) -> None:
    """A spec of ``item_id`` was written — an input, an override, or a formula authored on a
    line the model already holds (``None``: of lines not known, an inclusion lowering): the
    line's recorded adjacency and the readers index follow the write, the cell graph stands.
    A write whose lines are not known leaves both to a rebuild."""
    graph = self._graph
    if graph is None or self._cells_stale or graph.model is not self.doc.model:
      self._adjacency_stale = True
      return
    item = None if item_id is None else graph._loaded_model._index.get(item_id)
    if item is None:
      self._adjacency_stale = True
      self._readers = None
      return
    self._refresh_adjacency(graph, item)
    if self._readers is None:
      return
    if item.column is not None:
      self._readers = None
      return
    for cell in self._cells_of.pop(item.id, ()):
      self._unread(cell)
    self._read_cells(graph, item)

  def _refresh_adjacency(self, graph: DependencyGraph, item: LineItem) -> None:
    """``item``'s recorded item-level edges follow its specs as they stand: ``build`` records
    a line's edges from that line's specs alone, so a write to the one line drops and re-adds
    only its own edges. Left alone when the adjacency is already stale for another write,
    whose reach this one cannot repair."""
    if self._adjacency_stale:
      return
    for readers in graph.adj.values():
      readers.discard(item.id)
    graph.time_edges[item.id] = set()
    for spec in graph._iter_formula_specs(item):
      for _carrier, ref in graph._read_refs(spec):
        if ref.id not in graph.nodes:
          continue
        if ref.t == 0:
          graph.adj[ref.id].add(item.id)
        else:
          graph.time_edges[item.id].add(ref)

  def _read_cells(self, graph: DependencyGraph, item: LineItem) -> None:
    held: set[CellKey] = set()
    for period in _cell_periods(graph, item):
      operands = graph.cell_operands(item, period)
      if operands:
        cell = (item.id, period)
        held.add(cell)
        self._reads[cell] = operands
        for operand in operands:
          self.readers_of(operand).add(cell)
    self._cells_of[item.id] = held

  def _unread(self, cell: CellKey) -> None:
    for operand in self._reads.pop(cell, ()):
      self.readers_of(operand).discard(cell)

  def readers_of(self, cell: CellKey) -> set[CellKey]:
    assert self._readers is not None
    return self._readers.setdefault(cell, set())

  def _build(self) -> DependencyGraph:
    self._graph = _cell_graph(self.doc.model)
    self._cells_stale = self._adjacency_stale = False
    self._readers = None
    return self._graph


def _cell_periods(graph: DependencyGraph, item: LineItem) -> list[int]:
  """The periods ``item`` has a cell in: a fixed-cell line's anchor, else the time axis."""
  if item.column is not None:
    anchor = graph._fixed_cell_anchor_period(item)
    return [] if anchor is None else [int(anchor)]
  return [int(period) for period in graph._time_order()]


def upstream_cells(doc: ModelSheet, item_id: str, period: int) -> set[CellKey]:
  """Every cell ``(item_id, period)`` reads, at any depth, period by period
  (``DependencyGraph.cell_operands``: ``[-1]`` reads the prior period's cell, ``[last]`` the
  last projection period's); the cell itself excluded. An item-level walk (``dependents_map``) reaches every period of every line
  that feeds any period; this reaches only the cells that feed this one."""
  model = doc.model
  if not model._index:
    model.build_index()
  model.get_item(item_id)
  start = (item_id, int(period))
  seen = _upstream(_cell_graph(model), model, start)
  seen.discard(start)
  return seen


def _upstream(graph: DependencyGraph, model: FinancialModel, start: CellKey) -> set[CellKey]:
  """Every cell ``start`` reads at any depth through ``graph``'s live spec selection; ``start``
  itself only when it reads itself through a cycle."""
  seen: set[CellKey] = set()
  stack = [start]
  while stack:
    reader, reader_period = stack.pop()
    item = model._index.get(reader)
    if item is None:
      continue
    for operand in graph.cell_operands(item, reader_period):
      if operand not in seen:
        seen.add(operand)
        stack.append(operand)
  return seen


def _cell_readers(graph: DependencyGraph, model: FinancialModel) -> dict[CellKey, set[CellKey]]:
  """Every cell → the cells that read it directly (``DependencyGraph.cell_operands`` inverted)."""
  readers: dict[CellKey, set[CellKey]] = {}
  for item in model._index.values():
    for period in _cell_periods(graph, item):
      for operand in graph.cell_operands(item, period):
        readers.setdefault(operand, set()).add((item.id, period))
  return readers


def _reachable(readers: Mapping[CellKey, set[CellKey]], cells: Iterable[CellKey]) -> set[CellKey]:
  seen: set[CellKey] = set()
  stack = [(str(item_id), int(period)) for item_id, period in cells]
  while stack:
    for reader in readers.get(stack.pop(), ()):
      if reader not in seen:
        seen.add(reader)
        stack.append(reader)
  return seen


def downstream_cells(doc: ModelSheet, cells: Iterable[CellKey]) -> set[CellKey]:
  """Every cell that reads any of ``cells``, at any depth (the inverse of ``upstream_cells``);
  ``cells`` themselves excluded unless one reads another."""
  return _downstream(_BatchGraphs(doc), cells)


def _downstream(graphs: _BatchGraphs, cells: Iterable[CellKey]) -> set[CellKey]:
  return _reachable(graphs.readers(), cells)


#: Units whose values are pure numbers: their ``model_scale`` never scales a product.
_SCALE_FREE_UNITS = frozenset({Unit.percentage, Unit.ratio, Unit.multiple, Unit.days, Unit.date})
_SCALE_EXPONENTS: dict[str, int] = {"units": 0, "millions": 6}
_SUM_OPS = frozenset({"+", "-", "SUM", "AVG", "MAX", "MIN"})
#: A literal number's scale: it takes its sum's scale and adds nothing to a product's.
_FREE = "free"


def _as_ref(node: Any) -> LineItemRef | None:
  return line_item_ref_from_obj(node) if isinstance(node, (LineItemRef, Mapping)) else None


def _children(node: Mapping[str, Any]) -> list[Any]:
  return node.get("args") or [node.get(key) for key in ("left", "right", "arg") if key in node]


def _scale_name(exponent: int) -> str:
  names = {value: key for key, value in _SCALE_EXPONENTS.items()}
  return names.get(exponent, f"10^{exponent}")


def _line_exponent(item: LineItem) -> int:
  """A line's scale as a power of ten; a pure-number unit (``_SCALE_FREE_UNITS``) is 0."""
  return 0 if item.unit in _SCALE_FREE_UNITS else _SCALE_EXPONENTS.get(item.model_scale, 0)


def _sum_terms(node: Any) -> list[Any]:
  """A sum's terms, nested sums and negations flattened: ``a - (b + -c)`` → ``a, b, c``."""
  if _as_ref(node) is None and isinstance(node, Mapping):
    if node.get("op") == "NEG":
      return _sum_terms(node.get("arg"))
    if node.get("op") in _SUM_OPS:
      return [term for child in _children(node) for term in _sum_terms(child)]
  return [node]


def _expr_scale(node: Any, model: FinancialModel, warnings: list[str]) -> int | str | None:
  """``node``'s scale as a power of ten (``_FREE`` for a literal, ``None`` when unknown); a sum
  whose terms mix scales appends a warning to ``warnings``. A product adds its factors'
  scales, a quotient subtracts them, so ``count × USD millions`` is millions and
  ``USD millions / share_count millions`` is units."""
  ref = _as_ref(node)
  if ref is not None:
    item = model._index.get(ref.id)
    return None if item is None else _line_exponent(item)
  if isinstance(node, (int, float)) and not isinstance(node, bool):
    return _FREE
  if not isinstance(node, Mapping):
    return None
  op = node.get("op")
  if op in _SUM_OPS or op == "NEG":
    by_scale: dict[int, set[str]] = {}
    unknown = False
    for term in _sum_terms(node):
      scale = _expr_scale(term, model, warnings)
      if scale is None:
        unknown = True
      elif isinstance(scale, int):
        by_scale.setdefault(scale, set()).add(_expr_text(term)[0])
    if len(by_scale) > 1:
      warnings.append("mixed model_scale in one sum: " + "; ".join(
        f"{name}: {sorted(terms)}"
        for name, terms in sorted((_scale_name(scale), terms) for scale, terms in by_scale.items())
      ))
      return None
    if unknown:
      return None
    return next(iter(by_scale)) if by_scale else _FREE
  scales = [_expr_scale(child, model, warnings) for child in _children(node)]
  if None in scales or op not in {"*", "/", "^"}:
    return None
  if all(scale == _FREE for scale in scales):
    return _FREE
  exponents = [scale if isinstance(scale, int) else 0 for scale in scales]
  if op == "*":
    return sum(exponents)
  if op == "/":
    return exponents[0] - exponents[1]
  power = node.get("right")
  if exponents[0] == 0:
    return 0
  if isinstance(power, int) and not isinstance(power, bool):
    return exponents[0] * power
  return None


def scale_warnings(spec: FormulaSpec, model: FinancialModel, item_id: str) -> list[str]:
  """A warning per sum (``+ - SUM AVG MAX MIN``, at any depth) whose terms mix scales, and when
  the formula's result scale differs from the line's: scales as powers of ten, a product adding
  its factors' (``_expr_scale``), a pure-number unit (``_SCALE_FREE_UNITS``) contributing none.
  Warnings only: nothing is refused."""
  expr = formula_expr(spec)
  if expr is None:
    return []
  warnings: list[str] = []
  result = _expr_scale(expr, model, warnings)
  target = model._index.get(item_id)
  if target is not None and isinstance(result, int) and result != _line_exponent(target):
    warnings.append(
      f"{item_id} is model_scale {_scale_name(_line_exponent(target))} but its formula gives "
      f"{_scale_name(result)}: {_expr_text(expr)[0]}"
    )
  return warnings


def formula_report(spec: FormulaSpec, model: FinancialModel, item_id: str) -> dict[str, Any]:
  """How a formula was stored: ``typed`` (our graph evaluates it) or ``raw`` (only the native
  engine does), and any mixed-scale warnings."""
  raw = spec.type is FormulaType.raw
  return {
    "stored_as": "raw" if raw else "typed",
    "ours_supported": not raw,
    "warnings": scale_warnings(spec, model, item_id),
  }


def _agent_note(op: Mapping[str, Any], edit_index: int) -> dict[str, Any]:
  """An agent op's note: its ``source`` (text → ``{kind: agent, text}``; an object keeps its own
  ``kind``, default ``agent``; absent → ``{kind: agent}``), ``rationale`` when given, and the
  index of the edit that records it in ``doc.edits``."""
  given = op.get("source")
  if given is None:
    source: dict[str, Any] = {"kind": "agent"}
  elif isinstance(given, str):
    source = {"kind": "agent", "text": given}
  else:
    source = {"kind": "agent", **dict(given)}
  note: dict[str, Any] = {"source": source, "edit_index": edit_index}
  if op.get("rationale") is not None:
    note["rationale"] = op["rationale"]
  return note


def _formula_spec(graph: DependencyGraph, item: LineItem, period: int, recompute: set[str]) -> FormulaSpec | None:
  """The formula that produces ``(item, period)``'s value, ``None`` when none does — the recalc's
  rule (``_seed_inputs`` over ``_formula_owned_snapshots``): the period selects a non-constant spec
  and it is an override, or the row is recomputed (``formula_owned_recompute_ids``), or no stored
  value stands in for it."""
  spec = graph._spec_for_period(item, int(period))
  if spec is None or spec.type is FormulaType.constant:
    return None
  if (item.overrides or {}).get(int(period)) is spec or item.id in recompute:
    return spec
  stored = item.values.values.get(int(period)) if item.values is not None else None
  return spec if stored is None else None


def _record_superseded(doc: ModelSheet, rows: list[dict[str, Any]]) -> list[list[Any]]:
  """The columns an edit dropped (``{item_id, period, column}``) onto that edit (the last one)'s
  ``columns_superseded``; returns their cells."""
  if rows:
    doc.edits[-1].setdefault("columns_superseded", []).extend(rows)
  return [[row["item_id"], row["period"]] for row in rows]


def _formula_columns(
  doc: ModelSheet,
  item_id: str,
  meta: Mapping[str, Any],
  apply_to: ApplyTo,
  *,
  is_agent: bool = True,
  graphs: _BatchGraphs,
) -> tuple[set[CellKey], list[list[Any]]]:
  """Stamp the source on every cell the authored formula now produces (``_formula_spec``
  selects the line's ``apply_to`` formula, not an override, and no stored value stands in for it).
  Each replaced column rides on the edit; return the written and replaced cells."""
  graph = graphs.cells()
  recompute = formula_owned_recompute_ids(doc.model)
  item = doc.model.get_item(item_id)
  authored = [
    spec for spec, side in ((item.historical, "historical"), (item.projected, "projected"))
    if spec is not None and apply_to in (side, "both")
  ]
  written: set[CellKey] = set()
  replaced: list[dict[str, Any]] = []
  for period in _cell_periods(graph, item):
    selected = _formula_spec(graph, item, period, recompute)
    if selected is None or not any(selected is spec for spec in authored):
      continue
    prior = doc.semantic.get((item_id, period))
    if prior is not None:
      replaced.append({"item_id": item_id, "period": period, "column": prior})
    doc.semantic[(item_id, period)] = (
      stamp_cell_source({
        "cell": "formula", "origin": "manual_input", "formula": formula_text(selected),
        "operands": sorted({ref.id for ref in graph._extract_refs(selected)}), **dict(meta),
      })
      if is_agent else stamp_cell_source({**formula_column(selected), **dict(meta)})
    )
    written.add((item_id, period))
  return written, _record_superseded(doc, replaced)


def _input_cell(doc: ModelSheet, item_id: str, period: int) -> CellKey:
  """The cell an input at ``period`` lands in: a fixed-cell line's anchor, else ``period``."""
  item = doc.model.get_item(item_id)
  anchor = _fixed_cell_anchor_period(doc.model, item) if item.column is not None else None
  return (item_id, int(anchor if anchor is not None else period))


def _supersede_input_columns(doc: ModelSheet, edited: set[CellKey], graphs: _BatchGraphs) -> list[list[Any]]:
  """Drop the input column (``cell: input`` — a judgment row, a C fact, a construction input) of
  every cell downstream of ``edited`` whose value a formula produces (``_formula_spec``): its
  value is the formula's over the edit, so it carries the formula column (``origin: derived``,
  its operands) from now on, and the judgment column stays only on the input cell the judgment
  sets. A stored value the recalc serves keeps its column, and so does a column that already is the
  formula's (``origin: derived``: the construction helper's). The dropped columns ride on the edit
  (``columns_superseded``); returns their cells."""
  graph = graphs.cells()
  recompute = formula_owned_recompute_ids(doc.model)
  superseded: list[dict[str, Any]] = []
  for cell in sorted(_downstream(graphs, edited) - edited):
    column = doc.semantic.get(cell)
    if column is None or column.get("cell") != "input" or column.get("origin") == "derived":
      continue
    if _formula_spec(graph, doc.model.get_item(cell[0]), cell[1], recompute) is None:
      continue
    superseded.append({"item_id": cell[0], "period": cell[1], "column": doc.semantic.pop(cell)})
  return _record_superseded(doc, superseded)


def apply_ops(doc: ModelSheet, ops: Sequence[Mapping[str, Any]], *, stored: bool = False) -> list[dict[str, Any]]:
  """A batch of ``{op: add_line|set_formula|set_input|clear_input|remove_line|advance_period|decline|note, ...}``,
  in order (``advance_period``: ``advance_period``, the new actual year's provisional cells listed;
  ``decline`` / ``note``: ``record_declination`` / ``record_note``, recorded on ``doc.edits``, no cell written).

  Keys per op are ``OP_SPECS``; a missing or unknown key is an ``OpUsage`` naming them. A later
  op names an earlier ``add_line``'s line by its minted id (in any field or formula text) or
  by ``{"$ref": <op index>}`` in ``item_id`` / ``after_item_id``. The first invalid
  operation raises ``OpUsage`` carrying its zero-based ``op_index`` and original
  message. The batch is one ``transaction``: an op that raises leaves ``doc.model`` and
  ``doc.semantic`` as they were before the batch, and ``doc.edits`` holds the records appended up
  to and including the refused op. The scenario outputs refresh (``refresh_scenario_outputs``)
  runs on the committed document after it.

  Every op is the agent's: its edit carries ``agent`` (``_agent_note``: ``source``,
  ``rationale``, ``edit_index``) and each cell it writes the column ``{origin: manual_input,
  **agent}`` (``origin: cited_input`` for a ``set_input`` or an ``add_line`` input whose ``source`` is a cited object) —
  ``cell: input`` for a value, ``cleared`` for a cleared one (at a fixed-cell line's
  anchor), ``formula`` (with ``formula`` and ``operands``) where its formula now produces the value
  (``_formula_spec``, the recalc's rule). A formula-owned cell downstream of what it wrote loses a
  stale input column (``_supersede_input_columns``); every column an op drops or replaces rides on
  its edit and result as ``columns_superseded``. A ``set_input`` and a ``clear_input`` record the
  cell's prior column, override and value on their edit (``replaced``); a ``set_input`` that puts a
  cell the agent changed back to its constructed value restores the constructed column, override
  note and value (edit and result row: ``restored_baseline``). A ``set_formula`` clears the agent
  inputs (constant override and value) on the cells of its ``apply_to`` side(s) so its formula is
  what they evaluate, and records each as ``replaced`` on its edit and result row; overrides off
  that side and bench-written ones stay.

  ``stored`` marks ops the log already holds (``live_sheet.replay_agent_ops``): a source shape the
  door refuses for a new op (a partial ``judgment_row`` ref) replays as the manual input it was
  admitted as, never ends the replay.

  The refresh record of the previous batch (``release_refresh_record``) is replaced, not kept beside the
  new one: an op's ``edit_index`` is its position in the log a single replay of every op leaves.
  """
  from .target_build import refresh_scenario_outputs, release_refresh_record

  with transaction(doc):
    if ops:
      release_refresh_record(doc)
    results = _run_ops(doc, ops, stored=stored)
  if results:
    refresh_scenario_outputs(doc)
  return results


def _resolve_apply_to(where: str, raw_apply_to: Any, time: Any) -> ApplyTo:
  if raw_apply_to == "projected":
    return "projected"
  if raw_apply_to == "historical":
    return "historical"
  if raw_apply_to == "both":
    return "both"
  periods = set(raw_apply_to) if isinstance(raw_apply_to, Sequence) else set()
  proj = {int(period) for period in (time.projection_periods or time.projection_years or ())}
  hist = {int(period) for period in (time.historical_periods or time.historical_years or ())}
  if periods and periods == proj:
    return "projected"
  if periods and periods == hist:
    return "historical"
  if periods and periods == (proj | hist):
    return "both"
  if periods:
    raise OpUsage(
      f"{where}: set_formula applies to a whole side ('projected' {sorted(proj)}, 'historical' {sorted(hist)}, "
      f"or 'both'), not a subset of periods {sorted(periods)}; use set_input for single-period values"
    )
  accepted = OP_SPECS["set_formula"].optional["apply_to"]["anyOf"][0]["enum"]
  raise OpUsage(f"{where}: apply_to must be one of {', '.join(accepted)}, not {raw_apply_to!r}")


def _op_error_entry(index: int, raw: Any, exc: OpUsage | EditRefused) -> dict[str, Any]:
  item_id = raw.get("item_id") if isinstance(raw, Mapping) else None
  entry: dict[str, Any] = {
    "op_index": index,
    "item_id": item_id if isinstance(item_id, str) else None,
    "reason": exc.result.reason if isinstance(exc, EditRefused) else exc.reason,
    "message": str(exc),
  }
  if isinstance(exc, OpUsage) and exc.next_call is not None:
    entry["next_call"] = dict(exc.next_call)
  return entry


def _refuse_historical_statement_analyst_note(
  doc: ModelSheet,
  graphs: _BatchGraphs,
  index: int,
  cell: CellKey,
  raw_source: Any,
) -> None:
  if not isinstance(raw_source, Mapping) or raw_source.get("kind") not in ANALYST_NOTE_KINDS:
    return
  if not is_historical_statement_cell(doc.model, cell[0], cell[1]):
    return
  from schema.root_causes import segment_revenue_seed_subject, stage_door

  seed_subject = segment_revenue_seed_subject(graphs.cells(), doc.model, cell[0], cell[1])
  placeholder_hint = ""
  if seed_subject is not None:
    door = stage_door("historical_foundation", seed_subject)
    if str(seed_subject).startswith("tpl.fm."):
      from schema.stage_registry import STAGE_REGISTRY

      read_tool = STAGE_REGISTRY["historical_foundation"].producer
      placeholder_hint = (
        f" (fill historical_driver_seeds[].driver_key with the '<segment>.<driver>' key "
        f"from {read_tool}(research_file_id=<research_file_id>))"
      )
  else:
    item_obj = doc.model.get_item(cell[0]) if cell[0] in doc.model._index else None
    role = None if item_obj is None else item_obj.required_role
    door = stage_door("correspondence", cell[0], role=role)
  next_call = {"tool_name": door["tool"], "arguments": dict(door["arguments"])}
  raise OpUsage(
    f"ops[{index}] (set_input): source.kind {raw_source.get('kind')!r} is for off-sheet judgment inputs, "
    f"not historical statement line {cell[0]!r} in FY{cell[1]}; bind historical statement lines through "
    f"{door['tool']}{placeholder_hint} or pass a cited filing_line / kpi source.",
    op_index=index,
    reason="invalid_source_kind_for_statement_line",
    next_call=next_call,
  )

def _run_single_op(
  doc: ModelSheet, graphs: _BatchGraphs, index: int, raw: Any, results: list[dict[str, Any]], *, stored: bool = False,
) -> None:
  try:
    kind, op = _read_op(index, raw, results, stored=stored)
    agent = _agent_note(op, len(doc.edits))
    if kind == "add_line":
      item_id = add_line(
        doc,
        section=op["section"],
        label=op["label"],
        unit=op["unit"],
        after_item_id=op.get("after_item_id"),
        formula=op.get("formula"),
        inputs=op.get("inputs"),
        agent=agent,
      )
      graphs.structural()
      formula_cells: set[CellKey] = set()
      if op.get("formula") is not None:
        formula_cells, _ = _formula_columns(doc, item_id, agent, "both", graphs=graphs)
      for period in op.get("inputs") or {}:
        if _input_cell(doc, item_id, period) not in formula_cells:
          judgment = analyst_note_judgment(doc, agent["source"], item_id, period, agent.get("edit_index"))
          doc.semantic[_input_cell(doc, item_id, period)] = stamp_cell_source({
            "cell": "input", "origin": _input_origin(agent["source"]), **agent,
            **({} if judgment is None else {"source": judgment}),
          })
      item = doc.model.get_item(item_id)
      row: dict[str, Any] = {
        "op": kind, "item_id": item_id, "unit": op["unit"],
        "stored_unit": None if item.unit is None else item.unit.value, "model_scale": item.model_scale,
      }
      spec = item.projected
      if op.get("formula") is not None and spec is not None:
        row.update(formula_report(spec, doc.model, item_id))
      results.append(row)
    elif kind == "set_formula":
      item_id = op["item_id"]
      apply_to = _resolve_apply_to(f"ops[{index}] (set_formula)", op.get("apply_to", "both"), doc.model.time_structure)
      before = _dependents(graphs.adjacency())
      spec, written, replaced, cleared = _apply_set_formula_with_columns(
        doc, item_id, op["formula_expr"], apply_to=apply_to, agent=agent, graphs=graphs,
      )
      after = _dependents(graphs.adjacency())
      results.append({
        "op": kind,
        "item_id": item_id,
        "type": spec.type.value,
        **formula_report(spec, doc.model, item_id),
        "dependents_removed": sorted(
          operand for operand, readers in before.items() if item_id in readers and item_id not in after.get(operand, ())
        ),
        "orphaned_inputs": sorted(operand for operand, readers in before.items() if readers and not after.get(operand)),
        "columns_superseded": [*replaced, *_supersede_input_columns(doc, written, graphs)],
        **({"replaced": cleared} if cleared else {}),
      })
    elif kind in ("set_input", "clear_input"):
      if "values" in op:
        values = op["values"]
      elif "periods" in op:
        values = {period: op["value"] for period in op["periods"]}
      else:
        values = {op["period"]: op.get("value")}
      restored = False
      superseded: list[list[Any]] = []
      for period, value in values.items():
        cell = _input_cell(doc, op["item_id"], period)
        if kind == "set_input" and not stored and value is not None:
          _refuse_historical_statement_analyst_note(doc, graphs, index, cell, op.get("source"))
        marker = (doc.model.get_item(cell[0]).overrides or {}).get(cell[1])
        identities = _inclusion_role_items(doc) if marker is not None and marker.subtype == "not_separately_presented" else {}
        included_role = (
          marker.params.get("_not_separately_presented", {}).get("required_role")
          or (doc.semantic.get(cell) or {}).get("required_role")
          or doc.model.get_item(cell[0]).required_role
        ) if marker is not None and marker.subtype == "not_separately_presented" else None
        # Each cell's lineage points at the edit that actually records its displaced baseline.
        cell_agent = _agent_note(op, len(doc.edits))
        if kind == "set_input":
          set_input(doc, op["item_id"], period, value, agent=cell_agent)
        else:
          clear_input(doc, op["item_id"], period, agent=cell_agent)
        restored = bool(doc.edits[-1].get("restored_baseline")) or restored
        if included_role is not None:
          reconcile_inclusions(
            doc, role_items=identities, cleared_role_periods=((included_role, cell[1]),),
          )
        graphs.written(cell[0] if included_role is None else None)
        superseded.extend(_supersede_input_columns(doc, {cell}, graphs))
      row = {"op": kind, "item_id": op["item_id"]}
      row.update({"period": op["period"]} if "period" in op else {"periods": sorted(values)})
      if kind == "clear_input":
        row["cleared_by"] = "agent"
      if restored:
        row["restored_baseline"] = True
      results.append({**row, "columns_superseded": superseded})
    elif kind == "decline":
      record_declination(doc, op["item_id"], op["periods"], op["reason"], agent=agent)
      results.append({"op": kind, "item_id": op["item_id"], "periods": op["periods"]})
    elif kind == "note":
      record_note(doc, op["text"], agent=agent)
      results.append({"op": kind})
    elif kind == "advance_period":
      advanced = advance_period(doc, op["new_actual_year"], agent=agent, graphs=graphs)
      graphs.structural()
      results.append({
        "op": kind,
        **{key: advanced[key] for key in ("axis_before", "axis_after", "new_actual_year", "new_projection_year")},
        "provisional_projection": [[row["item_id"], row["period"]] for row in advanced["provisional_projection"]],
        "formula_periods_extended": len(advanced["formula_periods_extended"]),
        "fixed_cells_shifted": len(advanced["fixed_cells_shifted"]),
      })
    else:
      remove_line(doc, op["item_id"], agent=agent)
      graphs.structural()
      results.append({"op": kind, "item_id": op["item_id"]})
  except (OpUsage, EditRefused) as exc:
    exc.op_index = index
    raise
  except KeyError as exc:
    missing = exc.args[0] if exc.args and isinstance(exc.args[0], str) else str(exc)
    near = closest_items(doc.model, missing) if isinstance(missing, str) else []
    hint = f"; did you mean: {', '.join(near)}" if near else ""
    op_kind = raw.get("op") if isinstance(raw, Mapping) else None
    prefix = f"ops[{index}] ({op_kind})" if isinstance(op_kind, str) and op_kind in OP_SPECS else f"ops[{index}]"
    raise OpUsage(f"{prefix}: {missing!r} is not an item id of the document{hint}", op_index=index) from exc
  except ValueError as exc:
    raise OpUsage(str(exc), op_index=index) from exc


def _run_ops(doc: ModelSheet, ops: Sequence[Mapping[str, Any]], *, stored: bool = False) -> list[dict[str, Any]]:
  """``apply_ops``'s ops, in order, on one cell graph that is rebuilt only where an op changed
  what it was built from (``_BatchGraphs``)."""
  graphs = _BatchGraphs(doc)
  results: list[dict[str, Any]] = []
  for index, raw in enumerate(ops):
    try:
      _run_single_op(doc, graphs, index, raw, results, stored=stored)
    except (OpUsage, EditRefused) as first_exc:
      op_errors = [_op_error_entry(index, raw, first_exc)]
      valid_indices = list(range(index))
      if not stored and index + 1 < len(ops):
        probe = fork(doc)
        probe_graphs = _BatchGraphs(probe)
        probe_results = list(results)
        for rest_index in range(index + 1, len(ops)):
          while len(probe_results) < rest_index:
            probe_results.append({"op": "noop"})
          try:
            _run_single_op(probe, probe_graphs, rest_index, ops[rest_index], probe_results, stored=stored)
            valid_indices.append(rest_index)
          except (OpUsage, EditRefused) as rest_exc:
            op_errors.append(_op_error_entry(rest_index, ops[rest_index], rest_exc))
      first_exc.op_errors = op_errors
      first_exc.valid_op_indices = valid_indices
      raise
  return results


# --------------------------------------------------------------------------
# formula text
# --------------------------------------------------------------------------


_TOKEN_RE = re.compile(
  r"""
   (?P<ws>\s+)
  |(?P<num>(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?)
  |(?P<str>"(?:[^"]|"")*")
  |(?P<name>[A-Za-z_][A-Za-z0-9_.]*)
  |(?P<offset>\[\s*(?:last\s*(?:-\s*\d+)?|[-+]?\d+)\s*\])
  |(?P<cmp><>|<=|>=|=|<|>)
  |(?P<op>[-+*/^&%])
  |(?P<lp>\()
  |(?P<rp>\))
  |(?P<comma>,)
  """,
  re.X,
)
_GRAMMAR_FUNCTIONS = {"SUM": "SUM", "AVG": "AVG", "AVERAGE": "AVG", "MAX": "MAX", "MIN": "MIN"}
_RAW_CONSTANTS = frozenset({"TRUE", "FALSE"})


@dataclass(frozen=True)
class _Token:
  kind: str
  text: str
  start: int
  end: int


def _tokens(text: str) -> list[_Token]:
  out: list[_Token] = []
  position = 0
  while position < len(text):
    match = _TOKEN_RE.match(text, position)
    if match is None:
      hint = (
        " (use relative offset item_id[-1] or anchor item_id[last] / item_id[last-1], not [FY...] or @year)"
        if text[position] in {"[", "@"}
        else ""
      )
      raise ValueError(f"formula: unexpected {text[position]!r} at {position} in {text!r}{hint}")
    kind = match.lastgroup or ""
    if kind != "ws":
      out.append(_Token(kind, match.group(), match.start(), match.end()))
    position = match.end()
  return out


def _is_call(tokens: Sequence[_Token], index: int) -> bool:
  return index + 1 < len(tokens) and tokens[index + 1].kind == "lp"


def _offset(token: _Token) -> tuple[int, PeriodAnchor]:
  """``[N]`` → ``(N, "first")``; ``[last]`` / ``[last-N]`` → ``(0 | -N, "last")``."""
  inner = token.text.strip("[] ").replace(" ", "")
  if inner.startswith("last"):
    return int(inner.removeprefix("last") or 0), "last"
  return int(inner), "first"


def _ref_at(tokens: Sequence[_Token], index: int) -> LineItemRef:
  """The ``LineItemRef`` the name at ``index`` reads, with its following offset if any."""
  has_offset = index + 1 < len(tokens) and tokens[index + 1].kind == "offset"
  offset: tuple[int, PeriodAnchor] = _offset(tokens[index + 1]) if has_offset else (0, "first")
  t, anchor = offset
  return LineItemRef(id=tokens[index].text, t=t, period_anchor=anchor)


def _refs(tokens: Sequence[_Token], model: FinancialModel, self_id: str | None = None) -> list[tuple[int, LineItemRef]]:
  """Every item reference ``(token index, ref)``; an unknown bare name (other than ``self_id``)
  is an error, and so is an item id followed by ``(`` (a spreadsheet reads a call to an
  unknown function there, ``#NAME?``; the likely intent is a multiply)."""
  refs: list[tuple[int, LineItemRef]] = []
  for index, token in enumerate(tokens):
    if token.kind == "offset":
      if index == 0 or tokens[index - 1].kind != "name" or _is_call(tokens, index - 1):
        raise ValueError(f"formula: offset {token.text} must follow an item id")
      continue
    if token.kind != "name" or token.text.upper() in _RAW_CONSTANTS:
      continue
    is_item = token.text in model._index or token.text == self_id
    if _is_call(tokens, index):
      if is_item and token.text.upper() not in _GRAMMAR_FUNCTIONS:
        raise ValueError(
          f"formula: item id {token.text!r} is followed by '(' at {tokens[index + 1].start}, which a "
          f"spreadsheet reads as a call to an unknown function (#NAME?); for a multiply write "
          f"'{token.text} * (…)'"
        )
      continue
    if not is_item:
      near = closest_items(model, token.text)
      hint = f"; did you mean: {', '.join(near)}" if near else ""
      raise ValueError(f"formula: {token.text!r} is not an item id of the document{hint}")
    refs.append((index, _ref_at(tokens, index)))
  return refs


def _outside_grammar(tokens: Sequence[_Token]) -> str | None:
  for index, token in enumerate(tokens):
    if token.kind in {"str", "cmp"} or (token.kind == "op" and token.text in {"&", "%"}):
      return token.text
    if token.kind == "name" and _is_call(tokens, index) and token.text.upper() not in _GRAMMAR_FUNCTIONS:
      return token.text.upper()
    if token.kind == "name" and token.text.upper() in _RAW_CONSTANTS and not _is_call(tokens, index):
      return token.text.upper()
  return None


class _Parser:
  """Recursive descent: expr := term (('+'|'-') term)*; term := power (('*'|'/') power)*;
  power := unary ('^' unary)*; unary := ('-'|'+') unary | primary (spreadsheet precedence)."""

  def __init__(self, tokens: Sequence[_Token], refs: Mapping[int, LineItemRef], text: str) -> None:
    self._tokens = tokens
    self._refs = refs
    self._text = text
    self._index = 0

  def parse(self) -> Any:
    node = self._expr()
    if self._index != len(self._tokens):
      self._fail("unexpected trailing input")
    return node

  def _fail(self, message: str) -> NoReturn:
    at = self._tokens[self._index].start if self._index < len(self._tokens) else len(self._text)
    raise ValueError(f"formula: {message} at {at} in {self._text!r}")

  def _peek(self, kind: str, text: str | None = None) -> bool:
    if self._index >= len(self._tokens):
      return False
    token = self._tokens[self._index]
    return token.kind == kind and (text is None or token.text == text)

  def _take(self) -> _Token:
    token = self._tokens[self._index]
    self._index += 1
    return token

  def _expr(self) -> Any:
    node = self._term()
    chain: dict[str, Any] | None = None
    while self._peek("op", "+") or self._peek("op", "-"):
      operator = self._take().text
      right = self._term()
      if operator == "+":
        if chain is not None and chain is node:
          chain["args"].append(right)
        else:
          chain = {"op": "+", "args": [node, right]}
          node = chain
      else:
        node = {"op": "-", "left": node, "right": right}
    return node

  def _term(self) -> Any:
    node = self._power()
    chain: dict[str, Any] | None = None
    while self._peek("op", "*") or self._peek("op", "/"):
      operator = self._take().text
      right = self._power()
      if operator == "*":
        if chain is not None and chain is node:
          chain["args"].append(right)
        else:
          chain = {"op": "*", "args": [node, right]}
          node = chain
      else:
        node = {"op": "/", "left": node, "right": right}
    return node

  def _power(self) -> Any:
    node = self._unary()
    while self._peek("op", "^"):
      self._take()
      node = {"op": "^", "left": node, "right": self._unary()}
    return node

  def _unary(self) -> Any:
    if self._peek("op", "-"):
      self._take()
      inner = self._unary()
      return -inner if isinstance(inner, (int, float)) else {"op": "NEG", "arg": inner}
    if self._peek("op", "+"):
      self._take()
      return self._unary()
    return self._primary()

  def _primary(self) -> Any:
    if self._index >= len(self._tokens):
      return self._fail("expression ends early")
    token = self._tokens[self._index]
    if token.kind == "num":
      self._take()
      number = float(token.text)
      return int(number) if number.is_integer() and abs(number) < 2**53 else number
    if token.kind == "lp":
      self._take()
      node = self._expr()
      if not self._peek("rp"):
        return self._fail("missing ')'")
      self._take()
      return node
    if token.kind == "name" and self._index + 1 < len(self._tokens) and self._tokens[self._index + 1].kind == "lp":
      function = _GRAMMAR_FUNCTIONS[token.text.upper()]
      self._index += 2
      args = [self._expr()]
      while self._peek("comma"):
        self._take()
        args.append(self._expr())
      if not self._peek("rp"):
        return self._fail(f"missing ')' after {function} arguments")
      self._take()
      return {"op": function, "args": args}
    if token.kind == "name":
      ref = self._refs[self._index]
      self._take()
      if self._peek("offset"):
        self._take()
      return ref
    return self._fail(f"unexpected {token.text!r}")


def parse_formula(text: str, model: FinancialModel, *, self_id: str | None = None) -> FormulaSpec:
  """Infix text over item ids → ``arithmetic`` ``{expr}`` or, outside the grammar, ``raw``.

  ``self_id`` names the line the formula is for when it is not in ``model`` yet (``add_line``),
  so the formula may read it. A raw spec stores ``semantic_formula`` (the text, ids intact) and
  ``refs`` (its operands, so the graph keeps the dependency edges); ``render`` resolves the ids
  to cells per period.
  """
  if not model._index:
    model.build_index()
  source = text.strip().removeprefix("=").strip()
  tokens = _tokens(source)
  if not tokens:
    raise ValueError("formula: empty")
  refs = _refs(tokens, model, self_id)
  if _outside_grammar(tokens) is not None:
    depth = 0
    for token in tokens:
      depth += 1 if token.kind == "lp" else -1 if token.kind == "rp" else 0
      if depth < 0:
        raise ValueError(f"formula: unbalanced ')' in {source!r}")
    if depth:
      raise ValueError(f"formula: unbalanced '(' in {source!r}")
    unique = list(dict.fromkeys(ref for _index, ref in refs))
    return FormulaSpec(type=FormulaType.raw, params={"semantic_formula": source, "refs": unique})
  expr = _Parser(tokens, dict(refs), source).parse()
  return FormulaSpec(type=FormulaType.arithmetic, params={"expr": expr})


# --------------------------------------------------------------------------
# formula text out (the inverse of parse_formula)
# --------------------------------------------------------------------------


_ADD, _MUL, _POW, _UNARY, _ATOM = range(1, 6)


def _number_text(value: float) -> tuple[str, int]:
  number = float(value)
  text = str(int(number)) if number.is_integer() and abs(number) < 2**53 else repr(number)
  return text, _UNARY if number < 0 else _ATOM


def _ref_text(ref: LineItemRef) -> str:
  """``id``, ``id[t]``, or for a last-period anchor ``id[last]`` / ``id[last-N]`` (``_offset``'s forms)."""
  if ref.period_anchor == "last":
    return f"{ref.id}[last]" if ref.t == 0 else f"{ref.id}[last{ref.t:+d}]"
  return ref.id if ref.t == 0 else f"{ref.id}[{ref.t:+d}]"


def _wrap(node: Any, precedence: int, *, first: bool) -> str:
  text, inner = _expr_text(node)
  return f"({text})" if inner < precedence or (not first and inner == precedence) else text


def _expr_text(node: Any) -> tuple[str, int]:
  """One expression node → (infix text, its precedence); parsing the text gives the node back."""
  ref = line_item_ref_from_obj(node) if isinstance(node, (LineItemRef, Mapping)) else None
  if ref is not None:
    return _ref_text(ref), _ATOM
  if isinstance(node, (int, float)) and not isinstance(node, bool):
    return _number_text(node)
  if isinstance(node, (list, tuple)):
    return "[" + ", ".join(_expr_text(child)[0] for child in node) + "]", _ATOM
  if not isinstance(node, Mapping) or node.get("op") is None:
    return json.dumps(node, default=str), _ATOM
  op = node["op"]
  if op in {"+", "*"}:
    precedence = _ADD if op == "+" else _MUL
    args = node.get("args") or []
    return f" {op} ".join(_wrap(arg, precedence, first=index == 0) for index, arg in enumerate(args)), precedence
  if op in {"-", "/"}:
    precedence = _ADD if op == "-" else _MUL
    left = _wrap(node.get("left"), precedence, first=True)
    return f"{left} {op} {_wrap(node.get('right'), precedence, first=False)}", precedence
  if op == "^":
    return f"{_wrap(node.get('left'), _POW, first=True)}^{_wrap(node.get('right'), _UNARY, first=True)}", _POW
  if op == "NEG":
    return f"-{_wrap(node.get('arg'), _UNARY, first=True)}", _UNARY
  if op in {"SUM", "AVG", "MAX", "MIN"}:
    return f"{op}({', '.join(_expr_text(arg)[0] for arg in node.get('args') or [])})", _ATOM
  args = [f"{key}={_expr_text(value)[0]}" for key, value in node.items() if key != "op"]
  return f"{op}({', '.join(args)})", _ATOM


def _chain(op: str, nodes: Sequence[Any]) -> Any:
  """``nodes`` joined by ``op``: one n-ary node for ``+``/``*``, left-associative for ``-``/``/``."""
  if len(nodes) == 1:
    return nodes[0]
  if op in {"+", "*"}:
    return {"op": op, "args": list(nodes)}
  node = nodes[0]
  for right in nodes[1:]:
    node = {"op": op, "left": node, "right": right}
  return node


def formula_expr(spec: FormulaSpec) -> Any | None:
  """A typed spec as one expression tree of ``parse_formula``'s grammar, evaluating as the graph
  evaluates the spec; ``None`` for raw, valuation, ``SUM_RANGE`` and ``MEDIAN``."""
  params = spec.params or {}
  kind = spec.type
  if kind is FormulaType.arithmetic:
    if "expr" in params:
      return params["expr"]
    function = params.get("function")
    if function in {"SUM", "AVERAGE"} and params.get("items"):
      return {"op": "SUM" if function == "SUM" else "AVG", "args": list(params["items"])}
    operands = params.get("operands")
    if function is None and isinstance(operands, list) and operands:
      operator, values = "+", list(operands)
      if isinstance(operands[0], str) and operands[0] in {"+", "-", "*", "/"}:
        operator, values = operands[0], list(operands[1:])
      return _chain(operator, values) if values else None
    items = params.get("items")
    if function is None and isinstance(items, list) and items:
      return _chain("+", items)
    return None
  if kind is FormulaType.ref:
    node = params.get("source")
    if node is None:
      return None
    if params.get("adjustment") is not None:
      node = {"op": "+", "args": [node, float(params["adjustment"])]}
    return {"op": "NEG", "arg": node} if params.get("negate") else node
  if kind is FormulaType.driver:
    node = {"op": "*", "args": [params.get("base"), params.get("rate")]}
    for operation in driver_formula_params(params).scale_operations():
      if operation.operator == "divide":
        node = {"op": "/", "left": node, "right": operation.factor}
      else:
        node = {"op": "*", "args": [node, operation.factor]}
    return node
  if kind is FormulaType.ratio:
    node = {"op": "/", "left": params.get("numerator"), "right": params.get("denominator")}
    return {"op": "-", "left": node, "right": 1} if params.get("subtract_one") else node
  if kind is FormulaType.growth:
    return {"op": "*", "args": [params.get("base"), {"op": "+", "args": [1, params.get("rate")]}]}
  if kind is FormulaType.roll_forward:
    node = _chain("+", [params.get("beginning"), *(params.get("additions") or [])])
    return _chain("-", [node, *(params.get("subtractions") or [])])
  if kind is FormulaType.constant:
    value = params.get("value")
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None
  return None


def formula_text(spec: FormulaSpec | None) -> str | None:
  """``spec`` as infix text over item ids — the text ``set_formula`` would take to store an
  equivalent formula. A raw spec is its own text; a valuation (or ``SUM_RANGE``/``MEDIAN``)
  is ``<subtype or function>(<param>=<operand>, …)``, informative but not ours to parse."""
  if spec is None:
    return None
  params = spec.params or {}
  if spec.type is FormulaType.raw:
    return str(params.get("semantic_formula") or params.get("formula") or "")
  expr = formula_expr(spec)
  if expr is not None:
    return _expr_text(expr)[0]
  name = spec.subtype or params.get("function") or spec.type.value
  args = ", ".join(f"{key}={_expr_text(value)[0]}" for key, value in params.items() if key != "function")
  return f"{name}({args})"


# --------------------------------------------------------------------------
# recalc (ours)
# --------------------------------------------------------------------------


def scalar_anchors(model: FinancialModel) -> dict[str, int | None]:
  """Fixed-cell items — ``LineItem.column`` set, one cell such as a valuation output — → the
  period their cell anchors to (``renderer_columns._fixed_cell_anchor_period``)."""
  if not model._index:
    model.build_index()
  return {item.id: _fixed_cell_anchor_period(model, item) for item in model._index.values() if item.column is not None}


def _formula_owned_snapshots(graph: DependencyGraph, recompute: set[str]) -> None:
  """Drop every stored value a formula owns from ``graph``'s model: the cells of ``recompute``
  rows whose period selects a non-constant spec (a formula or a formula override).

  The graph seeds none of them, but it still reads them: a node of an item-level cycle (a line
  whose historical formula reads the line its projected formula feeds) takes its stored value
  when no user input is passed, a ``computed`` value wins over the evaluated one, and a missing
  operand falls back to it. Each serves the value the document last saved, so an edit upstream
  stops at that cell while its own override re-evaluates. With the snapshot gone the formula is
  the only source: it evaluates in the period order, or the cell is missing with its cause.
  """
  model = graph._loaded_model
  for item_id in recompute:
    item = model._index.get(item_id)
    if item is None or item.values is None:
      continue
    for period in list(item.values.values):
      spec = graph._spec_for_period(item, int(period))
      if spec is not None and spec.type is not FormulaType.constant:
        del item.values.values[period]


def _compute(
  model: FinancialModel, *,
  evaluation_reads: dict[tuple[str, int], tuple[tuple[str, int], ...]] | None = None,
  absences: Mapping[CellKey, EvidencedAbsentNumberValue] | None = None,
  evidenced_zeros: Mapping[CellKey, ZeroDisposition] | None = None,
) -> tuple[dict[str, dict[int, Value]], DependencyGraph]:
  """``compute_typed`` over ``model`` (mutated: pass a copy) with formula-owned rows recomputed
  and their stored snapshots dropped (``_formula_owned_snapshots``), and the graph that ran it.
  No input is substituted.
  """
  graph = DependencyGraph()
  graph.build(model)
  recompute = formula_owned_recompute_ids(model)
  _formula_owned_snapshots(graph, recompute)
  values = graph.compute_typed(
    {}, recompute=recompute, propagate_roots=set(),
    absences=dict(absences or {}), evidenced_zeros=dict(evidenced_zeros or {}),
  )
  if evaluation_reads is not None:
    evaluation_reads.clear()
    evaluation_reads.update(graph.last_eval_reads)
  return values, graph




def sum_may_omit_absence(disposition: AbsenceDisposition) -> bool:
  """Whether a plain balance-sheet sum may omit an evidenced absence of this disposition.

  A cited ``not_applicable`` and a period the agent declared its filing does not print the
  line (``not_disclosed_in_vintage``) are historical absences with no amount in any present
  term. A ``not_separately_presented`` amount is already carried by its ``included_in``
  line, so it is never dropped as a sum term, and an ``optional_unreported_evidenced``
  absence has no such sum rule. An unknown value (``source_observation_absent``,
  ``source_value_unavailable``, ``required_dependency_missing``) is not an absence at all
  and always propagates.
  """
  return disposition in ("not_applicable_evidenced", NOT_DISCLOSED_IN_VINTAGE)


def absence_projects_as_explicit_zero(disposition: AbsenceDisposition) -> bool:
  """Whether a schedule policy projects a balance whose whole history carries this
  disposition as an explicit zero: only a cited ``not_applicable_evidenced``. A
  ``not_disclosed_in_vintage`` period is a historical absence, never a zero."""
  return disposition == "not_applicable_evidenced"


def _column_source(column: Mapping[str, Any]) -> Mapping[str, Any]:
  source = column.get("source")
  return source if isinstance(source, Mapping) else {}


def _column_citation(column: Mapping[str, Any]) -> AbsenceCitation | None:
  """The cell's filing citation for an absence, when its column carries a complete one
  (``accession`` / ``presentation_role_uri`` / ``absence_statement``); a partial citation
  a stored column holds names no evidence the algebra can carry, so the absence stands
  without one."""
  source = _column_source(column)
  for candidate in (column.get("absence_citation"), source.get("citation"), source.get("absence_citation")):
    if not isinstance(candidate, Mapping):
      continue
    fields = {name: candidate.get(name) for name in ("accession", "presentation_role_uri", "absence_statement")}
    if all(isinstance(value, str) and value for value in fields.values()):
      return AbsenceCitation.model_validate(fields)
  return None


def _absent_value(
  doc: ModelSheet, key: CellKey, disposition: AbsenceDisposition, column: Mapping[str, Any],
  *, included_in: str | None = None,
) -> EvidencedAbsentNumberValue:
  item_id, period = key[0], int(key[1])
  projection = {int(year) for year in (doc.model.time_structure.projection_periods or doc.model.time_structure.projection_years or ())}
  return EvidencedAbsentNumberValue(
    state="evidenced_absent",
    primary_reason=disposition,
    causes=(AbsentCause(
      state="evidenced_absent", node_id=item_id,
      period=FiscalPeriod(
        period_id=str(period), fiscal_year=period, period_type="annual", is_estimate=period in projection,
      ),
      expression_path="root", reason_code=disposition,
    ),),
    citation=_column_citation(column),
    included_in=included_in,
  )


def _stored_zero(doc: ModelSheet, key: CellKey) -> bool:
  item = doc.model._index.get(key[0])
  if item is None:
    return False
  if item.values is not None:
    cell = item.values.values.get(int(key[1]))
    if cell is not None and cell.value is not None:
      return float(cell.value) == 0.0
  spec = _render_spec_for_period(doc.model, item, int(key[1]))
  return spec is not None and spec.type is FormulaType.constant and float(spec.params.get("value") or 0.0) == 0.0


def cell_dispositions(
  doc: ModelSheet,
) -> tuple[dict[CellKey, EvidencedAbsentNumberValue], dict[CellKey, ZeroDisposition]]:
  """The one derivation of the cell algebra's inputs from the document's stamps and
  source columns: ``(absences, evidenced_zeros)``, each keyed ``(item_id, period)``.

  A ``mark_unavailable`` whose reason is an ``AbsenceDisposition`` is an evidenced
  absence (its citation from the cell's ``absence_citation`` column when present); a
  reason that is not one (``source_observation_absent``, ``source_value_unavailable``,
  …) stays the owed input it is today. ``mark_not_separately_presented`` is an absence
  with ``included_in``. A source column with ``derivation == closure_corroborated_zero``
  is a closure-corroborated zero, an observed fact of ``0.0`` an observed zero, and a
  projected zero formula whose ``derivation.strategy`` is ``terminal_zero_fx_policy``
  (``schema.build_company_grounded_schedules._zero_projection_formula`` records the
  strategy only there) that policy's evidenced zero.
  """
  absences: dict[CellKey, EvidencedAbsentNumberValue] = {}
  zeros: dict[CellKey, ZeroDisposition] = {}
  for key, column in sorted(doc.semantic.items()):
    cell = (key[0], int(key[1]))
    source = _column_source(column)
    if column.get("cell") == "unavailable":
      reason = column.get("reason")
      if isinstance(reason, str) and is_absence_disposition(reason):
        absences[cell] = _absent_value(doc, cell, reason, column)
    elif column.get("cell") == "not_separately_presented":
      included = source.get("included_in") or column.get("included_in")
      absences[cell] = _absent_value(
        doc, cell, "not_separately_presented", column,
        included_in=str(included) if included else None,
      )
    derivation = source.get("derivation") or column.get("derivation")
    if derivation == "closure_corroborated_zero":
      zeros[cell] = "closure_corroborated_zero"
    elif column.get("origin") == "observed" and _stored_zero(doc, cell):
      zeros[cell] = "observed_zero"
  time = doc.model.time_structure
  periods = [int(year) for year in (*(time.historical_periods or time.historical_years or ()), *(time.projection_periods or time.projection_years or ()))]
  for item in doc.model._index.values():
    for period in periods:
      spec = _render_spec_for_period(doc.model, item, period)
      if (
        spec is not None and spec.derivation is not None
        and canonical_strategy(spec.derivation.strategy) == "terminal_zero_fx_policy"
      ):
        zeros[(item.id, period)] = "terminal_zero_fx_policy"
  return absences, zeros


class OptionalSumOmission(DependencyGraphKnownNumberValue):
  """A balance-sheet sum of its present terms.

  ``absences`` are ``optional_if_unreported`` terms omitted because they were
  evidenced absent. The citation stays on each absence cell's semantic column.
  """

  derivation: Literal["optional_absence_omitted"] = "optional_absence_omitted"
  absences: tuple[tuple[str, int], ...]


def _item_role(doc: ModelSheet, model: FinancialModel, item_id: str, period: int) -> str | None:
  item = model._index.get(item_id)
  if item is not None and item.required_role:
    return item.required_role
  role = (doc.semantic.get((item_id, period)) or {}).get("required_role")
  return role if isinstance(role, str) else None


def _plain_sum_refs(spec: FormulaSpec) -> list[LineItemRef] | None:
  """A ``SUM`` of item refs — the plain-sum shape — or ``None``."""
  params = spec.params or {}
  if spec.type is not FormulaType.arithmetic or params.get("function") != "SUM":
    return None
  refs: list[LineItemRef] = []
  for item in params.get("items") or ():
    ref = line_item_ref_from_obj(item)
    if ref is None:
      return None
    refs.append(ref)
  return refs


def _absence_cells(
  model: FinancialModel, absences: Mapping[CellKey, EvidencedAbsentNumberValue],
) -> list[CellKey]:
  """Evidenced-absence input cells a sum may omit: no formula selected there."""
  graph = _cell_graph(model)
  out: list[CellKey] = []
  for key, value in sorted(absences.items()):
    item = model._index.get(key[0])
    if not sum_may_omit_absence(value.primary_reason) or item is None:
      continue
    spec = graph._spec_for_period(item, key[1])
    if spec is None or spec.type is FormulaType.constant:
      out.append(key)
  return out


def evidenced_absence_cells(doc: ModelSheet) -> frozenset[CellKey]:
  """Every ``(item_id, period)`` whose disposition a schedule policy projects as an
  explicit zero (``absence_projects_as_explicit_zero``: a cited
  ``not_applicable_evidenced``): never a ``not_disclosed_in_vintage`` period, which a
  vintage did not print and is no zero.

  The model holds no cell there, so these are what tells a cited absence from a
  missing value; ``apply_company_grounded_schedule_policies`` reads them to
  project a balance whose whole history is cited absent as an explicit zero.
  """
  absences, _zeros = cell_dispositions(doc)
  return frozenset(
    key for key, value in absences.items() if absence_projects_as_explicit_zero(value.primary_reason)
  )


def _omit_optional_sum_terms(
  doc: ModelSheet, model: FinancialModel, absences: set[CellKey],
) -> tuple[dict[CellKey, tuple[CellKey, ...]], bool]:
  """Drop ``optional_if_unreported`` evidenced absences from balance-sheet sums on ``model``.

  The reader is a plain ``SUM`` on the balance-sheet plane: a balance-sheet
  statement total, or an item stamped with a balance-sheet role (the cash and
  short-term investments line). A term whose own role is ``optional_if_unreported``
  contributes nothing when evidenced absent; every other absence propagates.
  Returns the sum cells rewritten and whether some reader was not such a sum. An
  unmatched reader keeps the absence, which then propagates, and is not given a
  second interpretation.
  """
  if not absences:
    return {}, False
  from schema.required_role_catalog import (
    RequiredRole, RequiredRoleUnknown, load_required_role_catalog,
  )

  catalog = load_required_role_catalog()
  totals = frozenset(catalog.statement_totals())

  def role_row(role: str | None) -> RequiredRole | None:
    try:
      return catalog.role(role) if role else None
    except RequiredRoleUnknown:
      return None

  graph = _cell_graph(model)
  readers = _cell_readers(graph, model)
  seen = {reader for absence in absences for reader in readers.get(absence, ())}
  omitted: dict[CellKey, tuple[CellKey, ...]] = {}
  unknown = False
  for reader_id, reader_period in sorted(seen):
    item = model._index.get(reader_id)
    spec = None if item is None else graph._spec_for_period(item, reader_period)
    if item is None or spec is None:
      unknown = True
      continue
    reader = role_row(_item_role(doc, model, reader_id, reader_period))
    on_balance_sheet = reader_id in totals or (
      reader is not None and reader.statement == "balance_sheet"
    )
    terms = _plain_sum_refs(spec) if on_balance_sheet else None
    if terms is not None:
      drop: list[CellKey] = []
      keep: list[LineItemRef] = []
      for ref in terms:
        key = (ref.id, reader_period)
        term = role_row(_item_role(doc, model, ref.id, reader_period))
        if ref.t == 0 and term is not None and term.optional_if_unreported and key in absences:
          drop.append(key)
        else:
          keep.append(ref)
      if not drop:
        continue
      if item.overrides is None:
        item.overrides = {}
      if keep:
        item.overrides[reader_period] = FormulaSpec(
          type=FormulaType.arithmetic, subtype=spec.subtype,
          params={"function": "SUM", "items": keep}, note=spec.note,
        )
      else:
        item.overrides[reader_period] = FormulaSpec(
          type=FormulaType.constant, params={"value": 0.0}, note=spec.note,
        )
      omitted[(reader_id, reader_period)] = tuple(sorted(drop))
      continue

    unknown = True
  return omitted, unknown


def lower_optional_absences(doc: ModelSheet, model: FinancialModel) -> tuple[dict[CellKey, tuple[CellKey, ...]], bool]:
  """The one formula lowering ``recalc``, ``render`` and the publisher share: ``model`` (an indexed
  copy of ``doc.model``) has each balance-sheet sum's ``optional_if_unreported`` evidenced absences dropped,
  so the workbook's native formula reads the same present terms ours computes."""
  absences, _zeros = cell_dispositions(doc)
  return _omit_optional_sum_terms(doc, model, set(_absence_cells(model, absences)))


@dataclass(frozen=True)
class Evaluation:
  """One evaluation of a document: ``values`` is what ``recalc`` returns, ``typed`` the
  ``compute_typed`` result those values were projected from (before the optional-omission
  overlay), ``graph`` the ``DependencyGraph`` that computed it, built on the
  lowered copy (``lower_optional_absences``: the one formula lowering the renderer and the
  publisher share)."""

  values: Values
  typed: Mapping[str, Mapping[int, Value]]
  graph: DependencyGraph


def evaluate(
  doc: ModelSheet, *,
  evaluation_reads: dict[tuple[str, int], tuple[tuple[str, int], ...]] | None = None,
) -> Evaluation:
  """Every (item, period) through ``compute_typed``, formula-owned rows recomputed.

  The call is the renderer's (``render_model``): ``recompute`` = rows whose formula owns their
  output, so a formula edit is evaluated rather than served from the stored value snapshot; the
  snapshot under those formulas is dropped first (``_formula_owned_snapshots``), so no saved
  value stands in for a formula anywhere downstream of an edit, cycles included.
  A fixed-cell item is one cell: only its anchor period is kept. The document's stamps
  become the algebra's inputs in the one derivation (``cell_dispositions``): a cell the
  document marks unavailable with an ``AbsenceDisposition`` reason, or not separately
  presented, is an ``evidenced_absent`` leaf; a non-absence unavailable reason stays an
  owed missing input, and reads, in the graph, as an absent input
  (``source_value_unavailable``) whose cause ``_declared_reason`` retypes to the reason the
  document recorded (a scenario-output stamp's cause names the source row it stands for).

  An evidenced absence (``not_applicable_evidenced``, or ``not_disclosed_in_vintage``
  the agent declared) stays evidenced absent on its own cell and propagates through
  required operands. Only a balance-sheet sum whose term role
  is ``optional_if_unreported`` drops the term (``_omit_optional_sum_terms``);
  ``OptionalSumOmission`` names those absences.
  A closure-corroborated ``Known(0)`` is an evidenced-zero input value, not an absence, and stays ``0``.
  ``evaluation_reads``, when supplied, receives this exact evaluation's selected
  operands for provenance projection without a second graph computation.
  """
  model = copy_model(doc.model)
  model.build_index()
  absences, evidenced_zeros = cell_dispositions(doc)
  # The owed stamps: an unavailable reason that is no absence disposition names why an owed
  # input is missing; the absences above are the graph's own inputs and need no retyping.
  reasons = {
    key: reason for key, reason in unavailable_reasons(doc).items() if not is_absence_disposition(reason)
  }
  omitted, _unknown = lower_optional_absences(doc, model)
  typed, graph = _compute(
    model, evaluation_reads=evaluation_reads, absences=absences, evidenced_zeros=evidenced_zeros,
  )
  anchors = scalar_anchors(model)

  def _kept(item_id: str, period: int) -> bool:
    anchor = anchors.get(item_id)
    return anchor is None or anchor == period

  values: Values = {
    (item_id, int(period)): _declared_reason(value, reasons) if reasons else value
    for item_id, by_period in typed.items()
    for period, value in by_period.items()
    if _kept(item_id, int(period))
  }
  for key, dropped in omitted.items():
    value = values.get(key)
    if isinstance(value, DependencyGraphKnownNumberValue):
      values[key] = OptionalSumOmission(
        state="known", value=value.value, provenance=value.provenance, absences=dropped,
        evidenced_zero=value.evidenced_zero,
      )
  return Evaluation(values=values, typed=typed, graph=graph)


def recalc(
  doc: ModelSheet, *,
  evaluation_reads: dict[tuple[str, int], tuple[tuple[str, int], ...]] | None = None,
) -> Values:
  """The values of ``evaluate(doc)``."""
  return evaluate(doc, evaluation_reads=evaluation_reads).values


def operand_value(value: Value | None) -> float | None:
  """A cell's number when known. An evidenced absence is not a number."""
  return value.value if value is not None and value.state == "known" else None


def _is_absence_cell(key: CellKey, value: Value) -> bool:
  if value.state != "evidenced_absent" or not value.causes:
    return False
  cause = value.causes[0]
  return (cause.node_id, int(cause.period.fiscal_year)) == key


def evidenced_absence_summary(values: Values, doc: ModelSheet | None = None) -> dict[str, Any]:
  """Cited absences in one recalc, and which a balance-sheet sum omitted.

  ``optionality_unknown`` is true when a reader was not a balance-sheet sum
  with an ``optional_if_unreported`` term, so that absence propagated. Without ``doc``
  the flag is true whenever an absence was not omitted: the summary cannot
  re-read the catalog.
  """
  absence_cells = sorted(key for key, value in values.items() if _is_absence_cell(key, value))
  derived = [value for value in values.values() if isinstance(value, OptionalSumOmission)]
  omitted = sorted({absence for value in derived for absence in value.absences})
  unknown = any(key not in set(omitted) for key in absence_cells)
  if doc is not None and absence_cells:
    _, unknown = _omit_optional_sum_terms(doc, copy_model(doc.model), set(absence_cells))
  return {
    "absences": len(absence_cells),
    "omitted": len(omitted),
    "dependents": len(derived),
    "optionality_unknown": unknown,
    "omitted_from": [
      [key[0], key[1], absence[0], absence[1]]
      for key, value in sorted(values.items())
      if isinstance(value, OptionalSumOmission)
      for absence in value.absences
    ],
  }



def _declared_reason(value: Value, reasons: Mapping[tuple[str, int], str]) -> Value:
  """A missing value retyped by the document's unavailable stamps: a cause whose cell a
  stamp declares takes the stamp's reason, and keeps naming the stamped cell itself."""
  if value.state != "missing":
    return value
  causes = []
  for cause in value.causes:
    if cause.reason_code != "source_value_unavailable":
      causes.append(cause)
      continue
    key = (cause.node_id, int(cause.period.fiscal_year))
    if key not in reasons:
      causes.append(cause)
      continue
    causes.append(cause.model_copy(update={"reason_code": reasons[key]}))
  if all(new is old for new, old in zip(causes, value.causes)):
    return value
  return type(value).model_validate({
    **value.model_dump(),
    "primary_reason": causes[0].reason_code,
    "causes": tuple(causes),
  })


def value_row(value: Value) -> dict[str, Any]:
  """JSON form: state, value, primary reason, and the first cause's node-period. An
  evidenced-absent cell also names its ``included_in`` line and filing ``citation``
  when it carries them. A sum that omitted a declared-optional absence names those
  cells (``optional_absence_omitted``)."""
  row: dict[str, Any] = {"state": value.state, "value": value.value if value.state == "known" else None}
  if value.state != "known":
    row["reason"] = value.primary_reason
    cause = value.causes[0]
    row["cause"] = [cause.node_id, cause.period.fiscal_year]
  if value.state == "evidenced_absent":
    if value.included_in is not None:
      row["included_in"] = value.included_in
    if value.citation is not None:
      row["citation"] = value.citation.model_dump(mode="json")
  if isinstance(value, OptionalSumOmission):
    row.update(derivation=value.derivation, evidenced_absences=[list(key) for key in value.absences])
  return row


def _same(left: Value | None, right: Value | None, tolerance: float) -> bool:
  if left is None or right is None:
    return left is right
  if left.state != right.state:
    return False
  if left.state == "known" and right.state == "known":
    return writer_numeric_equal(left.value, right.value, tolerance)
  return value_row(left) == value_row(right)


def value_diff(before: Values, after: Values, *, tolerance: float = 1e-9) -> list[dict[str, Any]]:
  """Every (item, period) whose state, value or reason differs between two recalcs."""
  rows: list[dict[str, Any]] = []
  for key in sorted(set(before) | set(after)):
    left, right = before.get(key), after.get(key)
    if _same(left, right, tolerance):
      continue
    rows.append({
      "item": key[0],
      "period": key[1],
      "before": None if left is None else value_row(left),
      "after": None if right is None else value_row(right),
    })
  return rows


def state_counts(values: Values) -> dict[str, Any]:
  states = Counter(value.state for value in values.values())
  reasons = Counter(value.primary_reason for value in values.values() if value.state != "known")
  return {"cells": len(values), "states": dict(states), "reasons": dict(reasons)}


def values_digest(values: Values) -> str:
  """sha256 over the sorted ``[item, period, state, value]`` rows: equal digests, equal recalcs."""
  rows = [
    [item, period, value.state, value.value if value.state == "known" else None]
    for (item, period), value in sorted(values.items(), key=lambda pair: pair[0])
  ]
  return hashlib.sha256(json.dumps(rows, separators=(",", ":")).encode("utf-8")).hexdigest()


def override_pins(
  doc: ModelSheet, values: Values, *, items: Collection[str] | None = None,
) -> dict[tuple[str, int], dict[str, Any]]:
  """Every (item, period) a constant override pins while a formula is declared there: the
  override, and what that formula gives over the recalc's operand values (``None`` when it
  cannot be evaluated). ``formula_state`` is ``known`` when it evaluates; otherwise
  ``formula_state`` / ``formula_reason`` / ``formula_cause`` are ours' for that cell with the
  override absent. The override is what the product serves; nothing here changes that.
  ``items`` narrows the pins to those lines (a reader of a few lines does not unpin and
  re-evaluate the whole document); every line when ``None``."""
  model = copy_model(doc.model)
  model.build_index()
  pinned: dict[tuple[str, int], float | None] = {}
  for item in model._index.values():
    if items is not None and item.id not in items:
      continue
    for period, spec in list((item.overrides or {}).items()):
      if spec.type is not FormulaType.constant:
        continue
      assert item.overrides is not None
      del item.overrides[period]
      if _render_spec_for_period(model, item, int(period)) is None:
        continue
      value = spec.params.get("value")
      pinned[(item.id, int(period))] = None if value is None else float(value)
  if not pinned:
    return {}
  graph = DependencyGraph()
  graph.build(model)
  results: dict[str, dict[int, float | None]] = {}
  for (item_id, period), value in values.items():
    operand = operand_value(value)
    if operand is not None:
      results.setdefault(item_id, {})[period] = operand
  time_order = graph._time_order()
  time_index = {period: index for index, period in enumerate(time_order)}
  out: dict[tuple[str, int], dict[str, Any]] = {}
  for (item_id, period), override in pinned.items():
    try:
      formula = graph._eval(item_id, period, results, time_index, time_order)
    except (ArithmeticError, KeyError, TypeError, ValueError):
      formula = None
    out[(item_id, period)] = {
      "override": override, "formula": formula, "formula_state": "known" if formula is not None else None,
      "formula_reason": None,
    }
  unevaluated = [key for key, pin in out.items() if pin["formula"] is None]
  if unevaluated:
    typed, _ = _compute(model)
    for item_id, period in unevaluated:
      value = typed.get(item_id, {}).get(period)
      if value is None:
        continue
      row = value_row(value)
      out[(item_id, period)].update(
        formula=row["value"], formula_state=row["state"], formula_reason=row.get("reason"),
        formula_cause=row.get("cause"),
      )
  return out


def override_pinned_items(
  pins: Mapping[tuple[str, int], Mapping[str, Any]], dependents: Mapping[str, set[str]],
) -> list[dict[str, Any]]:
  """The items whose declared formula disagrees with the constant override that wins, per period;
  ``orphaned`` when no line reads the item (``dependents``: ``dependents_map``)."""
  by_item: dict[str, dict[int, dict[str, float | None]]] = {}
  for (item_id, period), pin in sorted(pins.items()):
    override, formula = pin["override"], pin["formula"]
    if override is None or formula is None or math.isclose(override, formula, rel_tol=1e-6, abs_tol=1e-9):
      continue
    by_item.setdefault(item_id, {})[period] = {"override": override, "formula": formula, "difference": override - formula}
  return [
    {"item": item_id, "orphaned": not dependents.get(item_id), "periods": periods}
    for item_id, periods in by_item.items()
  ]


def override_hides_missing_formula(pins: Mapping[tuple[str, int], Mapping[str, Any]]) -> list[dict[str, Any]]:
  """The pinned cells whose declared formula ours cannot evaluate: the constant override keeps
  each known, so without this list the formula's failure is invisible."""
  return [
    {
      "item": item_id, "period": period, "override_value": pin["override"],
      "formula_state": pin["formula_state"], "formula_reason": pin["formula_reason"],
      "formula_cause": pin.get("formula_cause"),
    }
    for (item_id, period), pin in sorted(pins.items())
    if pin["formula_state"] != "known"
  ]


# --------------------------------------------------------------------------
# inspect
# --------------------------------------------------------------------------


def _cleared_cells(doc: ModelSheet) -> set[tuple[str, int]]:
  cleared: set[tuple[str, int]] = set()
  for edit in doc.edits:
    op, result = edit.get("op") or {}, edit.get("result") or {}
    if op.get("type") == "clear_input" and result.get("cleared_by") == "agent":
      cleared.add((str(op["item_id"]), int(op["period"])))
  return cleared


def _formula_entry(spec: FormulaSpec | None) -> dict[str, Any] | None:
  if spec is None:
    return None
  entry: dict[str, Any] = {"type": spec.type.value, "text": formula_text(spec)}
  if spec.subtype is not None:
    entry["subtype"] = spec.subtype
  if spec.derivation is not None:
    entry["derivation"] = spec.derivation.strategy
  return entry


def _periods_entry(item: LineItem, spec: FormulaSpec | None, periods: Sequence[int]) -> dict[str, Any] | None:
  """``spec``'s entry, or — no formula for these periods — ``{type: input, periods}`` for the
  periods holding a stored value (``override_periods`` beside it for constant overrides), or
  ``{type: override, periods}`` when only overrides do; ``None`` when nothing does."""
  if spec is not None:
    return _formula_entry(spec)
  stored = item.values.values if item.values is not None else {}
  inputs = [period for period in periods if stored.get(period) is not None and stored[period].value is not None]
  overrides = [period for period in periods if period in (item.overrides or {})]
  if inputs:
    return {"type": "input", "periods": inputs, **({"override_periods": overrides} if overrides else {})}
  return {"type": "override", "periods": overrides} if overrides else None


def _cell_status(
  model: FinancialModel,
  item: LineItem,
  period: int,
  value: Value | None,
  *,
  recompute: set[str],
  pin: Mapping[str, Any] | None,
  cleared: bool,
  semantic: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
  """How one (item, period) is produced, mirroring the graph's seed order: a constant override
  (``override``), a formula the recalc evaluates (``formula``), a stored value beside a formula
  the recalc does not own (``snapshot``), a stored value alone (``input``), or nothing
  (``missing``); with ours' state, value, reason and origin, and the cell's semantic column
  (``semantic``: ``origin``, ``source``, an agent op's ``rationale`` and ``edit_index``) when a
  writer recorded one."""
  override = (item.overrides or {}).get(period)
  spec = _render_spec_for_period(model, item, period)
  stored = item.values.values.get(period) if item.values is not None else None
  extra: dict[str, Any] = {}
  if override is not None and override.type is FormulaType.constant:
    status = "override"
    extra["override_value"] = override.params.get("value")
    if pin is not None:
      extra["formula_value"] = pin["formula"]
      if pin["formula_state"] != "known":
        extra["formula_state"], extra["formula_reason"] = pin["formula_state"], pin["formula_reason"]
  elif override is not None:
    status = "formula"
    extra["override_formula"] = formula_text(override)
  elif spec is not None and spec.type is not FormulaType.constant and item.id in recompute:
    status = "formula"
  elif stored is not None:
    status = "snapshot" if spec is not None else "input"
  elif spec is not None:
    status = "formula"
  else:
    status = "missing"
  if stored is not None and status in {"input", "snapshot"}:
    extra["provenance"] = stored.provenance.value
    if stored.edgar_provenance is not None:
      extra["source_tag"] = stored.edgar_provenance.metric_tag
  if override is not None and override.edgar_provenance is not None:
    extra["source_tag"] = override.edgar_provenance.metric_tag
  if value is not None and value.state == "known":
    extra["origin"] = value.provenance.value
  if cleared and override is None and stored is None:
    extra["cleared_by"] = "agent"
  if semantic is not None:
    extra["semantic"] = dict(semantic)
  ours = value_row(value) if value is not None else {"state": None, "value": None}
  return {"period": period, "status": status, **ours, **extra}


@dataclass(frozen=True)
class Inspected:
  """``read_lines``' answer: the ``lines`` found, in document order, and the requested ``not_found`` ids
  (sorted) that are no line of the document."""

  lines: list[dict[str, Any]]
  not_found: tuple[str, ...]


def _unknown_items(model: FinancialModel, items: Sequence[str] | None) -> tuple[str, ...]:
  if not model._index:
    model.build_index()
  return tuple(sorted(set(items or ()) - set(model._index)))


def _id_leaf_tokens(item_id: str) -> list[str]:
  return [token for token in item_id.rsplit(".", 1)[-1].split("_") if token]


def closest_items(model: FinancialModel, item_id: str, *, limit: int = 3, floor: float = 0.5) -> list[str]:
  """The ids of the lines of ``model`` nearest ``item_id``, best first (at most ``limit``; none scoring
  under ``floor``), for an id the document does not hold.

  The matcher is plain and local. An id that is a prefix of a line's id (or the other way) comes first,
  then the lines of the same section (``item_id`` without its leaf), then the rest; within each, by
  leaf score. The leaf names (after the last ``.``) are split on ``_`` and each token is compared to the
  other leaf's tokens by ``difflib.SequenceMatcher.ratio`` (a character edit-distance ratio, so
  ``pretax`` ~ ``taxes``; a ratio under ``floor`` counts as no match). The leaf score is the harmonic
  mean of the two coverages (the mean best match of ``item_id``'s tokens over the line's, and the
  line's over ``item_id``'s), so a line saying the same thing in more words ranks below one saying it in
  the same ones. Ties break on the id."""
  from difflib import SequenceMatcher

  if not model._index:
    model.build_index()
  wanted = _id_leaf_tokens(item_id)
  parent = item_id.rpartition(".")[0]
  similar: dict[tuple[str, str], float] = {}

  def ratio(left: str, right: str) -> float:
    key = (left, right)
    if key not in similar:
      value = SequenceMatcher(None, left, right).ratio()
      similar[key] = value if value >= floor else 0.0
    return similar[key]

  scored: list[tuple[bool, bool, float, str]] = []
  for candidate in model._index:
    tokens = _id_leaf_tokens(candidate)
    if not wanted or not tokens:
      continue
    toward = sum(max(ratio(token, other) for other in tokens) for token in wanted) / len(wanted)
    back = sum(max(ratio(token, other) for other in wanted) for token in tokens) / len(tokens)
    score = 0.0 if toward + back == 0 else 2 * toward * back / (toward + back)
    prefixed = candidate.startswith(item_id) or item_id.startswith(candidate)
    if prefixed or score >= floor:
      scored.append((prefixed, candidate.rpartition(".")[0] == parent, score, candidate))
  scored.sort(key=lambda row: (not row[0], not row[1], -row[2], row[3]))
  return [candidate for _prefixed, _same, _score, candidate in scored[:limit]]


def read_lines(
  doc: ModelSheet,
  values: Values | None = None,
  *,
  section: str | None = None,
  items: Sequence[str] | None = None,
) -> Inspected:
  """Each line of ``section`` (or each of ``items``; neither: every line), in document order:
  identity, unit and scale, historical/projected formulas as infix text (``_periods_entry``),
  operands and dependents (item ids), its formulas' scale ``warnings`` when any, and per period
  how its cell is produced (``_cell_status``). An id of ``items`` that is no line is returned in
  ``not_found`` and read nothing (``closest_items`` names its nearest lines)."""
  model = doc.model
  if not model._index:
    model.build_index()
  section_ids = sorted({located.id for _sheet, located in _sections(model)})
  if section is not None and section not in section_ids:
    raise ValueError(f"section {section!r} is not a section of the document; sections: {section_ids}")
  unknown = _unknown_items(model, items)
  wanted = None if items is None else set(items) - set(unknown)
  if wanted is not None and not wanted:
    return Inspected([], unknown)
  values = recalc(doc) if values is None else values
  recompute = formula_owned_recompute_ids(model)
  dependents = dependents_map(model)
  operands: dict[str, set[str]] = {}
  for operand, readers in dependents.items():
    for reader in readers:
      operands.setdefault(reader, set()).add(operand)
  pins = override_pins(doc, values, items=wanted)
  cleared = _cleared_cells(doc)
  anchors = scalar_anchors(model)
  historical = [int(period) for period in model.time_structure.historical_periods]
  projection = [int(period) for period in model.time_structure.projection_periods]
  periods = [*historical, *projection]
  lines: list[dict[str, Any]] = []
  for sheet_name, located in _sections(model):
    if section is not None and located.id != section:
      continue
    for item in located.line_items:
      if wanted is not None and item.id not in wanted:
        continue
      line: dict[str, Any] = {
        "id": item.id,
        "label": item.label,
        "sheet": sheet_name,
        "section": located.id,
        "row": item.row,
        "item_type": item.item_type.value,
        "unit": None if item.unit is None else item.unit.value,
        "model_scale": item.model_scale,
        "historical": _periods_entry(item, item.historical, historical),
        "projected": _periods_entry(item, item.projected, projection),
        "operands": sorted(operands.get(item.id, ())),
        "dependents": sorted(dependents.get(item.id, ())),
      }
      warnings = list(dict.fromkeys(
        warning for spec in (item.historical, item.projected) if spec is not None
        for warning in scale_warnings(spec, model, item.id)
      ))
      if warnings:
        line["warnings"] = warnings
      anchor = anchors.get(item.id)
      if item.id in anchors:
        line["scalar"] = {"column": item.column, "anchor_period": anchor}
      if item.item_type in {ItemType.header, ItemType.spacer}:
        line["cells"] = []
      else:
        line["cells"] = [
          _cell_status(
            model, item, period, values.get((item.id, period)),
            recompute=recompute, pin=pins.get((item.id, period)), cleared=(item.id, period) in cleared,
            semantic=doc.semantic.get((item.id, period)),
          )
          for period in ([anchor] if anchor is not None else periods)
        ]
      lines.append(line)
  return Inspected(lines, unknown)


# --------------------------------------------------------------------------
# render, native recalc, compare
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class RenderedCell:
  sheet: str
  cell: str
  kind: str
  formula: Any
  tolerance: float
  family: str


@dataclass(frozen=True)
class Rendered:
  path: Path
  cells: dict[tuple[str, int], RenderedCell]
  marked_unavailable: int
  seconds: float
  #: Every write of the render plan, ``(sheet, cell)`` → value or formula text.
  writes: dict[tuple[str, str], Any] = field(default_factory=dict)



def _family(model: FinancialModel, item: LineItem, period: int) -> str:
  spec = _render_spec_for_period(model, item, period)
  if spec is None:
    return "input"
  return spec.type.value if spec.subtype is None else f"{spec.type.value}/{spec.subtype}"


def _mappers(model: FinancialModel) -> dict[str, AbsoluteColumnMapper]:
  out: dict[str, AbsoluteColumnMapper] = {}
  for sheet_name, sheet in model.sheets.items():
    layout = sheet.layout or SheetLayout(label_column="A", first_data_column="B")
    out[sheet_name] = AbsoluteColumnMapper(
      model.time_structure,
      (layout.first_data_column or "B").upper(),
      period_scope=layout.period_scope,
    )
  return out


def _resolve_semantic_text(text: str, compiler: ExcelFormulaCompiler, period: int) -> str:
  """A raw spec's semantic text with each item id resolved to its cell for ``period``.

  An item id stays a reference even before ``(`` (``parse_formula`` refuses that form; a
  stored one renders as the cell it reads, never a bare name the spreadsheet cannot resolve).
  """
  tokens = _tokens(text)
  pieces: list[str] = []
  cursor = 0
  skip_offset = False
  for index, token in enumerate(tokens):
    if token.kind == "offset" and skip_offset:
      pieces.append(text[cursor:token.start])
      cursor = token.end
      skip_offset = False
      continue
    skip_offset = False
    if token.kind != "name" or token.text.upper() in _RAW_CONSTANTS:
      continue
    if _is_call(tokens, index) and (
      token.text.upper() in _GRAMMAR_FUNCTIONS or token.text not in compiler._item_row_map
    ):
      continue
    has_offset = index + 1 < len(tokens) and tokens[index + 1].kind == "offset"
    ref = _ref_at(tokens, index)
    try:
      cell = compiler._resolve_ref(ref.id, ref.t, period, period_anchor=ref.period_anchor)
    except (KeyError, ValueError):
      cell = "NA()"
    pieces.append(text[cursor:token.start])
    pieces.append(cell)
    cursor = token.end
    skip_offset = has_offset
  pieces.append(text[cursor:])
  return "=" + "".join(pieces)


def _materialize_semantic_raw(model: FinancialModel) -> int:
  """Per-period ``raw`` overrides with cells resolved, for every raw spec over semantic ids."""
  mappers = _mappers(model)
  written = 0
  for sheet_name, section in list(_sections(model)):
    compiler = ExcelFormulaCompiler(model, mappers, sheet_name)
    for item in section.line_items:
      if item.column is not None:
        anchor = _fixed_cell_anchor_period(model, item)
        periods = [] if anchor is None else [int(anchor)]
      else:
        periods = [int(period) for period in mappers[sheet_name].all_periods()]
      for period in periods:
        spec = _render_spec_for_period(model, item, period)
        if spec is None or spec.type is not FormulaType.raw or "semantic_formula" not in spec.params:
          continue
        formula = _resolve_semantic_text(str(spec.params["semantic_formula"]), compiler, period)
        if item.overrides is None:
          item.overrides = {}
        item.overrides[period] = FormulaSpec(type=FormulaType.raw, params={"formula": formula})
        written += 1
  return written


def render(
  doc: ModelSheet, path: Path, *, ours: Values | None = None, mark_unavailable: bool = True,
  gaps: Collection[CellKey] = (),
) -> Rendered:
  """The product renderer and XLSX writer over the document; ``path`` receives the workbook.

  With ``mark_unavailable`` every non-known (item, period) of ours that the row owns
  (``publish.document_owns_cell`` → ``schema.renderer_columns.item_owns_cell``, with ``gaps`` the
  source bundle's ``owned_cell_gaps``) and would render blank is written ``=NA()`` through the renderer's own
  ``required_unavailable`` path, so a missing input — including a cited evidenced absence —
  is ``#N/A`` in the cell and its dependents follow natively. The reason stays in ours.
  """
  started = time.perf_counter()
  model = copy_model(doc.model)
  model.build_index()
  lower_optional_absences(doc, model)
  unavailable: frozenset[tuple[str, int]] = frozenset()
  if mark_unavailable:
    from model_build.publish import document_owns_cell

    values = ours if ours is not None else recalc(doc)
    unavailable = frozenset(
      key for key, value in values.items()
      if value.state != "known" and document_owns_cell(doc, model, key[0], key[1], gaps)
    )
  families = {
    (item.id, int(period)): _family(model, item, int(period))
    for _sheet, section in _sections(model)
    for item in section.line_items
    for period in model.time_structure.historical_periods + model.time_structure.projection_periods
  }
  _materialize_semantic_raw(model)
  context = AnnualRenderContext(version_label=f"Bench document {doc.source}", required_unavailable=unavailable)
  plan = render_model(model, annual_context=context)
  target = Path(path).expanduser()
  target.parent.mkdir(parents=True, exist_ok=True)
  target.write_bytes(render_xlsx_bytes(plan, finalized_at=_RENDER_TIME))
  writes = {(write.sheet, write.cell): write.value for write in plan.writes}
  cells: dict[tuple[str, int], RenderedCell] = {}
  marked = 0
  for entry in plan.semantic_cells:
    if entry.cell_kind == "presentation":
      continue
    if entry.cell_kind == "unavailable" and writes.get((entry.sheet, entry.cell)) == "=NA()":
      marked += 1
    key = (entry.semantic_item_id, int(entry.period))
    cells[key] = RenderedCell(
      sheet=entry.sheet,
      cell=entry.cell,
      kind=entry.cell_kind,
      formula=writes.get((entry.sheet, entry.cell)),
      tolerance=float(entry.tolerance),
      family=families.get(key, "fixed_cell"),
    )
  return Rendered(
    path=target, cells=cells, marked_unavailable=marked, seconds=time.perf_counter() - started, writes=writes,
  )


class NativeRecalcError(RuntimeError):
  pass


def soffice_path(explicit: str | None = None) -> str:
  """``explicit``, ``$BENCH_SOFFICE``, ``soffice`` on PATH, or the macOS install."""
  for candidate in (explicit, os.environ.get("BENCH_SOFFICE"), shutil.which("soffice"), *_SOFFICE_CANDIDATES):
    if candidate and Path(candidate).exists():
      return candidate
  raise NativeRecalcError("no LibreOffice soffice found (set BENCH_SOFFICE)")


def native_recalc(
  xlsx_path: Path,
  *,
  out_path: Path | None = None,
  soffice: str | None = None,
  timeout: float = 600.0,
) -> dict[tuple[str, str], Any]:
  """LibreOffice recalculates the workbook headless; every cell's value, by (sheet, cell).

  ``soffice --convert-to xlsx`` loads the workbook (formula cells carry no cached value, so
  LibreOffice computes each one) and writes cached results, read back with openpyxl
  ``data_only``. Every formula cell of the input is in the result, ``None`` when LibreOffice
  left it uncomputed; errors read as their text (``#N/A``, ``#DIV/0!``…). A private profile
  per call keeps the run independent of any running LibreOffice.
  """
  import openpyxl

  source = Path(xlsx_path).expanduser().resolve()
  executable = soffice_path(soffice)
  with tempfile.TemporaryDirectory(prefix="bench-native-") as scratch:
    scratch_dir = Path(scratch)
    out_dir = scratch_dir / "out"
    command = [
      executable,
      f"-env:UserInstallation={(scratch_dir / 'profile').as_uri()}",
      "--headless", "--norestore", "--nolockcheck", "--calc",
      "--convert-to", "xlsx:Calc Office Open XML",
      "--outdir", str(out_dir),
      str(source),
    ]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=timeout, check=False)
    produced = out_dir / source.name
    if completed.returncode != 0 or not produced.is_file():
      raise NativeRecalcError(
        f"soffice exit {completed.returncode}; no {produced.name}: {(completed.stderr or completed.stdout)[-2000:]}"
      )
    if out_path is not None:
      target = Path(out_path).expanduser()
      target.parent.mkdir(parents=True, exist_ok=True)
      shutil.copyfile(produced, target)
    computed = openpyxl.load_workbook(produced, data_only=True)
    formulas = openpyxl.load_workbook(source, data_only=False)
    values: dict[tuple[str, str], Any] = {}
    for sheet in computed.worksheets:
      for row in sheet.iter_rows():
        for cell in row:
          if cell.value is not None:
            values[(sheet.title, cell.coordinate)] = cell.value
    for sheet in formulas.worksheets:
      for row in sheet.iter_rows():
        for cell in row:
          if isinstance(cell.value, str) and cell.value.startswith("="):
            values.setdefault((sheet.title, cell.coordinate), None)
    computed.close()
    formulas.close()
  return values


@dataclass(frozen=True)
class Diff:
  item: str
  period: int
  sheet: str
  cell: str
  kind: str
  family: str
  ours: dict[str, Any]
  native: Any
  formula: Any


DIFF_KINDS = (
  "value_mismatch",
  "ours_known_native_unavailable",
  "ours_unavailable_native_value",
  "ours_unsupported_native_value",
  "native_error_vs_ours_unavailable",
)
#: Native errors a broken formula raises (an unresolved name, a dangling reference, an
#: operand of the wrong type), never a missing input propagating (``#N/A``, ``#DIV/0!``).
_BROKEN_FORMULA_ERRORS = frozenset({"#NAME?", "#REF!", "#VALUE!"})


def _numeric(value: Any) -> float | None:
  if isinstance(value, bool) or not isinstance(value, (int, float)):
    return None
  return float(value)


def _is_error(value: Any) -> bool:
  return isinstance(value, str) and value.startswith("#") and value[1:2].isupper()


def _cell_kind(mine: Value, theirs: Any, cell: RenderedCell, rel_tolerance: float) -> str:
  number = _numeric(theirs)
  if mine.state == "known":
    if number is None:
      return "ours_known_native_unavailable"
    tolerance = max(cell.tolerance, rel_tolerance * max(1.0, abs(mine.value)))
    return "agree_value" if abs(number - mine.value) <= tolerance else "value_mismatch"
  if number is not None:
    return "ours_unsupported_native_value" if mine.state == "unsupported" else "ours_unavailable_native_value"
  if isinstance(theirs, str) and theirs in _BROKEN_FORMULA_ERRORS:
    return "native_error_vs_ours_unavailable"
  return "agree_unavailable" if _is_error(theirs) else "agree_nonnumeric"


def _classified(
  ours: Values,
  native: Mapping[tuple[str, str], Any],
  provenance: Mapping[tuple[str, int], RenderedCell],
  rel_tolerance: float,
) -> Iterator[tuple[tuple[str, int], RenderedCell, Value, Any, str]]:
  for key in sorted(provenance):
    cell = provenance[key]
    mine = ours.get(key)
    if mine is None:
      continue
    theirs = native.get((cell.sheet, cell.cell))
    yield key, cell, mine, theirs, _cell_kind(mine, theirs, cell, rel_tolerance)


def compare(
  ours: Values,
  native: Mapping[tuple[str, str], Any],
  provenance: Mapping[tuple[str, int], RenderedCell],
  *,
  rel_tolerance: float = 1e-9,
) -> list[Diff]:
  """Disagreements of ours and native over every rendered semantic cell.

  Tolerance per cell is ``max(cell.tolerance, rel_tolerance·max(1,|ours|))``. Ours non-known
  agrees with a native ``#N/A``/``#DIV/0!``-style error, blank or text, but a native
  ``#NAME?``/``#REF!``/``#VALUE!`` beside it is ``native_error_vs_ours_unavailable`` (the
  rendered formula is broken, whatever ours holds); an ``unsupported`` ours beside a native
  number is its own kind (the formula only the native engine evaluates). A cited evidenced
  absence is missing, so it agrees with a native ``#N/A`` like any other unavailable cell.
  """
  return [
    Diff(
      item=key[0], period=key[1], sheet=cell.sheet, cell=cell.cell, kind=kind,
      family=cell.family, ours=value_row(mine), native=theirs, formula=cell.formula,
    )
    for key, cell, mine, theirs, kind in _classified(ours, native, provenance, rel_tolerance)
    if not kind.startswith("agree_")
  ]


def agreement(
  ours: Values,
  native: Mapping[tuple[str, str], Any],
  provenance: Mapping[tuple[str, int], RenderedCell],
  *,
  rel_tolerance: float = 1e-9,
) -> dict[str, int]:
  """Every compared cell's kind, agreements included (``agree_value`` = numerically equal)."""
  return dict(Counter(kind for *_rest, kind in _classified(ours, native, provenance, rel_tolerance)))


def summarize(diffs: Sequence[Diff], *, compared: int, sample: int = 10) -> dict[str, Any]:
  return {
    "compared_cells": compared,
    "disagreements": len(diffs),
    "by_kind": dict(Counter(diff.kind for diff in diffs)),
    "by_family": dict(Counter(diff.family for diff in diffs)),
    "by_kind_family": dict(Counter(f"{diff.kind}:{diff.family}" for diff in diffs)),
    "sample": [diff.__dict__ for diff in diffs[:sample]],
  }


def compare_rows(
  ours: Values,
  native: Mapping[tuple[str, str], Any],
  provenance: Mapping[tuple[str, int], RenderedCell],
  *,
  rel_tolerance: float = 1e-9,
) -> list[dict[str, Any]]:
  """Every compared (item, period): its cell and formula, ours, native, the kind and ``agree``."""
  return [
    {
      "item": key[0], "period": key[1], "cell": f"{cell.sheet}!{cell.cell}", "family": cell.family,
      "formula": cell.formula, "ours": value_row(mine), "native": theirs, "kind": kind,
      "agree": kind.startswith("agree_"),
    }
    for key, cell, mine, theirs, kind in _classified(ours, native, provenance, rel_tolerance)
  ]


def native_rows(rendered: Rendered, native: Mapping[tuple[str, str], Any]) -> list[dict[str, Any]]:
  """Every rendered semantic (item, period): its cell, render kind, formula and native value."""
  return [
    {
      "item": key[0], "period": key[1], "cell": f"{cell.sheet}!{cell.cell}", "render_kind": cell.kind,
      "formula": cell.formula, "native": native.get((cell.sheet, cell.cell)),
    }
    for key, cell in sorted(rendered.cells.items())
  ]


def write_jsonl(path: Path, rows: Iterable[Mapping[str, Any]]) -> int:
  """``rows`` one JSON object per line (non-JSON values as text); returns the row count."""
  path.parent.mkdir(parents=True, exist_ok=True)
  count = 0
  with path.open("w", encoding="utf-8") as handle:
    for row in rows:
      handle.write(json.dumps(row, separators=(",", ":"), ensure_ascii=False, default=str) + "\n")
      count += 1
  return count


def ours_vs_native(
  doc: ModelSheet, workdir: Path, name: str, *, soffice: str | None = None,
  recalculate: Callable[[ModelSheet], Values] | None = None,
) -> dict[str, Any]:
  """recalc, render, native recalc and compare one document; wall time per step."""
  workdir = Path(workdir).expanduser()
  started = time.perf_counter()
  ours = recalc(doc) if recalculate is None else recalculate(doc)
  ours_seconds = time.perf_counter() - started
  rendered = render(doc, workdir / f"{name}.xlsx", ours=ours)
  started = time.perf_counter()
  native = native_recalc(rendered.path, out_path=workdir / f"{name}.native.xlsx", soffice=soffice)
  native_seconds = time.perf_counter() - started
  diffs = compare(ours, native, rendered.cells)
  uncomputed = sum(1 for value in native.values() if value is None)
  return {
    "xlsx": str(rendered.path),
    "native_xlsx": str(workdir / f"{name}.native.xlsx"),
    "seconds": {"ours_recalc": ours_seconds, "render": rendered.seconds, "native_recalc": native_seconds},
    "ours": state_counts(ours),
    "rendered_cells": len(rendered.cells),
    "rendered_cell_kinds": dict(Counter(cell.kind for cell in rendered.cells.values())),
    "formula_families": dict(Counter(cell.family for cell in rendered.cells.values() if cell.kind == "formula")),
    "marked_unavailable": rendered.marked_unavailable,
    "native_formula_cells_uncomputed": uncomputed,
    "native_errors": dict(Counter(value for value in native.values() if _is_error(value))),
    "agreement": agreement(ours, native, rendered.cells),
    "compare": summarize(diffs, compared=len(rendered.cells)),
    "_ours": ours,
    "_diffs": diffs,
    "_rendered": rendered,
    "_native": native,
  }


# --------------------------------------------------------------------------
# scripted journeys (engine spec J1, J3, J5, J6, J8)
# --------------------------------------------------------------------------


_IS = "tpl.fm.income_statement"
OPEX_PARTS = (f"{_IS}.sales_and_marketing", f"{_IS}.research_and_development", f"{_IS}.general_and_administrative")
OPEX_SUBTOTAL = f"{_IS}.total_operating_expenses"
OPERATING_INCOME = f"{_IS}.operating_income"


def _known(values: Values, item: str, period: int) -> float | None:
  value = values.get((item, period))
  return value.value if value is not None and value.state == "known" else None


def _timed(fn: Callable[[], Any]) -> tuple[Any, float]:
  started = time.perf_counter()
  result = fn()
  return result, time.perf_counter() - started


def _delta(after: Values, before: Values, item: str, period: int) -> float | None:
  now, then = _known(after, item, period), _known(before, item, period)
  return None if now is None or then is None else now - then


def _projected_spec(doc: ModelSheet, item_id: str) -> dict[str, Any] | None:
  spec = doc.model.get_item(item_id).projected
  return None if spec is None else spec.model_dump(mode="json")


def _floor_matches(values: Values, floor: str, operand: str, periods: Sequence[int]) -> bool:
  """``floor`` = MAX(0, operand·0.79) wherever ``operand`` is known."""
  for period in periods:
    operand_value, floor_value = _known(values, operand, period), _known(values, floor, period)
    if operand_value is None:
      continue
    if floor_value is None or not math.isclose(floor_value, max(0.0, operand_value * (1 - 0.21))):
      return False
  return True


def run_journeys(
  load: Callable[[], ModelSheet], workdir: Path, *, native: bool = True, soffice: str | None = None,
  edit_document: Callable[[ModelSheet, Sequence[Mapping[str, Any]]], list[dict[str, Any]]] | None = None,
  inspect_document: Callable[..., list[dict[str, Any]]] | None = None,
  recalculate: Callable[[ModelSheet], Values] | None = None,
) -> dict[str, Any]:
  """Run J1, J3, J5, J6, J8 on ``load``'s latest committed document; return observations.

  Injectable edit, inspect and recalc operations default to the document functions.
  Each detached ``fork`` is an independent branch: product-door adapters must bind
  distinct documents independently and refresh edited documents from persisted state.
  Inspections receive computed values so the default does not recalculate implicitly.
  """
  edit_document = apply_ops if edit_document is None else edit_document
  inspect_document = (
    (lambda doc, values, *, items: read_lines(doc, values, items=items).lines) if inspect_document is None else inspect_document
  )
  recalculate = recalc if recalculate is None else recalculate
  workdir = Path(workdir).expanduser()
  workdir.mkdir(parents=True, exist_ok=True)
  report: dict[str, Any] = {}
  base, load_seconds = _timed(load)
  baseline, base_recalc_seconds = _timed(lambda: recalculate(base))
  ts = base.model.time_structure
  historical = [int(period) for period in ts.historical_periods]
  projection = [int(period) for period in ts.projection_periods]
  report["load"] = {"source": base.source, "seconds": load_seconds, "items": len(base.model._index),
                    "ours_recalc_seconds": base_recalc_seconds, "ours": state_counts(baseline)}

  # J1 — a new income-statement line with historical inputs, summed into the opex subtotal.
  j1 = fork(base)
  inputs = {period: 10.0 + index for index, period in enumerate(historical)}
  new_line = edit_document(j1, [{
    "op": "add_line", "section": "income_statement", "label": "Administrative depreciation",
    "unit": "USD millions", "after_item_id": OPEX_PARTS[-1], "inputs": inputs,
  }])[0]["item_id"]
  edit_document(j1, [
    {"op": "set_formula", "item_id": new_line, "formula_expr": f"{new_line}[-1]", "apply_to": "projected"},
    {"op": "set_formula", "item_id": OPEX_SUBTOTAL,
     "formula_expr": f"SUM({', '.join((*OPEX_PARTS, new_line))})", "apply_to": "both"},
  ])
  after_j1, j1_seconds = _timed(lambda: recalculate(j1))
  periods = historical + projection
  moved = {period: _delta(after_j1, baseline, OPEX_SUBTOTAL, period) for period in periods}
  report["J1"] = {
    "new_line": new_line,
    "inspection": inspect_document(j1, after_j1, items=[new_line]),
    "line_unit": [j1.model.get_item(new_line).unit, j1.model.get_item(new_line).model_scale],
    "inputs": inputs,
    "subtotal": OPEX_SUBTOTAL,
    "subtotal_delta": moved,
    "subtotal_moved_by_input": all(
      (delta := moved[period]) is not None and math.isclose(delta, inputs[period], abs_tol=1e-9)
      for period in historical
    ),
    "operating_income_delta": {period: _delta(after_j1, baseline, OPERATING_INCOME, period) for period in periods},
    "changed_cells": len(value_diff(baseline, after_j1)),
    "ours_recalc_seconds": j1_seconds,
  }

  # J3 — one formula in our grammar, one only the native engine evaluates.
  j3 = fork(j1)
  floor = edit_document(j3, [{
    "op": "add_line", "section": "income_statement", "label": "After-tax operating floor",
    "unit": "USD millions", "after_item_id": OPERATING_INCOME,
    "formula": f"MAX(0, {OPERATING_INCOME}*(1-0.21))",
  }])[0]["item_id"]
  flag = edit_document(j3, [{
    "op": "add_line", "section": "income_statement", "label": "Operating income positive flag",
    "unit": "count", "after_item_id": floor, "formula": f"IF({OPERATING_INCOME}>0, 1, 0)",
  }])[0]["item_id"]
  after_j3 = recalculate(j3)
  report["J3"] = {
    "inspection": inspect_document(j3, after_j3, items=[floor, flag]),
    "grammar_line": floor,
    "grammar_spec": _projected_spec(j3, floor),
    "grammar_values": {period: value_row(after_j3[(floor, period)]) for period in periods},
    "grammar_matches_max0_oi_x_0.79": _floor_matches(after_j3, floor, OPERATING_INCOME, periods),
    "raw_line": flag,
    "raw_spec": _projected_spec(j3, flag),
    "raw_values": {period: value_row(after_j3[(flag, period)]) for period in periods},
  }

  # J5 — a historical input removed: Missing, dependents follow; save, reopen, recalc equal.
  j5 = fork(j1)
  # The last historical year: its dependents include the carried-forward projections.
  cleared = (new_line, historical[-1])
  edit_document(j5, [{"op": "clear_input", "item_id": cleared[0], "period": cleared[1]}])
  after_j5, j5_seconds = _timed(lambda: recalculate(j5))
  turned = [row for row in value_diff(after_j1, after_j5) if row["after"] and row["after"]["state"] != "known"]
  dependents = [row for row in turned if (row["item"], row["period"]) != cleared]
  causes = Counter(tuple(row["after"].get("cause") or ()) for row in turned)
  reasons = Counter(row["after"].get("reason") for row in turned)
  removal: dict[str, Any]
  try:
    edit_document(fork(j5), [{"op": "remove_line", "item_id": new_line}])
    removal = {"status": "ok"}
  except EditRefused as refused:
    removal = {"status": "refused", "reason": refused.result.reason, "details": refused.result.details}
  saved = save(j5, workdir / "j5-document.json", values=after_j5)
  reopened = reopen(saved)
  after_reopen = recalculate(reopened)
  report["J5"] = {
    "cleared": list(cleared),
    "inspection": inspect_document(j5, after_j5, items=[new_line]),
    "cleared_value": value_row(after_j5[cleared]),
    "turned_unavailable": len(turned),
    "dependents_unavailable": len(dependents),
    "dependent_sample": [f"{row['item']}@{row['period']}" for row in dependents[:12]],
    "reasons": dict(reasons),
    "causes": {f"{cause[0]}@{cause[1]}" if cause else "": count for cause, count in causes.items()},
    "every_dependent_caused_by_cleared_input": all(tuple(row["after"].get("cause") or ()) == cleared for row in dependents),
    "required_dependency_missing_seen": reasons.get("required_dependency_missing", 0),
    "remove_line_with_dependents": removal,
    "saved": str(saved),
    "reopen_recalc_equal": not value_diff(after_j5, after_reopen),
    "ours_recalc_seconds": j5_seconds,
  }
  if native:
    j5_native = ours_vs_native(j5, workdir, "j5", soffice=soffice, recalculate=recalculate)
    cells = j5_native["_rendered"].cells
    native_values = j5_native["_native"]
    turned_keys = [(row["item"], row["period"]) for row in turned]
    shown = {
      f"{item}@{period}": {"cell": f"{cells[(item, period)].sheet}!{cells[(item, period)].cell}",
                           "formula": cells[(item, period)].formula,
                           "native": native_values.get((cells[(item, period)].sheet, cells[(item, period)].cell))}
      for item, period in turned_keys if (item, period) in cells
    }
    report["J5"]["native"] = {
      "cells": shown,
      "all_native_errors": all(_is_error(row["native"]) for row in shown.values()),
      "seconds": j5_native["seconds"],
      "compare": {key: j5_native["compare"][key] for key in ("compared_cells", "disagreements", "by_kind_family")},
    }

  # J6 — try an edit on a copy, compare with the original, discard.
  trial = fork(base)
  driver = OPEX_PARTS[0]
  last_historical = historical[-1]
  original_value = _known(baseline, driver, last_historical) or 0.0
  edit_document(trial, [{
    "op": "set_input", "item_id": driver, "period": last_historical, "value": original_value + 50.0,
  }])
  trial_values = recalculate(trial)
  trial_diff = value_diff(baseline, trial_values)
  base_after = recalculate(base)
  report["J6"] = {
    "inspection": inspect_document(trial, trial_values, items=[driver]),
    "edit": {"item": driver, "period": last_historical, "from": original_value, "to": original_value + 50.0},
    "changed_cells": len(trial_diff),
    "changed_items": sorted({row["item"] for row in trial_diff})[:25],
    "sample": trial_diff[:8],
    "discarded_original_unchanged": not value_diff(baseline, base_after),
  }

  # J8 — original and edited document: ours vs native, cell by cell.
  if native:
    original = ours_vs_native(base, workdir, "original", soffice=soffice, recalculate=recalculate)
    edited = ours_vs_native(j3, workdir, "edited", soffice=soffice, recalculate=recalculate)
    lines = _native_lines(edited["_rendered"], edited["_native"], (floor, flag), periods)
    for result in (original, edited):
      for private in ("_ours", "_diffs", "_rendered", "_native"):
        result.pop(private)
    report["J8"] = {"original": original, "edited": edited}
    report["J3"]["native"] = lines
    report["J3"]["native_compare"] = edited["compare"]["by_kind_family"]
  return report


def _native_lines(
  rendered: Rendered,
  native: Mapping[tuple[str, str], Any],
  items: Sequence[str],
  periods: Sequence[int],
) -> dict[str, Any]:
  """Each item-period's rendered cell, its formula, and the native engine's value."""
  out: dict[str, Any] = {}
  for item in items:
    out[item] = {}
    for period in periods:
      cell = rendered.cells.get((item, period))
      if cell is not None:
        out[item][period] = {"cell": f"{cell.sheet}!{cell.cell}", "formula": cell.formula,
                             "native": native.get((cell.sheet, cell.cell))}
  return out
