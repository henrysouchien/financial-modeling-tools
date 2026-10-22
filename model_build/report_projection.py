"""Deterministic report projection for one persisted build-run arm.

The projector consumes only the exact immutable ModelVersion named by the run
binding. It never reads CurrentModelPointer, a clock, or any mutable output
store. The FMS-facing dictionary is the closed server projection, which
carries the ValidationReport findings and the uncleared-warning codes; the
verdict is the ModelVersion's build status. Run output status stays
separately typed so it cannot be mistaken for workbook-derived state.

The hydrated artifacts and the persisted run arm arrive as the Protocols
below; the product's hydrate-then-project adapter is
``api/fms/subcommands/report_build_model.py::build_persisted_build_report_projection``.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Mapping
from dataclasses import dataclass
import hashlib
from io import BytesIO
from typing import Any, Literal, NoReturn, Protocol

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from schema.build_model_product_contracts import (
  ModelSpec,
  ModelVersion,
  ProvenanceMap,
  SourceBundle,
  ValidationReport,
)
from schema.finding_vocabulary import (
  FINDING_SEVERITY_MAP,
  BuildVerdict,
  error_findings,
  uncleared_warnings,
)
from schema.handoff import ModelRenderState, SeedCoverageRecord
from schema.model_semantic_document import (
  ModelSemanticDocument,
  RecordedValidationStatus,
)
from schema.render_state import render_state_hash


RUN_BINDING_INTEGRITY_ERROR = "run_binding_integrity_failed"
PERSISTED_BUILD_REPORT_PROJECTION_SCHEMA_VERSION = "1.0"
PERSISTED_BUILD_REPORT_SERVER_FIELD_KEYS = frozenset(
  {
    "annual_presentation",
    "control_snapshot_projection",
    "diagnostic_report",
    "findings",
    "model_spec",
    "normalized",
    "readiness_cross_check",
    "semantic_rows",
    "source",
    "warning_codes",
  }
)

BuildReportVerdict = Literal[
  "BUILD_CLEAN",
  "BUILD_WITH_WARNINGS",
  "BUILD_BLOCKED",
]
BuildOutputStatus = Literal[
  "pending",
  "published",
  "error",
  "skipped",
  "disabled",
  "skipped_blocked",
]
_VERDICT_BY_BUILD_STATUS: Mapping[str, BuildReportVerdict] = {
  "blocked": "BUILD_BLOCKED",
  "ready_with_warnings": "BUILD_WITH_WARNINGS",
  "ready": "BUILD_CLEAN",
}


class SourceBundleLineage(Protocol):
  """The SourceBundle identity the hydrator verified for the exact version."""

  @property
  def bundle_id(self) -> str: ...

  @property
  def identity_digest(self) -> str: ...

  @property
  def content_sha256(self) -> str: ...


class ExactBuildArtifacts(Protocol):
  """Verified bytes and decoded contracts for one exact ModelVersion."""

  @property
  def model_version(self) -> ModelVersion: ...

  @property
  def workbook_bytes(self) -> bytes: ...

  @property
  def semantic_document(self) -> ModelSemanticDocument: ...

  @property
  def provenance(self) -> ProvenanceMap: ...

  @property
  def validation_report(self) -> ValidationReport: ...

  @property
  def source_bundle(self) -> SourceBundle: ...

  @property
  def model_spec(self) -> ModelSpec: ...

  @property
  def source_bundle_lineage(self) -> SourceBundleLineage: ...


class BuildRunOutput(Protocol):
  """One run output binding (model insights or price target)."""

  def to_dict(self) -> Mapping[str, object]: ...


class PersistedBuildRun(Protocol):
  """The persisted build-run arm whose exact ModelVersion is projected."""

  @property
  def state(self) -> str: ...

  @property
  def user_id(self) -> str: ...

  @property
  def research_file_id(self) -> int: ...

  @property
  def ticker(self) -> str: ...

  @property
  def bundle_id(self) -> str: ...

  @property
  def model_version(self) -> ModelVersion: ...

  @property
  def mi(self) -> BuildRunOutput: ...

  @property
  def pt(self) -> BuildRunOutput: ...


class PersistedBuildReportProjectionError(RuntimeError):
  """Typed failure for disagreement inside one exact run projection."""

  code = RUN_BINDING_INTEGRITY_ERROR
  retryable = False

  def __init__(self, reason: str) -> None:
    super().__init__(reason)
    self.reason = reason


class _ProjectionDTO(BaseModel):
  model_config = ConfigDict(extra="forbid", frozen=True)


class NormalizedBuildProjectionV1(_ProjectionDTO):
  schema_version: Literal["1.0"] = "1.0"
  model_id: str = Field(min_length=1)
  model_version: int = Field(ge=1)
  ticker: str = Field(min_length=1)
  workbook_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
  workbook_semantics_hash: str = Field(min_length=1)
  worksheet_count: int = Field(ge=1)
  semantic_cell_count: int = Field(ge=0)
  provenance_cell_count: int = Field(ge=0)


class SourceBuildProjectionV1(_ProjectionDTO):
  schema_version: Literal["1.0"] = "1.0"
  bundle_id: str = Field(min_length=1)
  identity_digest: str = Field(pattern=r"^[0-9a-f]{64}$")
  content_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
  source_bundle_digest: str = Field(pattern=r"^[0-9a-f]{64}$")
  source_refs: tuple[str, ...] = ()
  observation_count: int = Field(ge=0)
  assumption_count: int = Field(ge=0)


class ControlCoverageProjectionV1(_ProjectionDTO):
  control_id: str = Field(min_length=1)
  binding_id: str | None = None
  scenario: str | None = None
  resolved_item_id: str = Field(min_length=1)
  periods: tuple[int, ...] = ()


class ControlSnapshotProjectionV1(_ProjectionDTO):
  schema_version: Literal["1.0"] = "1.0"
  source_authority: Literal["canonical_control_snapshot"] = (
    "canonical_control_snapshot"
  )
  available_curve_count: int = Field(ge=0)
  projected_curve_count: int = Field(ge=0)
  coverage: tuple[ControlCoverageProjectionV1, ...] = ()


class SemanticRowsProjectionV1(_ProjectionDTO):
  schema_version: Literal["1.0"] = "1.0"
  row_count: int = Field(ge=0)
  cell_count: int = Field(ge=0)
  sheets: tuple[str, ...] = ()
  lineage_status_counts: dict[str, int]


class AnnualPresentationProjectionV1(_ProjectionDTO):
  schema_version: Literal["1.0"] = "1.0"
  status: BuildVerdict
  workbook_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
  render_state_hash: str = Field(min_length=1)
  expected_axis: tuple[int, ...]
  worksheet_names: tuple[str, ...] = Field(min_length=1)
  semantic_cell_count: int = Field(ge=0)


class ReadinessCrossCheckProjectionV1(_ProjectionDTO):
  schema_version: Literal["1.0"] = "1.0"
  status: BuildVerdict
  model_identity_matches: Literal[True] = True
  workbook_hash_matches: Literal[True] = True
  semantic_validation_status: RecordedValidationStatus
  complete_lineage_cell_count: int = Field(ge=0)
  incomplete_lineage_cell_count: int = Field(ge=0)
  blocking_codes: tuple[str, ...] = ()
  warning_codes: tuple[str, ...] = ()


class PersistedBuildReportServerProjectionV1(_ProjectionDTO):
  """The exact closed server field set accepted by run-bound FMS reports."""

  annual_presentation: AnnualPresentationProjectionV1
  control_snapshot_projection: ControlSnapshotProjectionV1
  diagnostic_report: dict[str, Any] | None
  #: ValidationReport.findings, in report order.
  findings: tuple[dict[str, Any], ...]
  model_spec: dict[str, Any]
  normalized: NormalizedBuildProjectionV1
  readiness_cross_check: ReadinessCrossCheckProjectionV1
  semantic_rows: SemanticRowsProjectionV1
  source: SourceBuildProjectionV1
  warning_codes: tuple[str, ...]


class BuildOutputProjectionV1(_ProjectionDTO):
  status: BuildOutputStatus
  error_code: str | None = Field(default=None, min_length=1)
  error_detail: str | None = Field(default=None, min_length=1)
  #: ``warning`` when ``error_code`` is a warning Finding code (a derivation
  #: the version publishes without), absent otherwise.
  severity: Literal["warning"] | None = None

  @model_validator(mode="after")
  def _error_code_shape(self) -> "BuildOutputProjectionV1":
    if (self.status == "error") != (self.error_code is not None):
      raise ValueError("error_code is required exactly for error status")
    if self.error_detail is not None and self.status != "error":
      raise ValueError("error_detail is only allowed for error status")
    rule = (
      None if self.error_code is None else FINDING_SEVERITY_MAP.get(self.error_code)
    )
    warns = rule is not None and rule.severity == "warning"
    if (self.severity == "warning") != warns:
      raise ValueError("severity is warning exactly for a warning finding code")
    return self


class PersistedBuildReportProjectionV1(_ProjectionDTO):
  """One complete projection result before FMS presentation composition."""

  schema_version: Literal["1.0"] = "1.0"
  verdict: BuildReportVerdict
  model_insights_status: BuildOutputProjectionV1
  price_target_status: BuildOutputProjectionV1
  server: PersistedBuildReportServerProjectionV1

  def server_projection(self) -> dict[str, Any]:
    payload = self.server.model_dump(mode="json")
    if frozenset(payload) != PERSISTED_BUILD_REPORT_SERVER_FIELD_KEYS:
      raise PersistedBuildReportProjectionError(
        "server projection field set differs from its versioned contract"
      )
    return payload


@dataclass(frozen=True, slots=True)
class _WorkbookFacts:
  sha256: str
  worksheet_names: tuple[str, ...]


def _integrity(reason: str) -> NoReturn:
  raise PersistedBuildReportProjectionError(reason)


def build_report_verdict(status: str) -> BuildReportVerdict:
  """The build-report verdict for one ModelVersion build status."""

  verdict = _VERDICT_BY_BUILD_STATUS.get(status)
  if verdict is None:
    _integrity(f"build status {status!r} has no report verdict")
  return verdict



def _json_object(value: object, *, field: str) -> dict[str, Any]:
  if isinstance(value, BaseModel):
    payload = value.model_dump(mode="json")
  elif isinstance(value, Mapping):
    payload = dict(value)
  else:
    _integrity(f"{field} is not a JSON object")
  if not isinstance(payload, dict):
    _integrity(f"{field} is not a JSON object")
  return payload


def _workbook_facts(workbook_bytes: bytes) -> _WorkbookFacts:
  if type(workbook_bytes) is not bytes or not workbook_bytes:
    _integrity("exact ModelVersion workbook bytes are absent")
  digest = hashlib.sha256(workbook_bytes).hexdigest()
  try:
    workbook = load_workbook(
      BytesIO(workbook_bytes),
      read_only=True,
      data_only=False,
      keep_links=False,
    )
  except Exception as exc:
    raise PersistedBuildReportProjectionError(
      "exact ModelVersion workbook bytes are not a readable XLSX"
    ) from exc
  try:
    worksheet_names = tuple(workbook.sheetnames)
  finally:
    workbook.close()
  if not worksheet_names or len(worksheet_names) != len(set(worksheet_names)):
    _integrity("exact ModelVersion workbook has an invalid worksheet set")
  return _WorkbookFacts(sha256=digest, worksheet_names=worksheet_names)


def _validate_exact_binding(
  artifacts: ExactBuildArtifacts,
  resolution: PersistedBuildRun,
  workbook: _WorkbookFacts,
) -> None:
  if resolution.state not in {"persisted_pending_outputs", "persisted_final"}:
    _integrity("report projection run arm has an invalid persisted state")
  version = artifacts.model_version
  semantic = artifacts.semantic_document
  validation = artifacts.validation_report
  provenance = artifacts.provenance
  lineage = artifacts.source_bundle_lineage
  if version != resolution.model_version:
    _integrity("hydrated ModelVersion differs from the exact run binding")
  if (
    version.user_id,
    version.research_file_id,
    version.ticker,
  ) != (
    resolution.user_id,
    resolution.research_file_id,
    resolution.ticker,
  ):
    _integrity("hydrated ModelVersion scope differs from the exact run binding")
  if lineage.bundle_id != resolution.bundle_id:
    _integrity("hydrated SourceBundle differs from the exact run binding")
  if workbook.sha256 != version.workbook.sha256:
    _integrity("verified workbook bytes disagree with the ModelVersion digest")
  if getattr(validation, "schema_version", None) != "2.2":
    _integrity("persisted report requires ValidationReport 2.2")
  if (
    semantic.model_id != version.model_id
    or semantic.model_version != version.version
    or semantic.ticker != version.ticker
    or artifacts.source_bundle.ticker != version.ticker
    or artifacts.model_spec.ticker != version.ticker
  ):
    _integrity("hydrated artifacts disagree on exact model identity")
  if (
    semantic.validation_status != validation.status
    or semantic.validation_blocker_count != validation.error_count
    or semantic.validation_warning_count != validation.uncleared_warning_count
    or version.status != validation.status
  ):
    _integrity("ModelVersion, semantic, and ValidationReport status disagree")
  if (
    provenance.source_bundle_digest != version.source_bundle_digest
    or provenance.model_spec_digest != version.model_spec_digest
  ):
    _integrity("provenance digests disagree with the exact ModelVersion")
  source_input = semantic.render_state.compiler_inputs.source_bundle
  if (
    source_input.bundle_id != lineage.bundle_id
    or source_input.identity_digest != lineage.identity_digest
  ):
    _integrity("semantic SourceBundle lineage disagrees with the run binding")
  semantic_sheets = {value.sheet for value in semantic.semantic_values}
  missing_sheets = sorted(semantic_sheets - set(workbook.worksheet_names))
  if missing_sheets:
    _integrity(
      "semantic cells name worksheets absent from exact workbook bytes: "
      f"{missing_sheets!r}"
    )


def _source_projection(
  artifacts: ExactBuildArtifacts,
) -> SourceBuildProjectionV1:
  provenance = artifacts.provenance
  cells = tuple(provenance.cells)
  return SourceBuildProjectionV1(
    bundle_id=artifacts.source_bundle_lineage.bundle_id,
    identity_digest=artifacts.source_bundle_lineage.identity_digest,
    content_sha256=artifacts.source_bundle_lineage.content_sha256,
    source_bundle_digest=provenance.source_bundle_digest,
    source_refs=tuple(
      sorted({str(ref) for cell in cells for ref in cell.source_refs})
    ),
    observation_count=len(
      {item for cell in cells for item in cell.observation_ids}
    ),
    assumption_count=len({item for cell in cells for item in cell.assumption_ids}),
  )


def control_snapshot_projection(
  render_state: ModelRenderState,
) -> ControlSnapshotProjectionV1:
  """The canonical control projection a version's render state records.

  The build report and the hosted ``model_summarize`` read both derive it
  here, so a reader of an exact version never depends on a stored copy.
  """

  records = tuple(render_state.seed_coverage_manifest)
  coverage = tuple(
    sorted(
      (
        ControlCoverageProjectionV1(
          control_id=record.control_id,
          binding_id=record.binding_id,
          scenario=record.scenario,
          resolved_item_id=record.resolved_item_id,
          periods=tuple(sorted(set(record.periods))),
        )
        for record in records
        if isinstance(record, SeedCoverageRecord)
        and record.control_id is not None
      ),
      key=lambda item: (
        item.control_id,
        item.binding_id or "",
        item.scenario or "",
        item.periods,
      ),
    )
  )
  curve_identities = {
    (item.control_id, item.binding_id, item.scenario) for item in coverage
  }
  return ControlSnapshotProjectionV1(
    available_curve_count=len(curve_identities),
    projected_curve_count=len(curve_identities),
    coverage=coverage,
  )


def _semantic_rows_projection(
  artifacts: ExactBuildArtifacts,
) -> SemanticRowsProjectionV1:
  values = tuple(artifacts.semantic_document.semantic_values)
  lineage_counts = Counter(cell.lineage_status for cell in artifacts.provenance.cells)
  return SemanticRowsProjectionV1(
    row_count=len({value.semantic_item_id for value in values}),
    cell_count=len(values),
    sheets=tuple(sorted({value.sheet for value in values})),
    lineage_status_counts=dict(sorted(lineage_counts.items())),
  )


def _axis(artifacts: ExactBuildArtifacts) -> tuple[int, ...]:
  time_structure = artifacts.semantic_document.model.time_structure
  historical = tuple(
    time_structure.historical_periods or time_structure.historical_years
  )
  projections = tuple(
    time_structure.projection_periods or time_structure.projection_years
  )
  return tuple(dict.fromkeys((*historical, *projections)))


def _readiness_projection(
  artifacts: ExactBuildArtifacts,
) -> ReadinessCrossCheckProjectionV1:
  counts = Counter(cell.lineage_status for cell in artifacts.provenance.cells)
  findings = artifacts.validation_report.findings
  return ReadinessCrossCheckProjectionV1(
    status=artifacts.validation_report.status,
    semantic_validation_status=artifacts.semantic_document.validation_status,
    complete_lineage_cell_count=(
      counts.get("complete", 0) + counts.get("compiler_derived", 0)
    ),
    incomplete_lineage_cell_count=counts.get("missing_source", 0),
    blocking_codes=_codes(error_findings(findings)),
    warning_codes=_codes(uncleared_warnings(findings)),
  )


def _codes(findings: tuple[Any, ...]) -> tuple[str, ...]:
  return tuple(sorted({finding.code for finding in findings}))


def _output_projection(
  value: BuildRunOutput,
  *,
  field: str,
) -> BuildOutputProjectionV1:
  try:
    payload = value.to_dict()
  except AttributeError:
    _integrity(f"{field} run output binding is malformed")
  else:
    try:
      return BuildOutputProjectionV1.model_validate(payload)
    except (TypeError, ValueError) as exc:
      raise PersistedBuildReportProjectionError(
        f"{field} run output binding is malformed"
      ) from exc


def _project_exact_persisted_build_report(
  artifacts: ExactBuildArtifacts,
  resolution: PersistedBuildRun,
) -> PersistedBuildReportProjectionV1:
  workbook = _workbook_facts(artifacts.workbook_bytes)
  _validate_exact_binding(artifacts, resolution, workbook)
  semantic = artifacts.semantic_document
  validation = artifacts.validation_report
  provenance = artifacts.provenance
  readiness = _readiness_projection(artifacts)
  mi = _output_projection(resolution.mi, field="model_insights_status")
  pt = _output_projection(resolution.pt, field="price_target_status")
  if resolution.state == "persisted_final" and (
    mi.status == "pending" or pt.status == "pending"
  ):
    _integrity("final run binding carries a pending output status")
  if resolution.state == "persisted_pending_outputs" and (
    mi.status in {"published", "error"} or pt.status in {"published", "error"}
  ):
    _integrity("pending run binding carries a terminal output status")
  annual = AnnualPresentationProjectionV1(
    status=validation.status,
    workbook_sha256=workbook.sha256,
    render_state_hash=render_state_hash(semantic.render_state),
    expected_axis=_axis(artifacts),
    worksheet_names=workbook.worksheet_names,
    semantic_cell_count=len(semantic.semantic_values),
  )
  server = PersistedBuildReportServerProjectionV1(
    annual_presentation=annual,
    control_snapshot_projection=control_snapshot_projection(
      artifacts.semantic_document.render_state
    ),
    diagnostic_report=(
      None
      if validation.diagnostic_report is None
      else _json_object(validation.diagnostic_report, field="diagnostic_report")
    ),
    findings=tuple(
      finding.model_dump(mode="json") for finding in validation.findings
    ),
    model_spec=_json_object(artifacts.model_spec, field="model_spec"),
    normalized=NormalizedBuildProjectionV1(
      model_id=artifacts.model_version.model_id,
      model_version=artifacts.model_version.version,
      ticker=artifacts.model_version.ticker,
      workbook_sha256=workbook.sha256,
      workbook_semantics_hash=semantic.render_state.workbook_semantics_hash,
      worksheet_count=len(workbook.worksheet_names),
      semantic_cell_count=len(semantic.semantic_values),
      provenance_cell_count=len(provenance.cells),
    ),
    readiness_cross_check=readiness,
    semantic_rows=_semantic_rows_projection(artifacts),
    source=_source_projection(artifacts),
    # An output a materializer could not derive is a warning on the
    # report, beside the version's own uncleared warnings.
    warning_codes=tuple(
      sorted(
        {
          *readiness.warning_codes,
          *(
            output.error_code
            for output in (mi, pt)
            if output.severity == "warning" and output.error_code
          ),
        }
      )
    ),
  )
  return PersistedBuildReportProjectionV1(
    verdict=build_report_verdict(artifacts.model_version.status),
    model_insights_status=mi,
    price_target_status=pt,
    server=server,
  )


def project_exact_persisted_build_report(
  artifacts: ExactBuildArtifacts,
  resolution: PersistedBuildRun,
) -> PersistedBuildReportProjectionV1:
  """Purely project one already-hydrated exact persisted build-run arm."""

  try:
    return _project_exact_persisted_build_report(artifacts, resolution)
  except PersistedBuildReportProjectionError:
    raise
  except (AttributeError, KeyError, TypeError, ValueError) as exc:
    raise PersistedBuildReportProjectionError(
      "exact persisted build artifacts cannot form a complete report projection"
    ) from exc


__all__ = [
  "AnnualPresentationProjectionV1",
  "BuildOutputProjectionV1",
  "BuildRunOutput",
  "ControlSnapshotProjectionV1",
  "ExactBuildArtifacts",
  "NormalizedBuildProjectionV1",
  "PERSISTED_BUILD_REPORT_PROJECTION_SCHEMA_VERSION",
  "PERSISTED_BUILD_REPORT_SERVER_FIELD_KEYS",
  "PersistedBuildReportProjectionError",
  "PersistedBuildReportProjectionV1",
  "PersistedBuildReportServerProjectionV1",
  "PersistedBuildRun",
  "RUN_BINDING_INTEGRITY_ERROR",
  "ReadinessCrossCheckProjectionV1",
  "SemanticRowsProjectionV1",
  "SourceBuildProjectionV1",
  "SourceBundleLineage",
  "build_report_verdict",
  "control_snapshot_projection",
  "project_exact_persisted_build_report",
]
